"""Read Pioneers HTML-based ``.xls`` exports and convert them to ``.xlsx``.

Original Pioneers ``گزارش لیست کالا.xls`` files are UTF-8 HTML documents
(``<Html><Body><Table>…``) saved with an ``.xls`` extension — not OLE binary
workbooks. This module parses that table and writes a single-sheet OpenXML
workbook that the existing ``PioneersExcelParser`` / writer can consume.

The source file is never modified. Presentation-only HTML attributes (fonts,
borders, background colors, ``dir`` / ``align``) are not preserved in the
``.xlsx`` output — cell values, column order, and row order are.

Cell text is preserved exactly (no trimming). Header matching normalizes with
``str.strip()`` only for lookup. Nested tables, ``colspan`` / ``rowspan``
other than 1, short/long rows, and cells over Excel's 32,767-character limit
are rejected with a clear error rather than silently reshaping data.
"""

from __future__ import annotations

from html.parser import HTMLParser
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

from gheymatyar.parsers.errors import PioneersExcelError
from gheymatyar.parsers.pioneers_excel_parser import (
    HEADER_CODE,
    HEADER_NAME,
    PRODUCT_SHEET_NAME,
)
from gheymatyar.parsers.workbook_format import (
    WorkbookFormat,
    detect_workbook_format,
    unsupported_format_message,
)

# Excel cell character limit (shared by .xlsx).
_MAX_EXCEL_CELL_CHARS = 32767


class _HtmlTableStructureError(Exception):
    """Internal parse abort for unsupported HTML table structure."""

    def __init__(self, message: str, *, row: int | None = None) -> None:
        self.row = row
        super().__init__(message)


class _TableCollector(HTMLParser):
    """Collect ``<tr>`` / ``<td>`` text from a single flat HTML table.

    Rejects nested tables, multiple top-level tables, and non-1 span attributes
    so column alignment cannot shift silently.
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._parts: list[str] | None = None
        self._in_td = False
        self._saw_table = False
        self._table_depth = 0
        self._finished_table = False
        self._row_number = 0  # 1-based among collected rows (header = 1)

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        name = tag.lower()
        if name == "table":
            if self._finished_table or (self._saw_table and self._table_depth == 0):
                raise _HtmlTableStructureError(
                    "Pioneers HTML .xls contains more than one <Table>"
                )
            if self._table_depth >= 1:
                raise _HtmlTableStructureError(
                    "Pioneers HTML .xls contains a nested <Table>",
                    row=self._row_number or None,
                )
            self._table_depth += 1
            self._saw_table = True
            return
        if self._table_depth != 1:
            return
        if name == "tr":
            if self._row is not None:
                raise _HtmlTableStructureError(
                    "Pioneers HTML .xls has a nested or unclosed <tr>",
                    row=self._row_number or None,
                )
            self._row = []
            self._row_number += 1
            return
        if name == "td" and self._row is not None:
            if self._in_td:
                raise _HtmlTableStructureError(
                    "Pioneers HTML .xls has a nested or unclosed <td>",
                    row=self._row_number,
                )
            self._reject_span_attrs(attrs, row=self._row_number)
            self._in_td = True
            self._parts = []
            return
        if name == "th":
            raise _HtmlTableStructureError(
                "Pioneers HTML .xls uses unsupported <th> cells "
                "(expected <td> only)",
                row=self._row_number or None,
            )

    def handle_endtag(self, tag: str) -> None:
        name = tag.lower()
        if name == "table":
            if self._table_depth == 1:
                if self._row is not None or self._in_td:
                    raise _HtmlTableStructureError(
                        "Pioneers HTML .xls closed <Table> with an open row/cell",
                        row=self._row_number or None,
                    )
                self._finished_table = True
            if self._table_depth > 0:
                self._table_depth -= 1
            return
        if self._table_depth != 1:
            return
        if name == "td":
            if self._row is None or not self._in_td:
                raise _HtmlTableStructureError(
                    "Pioneers HTML .xls has a stray </td>",
                    row=self._row_number or None,
                )
            # Preserve exact text — no strip().
            self._row.append("".join(self._parts or []))
            self._in_td = False
            self._parts = None
            return
        if name == "tr":
            if self._row is None:
                raise _HtmlTableStructureError(
                    "Pioneers HTML .xls has a stray </tr>",
                    row=self._row_number or None,
                )
            if self._in_td:
                raise _HtmlTableStructureError(
                    "Pioneers HTML .xls has an unclosed <td> in a row",
                    row=self._row_number,
                )
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        if self._in_td and self._parts is not None:
            self._parts.append(data)

    @staticmethod
    def _reject_span_attrs(
        attrs: list[tuple[str, str | None]], *, row: int
    ) -> None:
        mapping = {key.lower(): (value or "") for key, value in attrs}
        for attr in ("colspan", "rowspan"):
            raw = mapping.get(attr)
            if raw is None or raw == "":
                continue
            if raw.strip() != "1":
                raise _HtmlTableStructureError(
                    f"Pioneers HTML .xls has unsupported {attr}={raw!r} "
                    f"(only {attr}=1 is accepted)",
                    row=row,
                )


def read_pioneers_html_xls_rows(path: str | Path) -> list[list[str]]:
    """Parse a Pioneers HTML ``.xls`` into a matrix of exact cell strings.

    Row 0 is the header row. Empty cells are ``""``. Leading/trailing spaces
    are preserved. Every data row must have exactly as many cells as the header.
    """
    source = Path(path)
    if not source.is_file():
        raise PioneersExcelError(
            "Pioneers HTML Excel file not found",
            filename=source,
            reason="path does not exist or is not a file",
        )

    raw = source.read_bytes()
    fmt = detect_workbook_format(raw)
    if fmt != WorkbookFormat.PIONEERS_HTML_XLS:
        raise PioneersExcelError(
            "Not a Pioneers HTML-based .xls export",
            filename=source,
            reason=unsupported_format_message(raw),
        )

    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise PioneersExcelError(
            "Pioneers HTML .xls is not valid UTF-8",
            filename=source,
            reason=str(exc),
        ) from exc

    collector = _TableCollector()
    try:
        collector.feed(text)
        collector.close()
    except _HtmlTableStructureError as exc:
        raise PioneersExcelError(
            str(exc),
            filename=source,
            row=exc.row,
            reason="unsupported HTML table structure",
        ) from exc
    except Exception as exc:  # noqa: BLE001 - wrap html.parser failures
        raise PioneersExcelError(
            "Failed to parse Pioneers HTML .xls table",
            filename=source,
            reason=str(exc),
        ) from exc

    if collector._row is not None or collector._in_td:
        raise PioneersExcelError(
            "Pioneers HTML .xls ended with an unclosed row or cell",
            filename=source,
            row=collector._row_number or None,
            reason="unsupported HTML table structure",
        )

    if not collector._saw_table or not collector.rows:
        raise PioneersExcelError(
            "Pioneers HTML .xls contains no table rows",
            filename=source,
            reason="expected a <Table> with a header row",
        )

    rows = collector.rows
    headers = rows[0]
    normalized_headers = [header.strip() for header in headers]
    if HEADER_CODE not in normalized_headers or HEADER_NAME not in normalized_headers:
        raise PioneersExcelError(
            "Pioneers HTML .xls is missing required product headers",
            filename=source,
            reason=f"required: {HEADER_CODE!r}, {HEADER_NAME!r}",
        )

    width = len(headers)
    if width == 0:
        raise PioneersExcelError(
            "Pioneers HTML .xls header row is empty",
            filename=source,
        )

    for index, row in enumerate(rows, start=1):
        if len(row) != width:
            raise PioneersExcelError(
                "Pioneers HTML .xls row column count does not match the header",
                filename=source,
                row=index,
                reason=f"expected {width} columns, found {len(row)}",
            )
        for col_index, value in enumerate(row, start=1):
            if len(value) > _MAX_EXCEL_CELL_CHARS:
                raise PioneersExcelError(
                    "Pioneers HTML .xls cell exceeds Excel's 32767-character limit",
                    filename=source,
                    row=index,
                    column=get_column_letter(col_index),
                    reason=f"length {len(value)}",
                )
    return rows


def convert_pioneers_html_xls_to_xlsx(
    source_path: str | Path,
    output_path: str | Path,
) -> Path:
    """Convert an HTML-based Pioneers ``.xls`` into a single-sheet ``.xlsx``.

    Every non-empty value is written as an explicit shared-string / text cell so
    values like ``=1+1`` or ``#N/A`` stay literal. Empty cells stay blank.
    The product sheet title is ``لیست کالا``.
    """
    source = Path(source_path)
    output = Path(output_path)
    if source.resolve() == output.resolve():
        raise PioneersExcelError(
            "Refusing to overwrite the source Pioneers HTML .xls",
            filename=source,
            reason="output_path must differ from source_path",
        )

    rows = read_pioneers_html_xls_rows(source)
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook(write_only=True)
    try:
        worksheet = workbook.create_sheet(title=PRODUCT_SHEET_NAME)
        for row in rows:
            worksheet.append([_text_cell(worksheet, value) for value in row])
        workbook.save(output)
    except PioneersExcelError:
        raise
    except Exception as exc:  # noqa: BLE001 - wrap openpyxl failures
        raise PioneersExcelError(
            "Failed to write converted Pioneers .xlsx workbook",
            filename=output,
            reason=str(exc),
        ) from exc
    finally:
        workbook.close()

    return output


def _text_cell(worksheet: object, value: str) -> WriteOnlyCell | None:
    """Return a write-only text cell, or ``None`` for a truly empty value."""
    if value == "":
        return None
    cell = WriteOnlyCell(worksheet, value=value)  # type: ignore[arg-type]
    # Force literal text — openpyxl otherwise treats leading '=' as a formula
    # and values like '#N/A' as Excel errors.
    cell.data_type = "s"
    return cell


def html_xls_preservation_notes() -> Sequence[str]:
    """Source-format features that the HTML→XLSX conversion cannot keep."""
    return (
        "Per-cell HTML styles (B Yekan font, borders, background colors)",
        "HTML dir=rtl / align attributes on <td> elements",
        "The HTML document wrapper itself (output is OpenXML .xlsx)",
        "Original .xls filename extension (working copy is always .xlsx)",
        "Nested tables, colspan/rowspan other than 1 (rejected, not flattened)",
    )
