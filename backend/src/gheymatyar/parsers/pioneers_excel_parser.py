"""Read-only parser for Pioneers product-list Excel exports.

Based on the inspected workbook ``گزارش لیست کالا - converted.xlsx``.
This module never saves or mutates the source workbook.
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from gheymatyar.domain.product import PioneersProduct
from gheymatyar.parsers.errors import (
    AmbiguousMappedHeaderError,
    InvalidProductRowError,
    MissingRequiredColumnError,
    PioneersExcelError,
    SheetNotFoundError,
)

# Exact headers from the real Pioneers export (see docs/pioneers_excel_schema.md).
PRODUCT_SHEET_NAME = "لیست کالا"
HEADER_ROW = 1

HEADER_CODE = "کد کالا"
HEADER_NAME = "نـــام کــالا"  # includes tatweel (U+0640)
HEADER_BARCODE = "بارکد"
HEADER_SALE_PRICE = "نرخ فروش اصلی"
HEADER_LAST_PURCHASE_PRICE = "نرخ آخرین خرید"
HEADER_CONSUMER_PRICE = "نرخ مصرف کننده"
HEADER_COLLEAGUE_PRICE = "نرخ همکاری"

REQUIRED_HEADERS: tuple[str, ...] = (HEADER_CODE, HEADER_NAME)

OPTIONAL_FIELD_HEADERS: Mapping[str, str] = {
    "barcode": HEADER_BARCODE,
    "sale_price": HEADER_SALE_PRICE,
    "last_purchase_price": HEADER_LAST_PURCHASE_PRICE,
    "consumer_price": HEADER_CONSUMER_PRICE,
    "colleague_price": HEADER_COLLEAGUE_PRICE,
}

# Headers whose column must be unambiguous (duplicates are fatal).
MAPPED_HEADERS: frozenset[str] = frozenset(
    (*REQUIRED_HEADERS, *OPTIONAL_FIELD_HEADERS.values())
)

_EXCEL_ERROR_VALUES = frozenset(
    {
        "#NULL!",
        "#DIV/0!",
        "#VALUE!",
        "#REF!",
        "#NAME?",
        "#NUM!",
        "#N/A",
        "#GETTING_DATA",
        "#SPILL!",
        "#CALC!",
        "#FIELD!",
        "#BLOCKED!",
        "#UNKNOWN!",
        "#CONNECT!",
    }
)

# Only a single-section zero-pad format is supported (e.g. "000000").
# Multi-section formats like '000000;000000;"ZERO"' are rejected.
_ZERO_PAD_FORMAT_RE = re.compile(r"^0+$")


@dataclass(frozen=True)
class InvalidProductRow:
    """A product-like Excel row that could not become a PioneersProduct.

    Kept for later review UI — parsing continues for remaining rows.
    """

    row_number: int
    reason: str
    column: str | None = None
    code: str | None = None
    name: str | None = None
    raw_values: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class PioneersParseResult:
    """Products plus observable parse diagnostics."""

    products: list[PioneersProduct]
    invalid_rows: list[InvalidProductRow] = field(default_factory=list)
    duplicate_codes: dict[str, tuple[int, ...]] = field(default_factory=dict)

    def __iter__(self):  # type: ignore[no-untyped-def]
        return iter(self.products)

    def __len__(self) -> int:
        return len(self.products)


@dataclass(frozen=True)
class _CellSnapshot:
    value: Any
    data_type: str
    number_format: str


class PioneersExcelParser:
    """Parse a Pioneers product workbook into ``PioneersProduct`` rows."""

    def parse(self, path: str | Path) -> PioneersParseResult:
        source = Path(path)
        if not source.is_file():
            raise PioneersExcelError(
                "Pioneers Excel file not found",
                filename=source,
                reason="path does not exist or is not a file",
            )

        workbook = None
        try:
            try:
                # data_only=False keeps stored values; we never save.
                workbook = load_workbook(
                    filename=source,
                    read_only=True,
                    data_only=False,
                    keep_vba=False,
                )
            except PioneersExcelError:
                raise
            except Exception as exc:  # noqa: BLE001 - wrap all openpyxl failures
                raise PioneersExcelError(
                    "Failed to open Pioneers Excel workbook",
                    filename=source,
                    reason=str(exc),
                ) from exc

            sheet_name: str | None = None
            try:
                worksheet = self._resolve_sheet(workbook, source)
                sheet_name = worksheet.title
                header_map = self._read_header_map(worksheet, source)
                self._validate_required_headers(
                    header_map, source, worksheet.title
                )
                products, invalid_rows = self._parse_rows(worksheet, header_map)
            except PioneersExcelError:
                raise
            except Exception as exc:  # noqa: BLE001 - wrap deferred openpyxl/XML failures
                raise PioneersExcelError(
                    "Failed to read Pioneers Excel worksheet",
                    filename=source,
                    sheet=sheet_name or PRODUCT_SHEET_NAME,
                    reason=str(exc),
                ) from exc

            duplicate_codes = self._find_duplicate_codes(products)
            return PioneersParseResult(
                products=products,
                invalid_rows=invalid_rows,
                duplicate_codes=duplicate_codes,
            )
        finally:
            if workbook is not None:
                workbook.close()

    def _resolve_sheet(self, workbook: Any, source: Path) -> Worksheet:
        """Select the product sheet from required headers, not from its title.

        Production exports use ``لیست کالا``. A development sample may use
        another title. Exactly one worksheet must contain the required headers.
        """
        matches: list[Worksheet] = []
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            if self._sheet_has_required_headers(worksheet):
                matches.append(worksheet)
        if len(matches) == 1:
            return matches[0]
        required = ", ".join(REQUIRED_HEADERS)
        available = ", ".join(workbook.sheetnames) or "(none)"
        if len(matches) > 1:
            found = ", ".join(worksheet.title for worksheet in matches)
            raise PioneersExcelError(
                "More than one sheet matches the Pioneers product headers",
                filename=source,
                reason=(
                    f"matching sheets: {found}. "
                    f"required on row {HEADER_ROW}: {required}"
                ),
            )
        raise SheetNotFoundError(
            "No sheet matches the Pioneers product headers",
            filename=source,
            reason=(
                f"available sheets: {available}. "
                f"required on row {HEADER_ROW}: {required}"
            ),
        )

    @staticmethod
    def _sheet_has_required_headers(worksheet: Worksheet) -> bool:
        found: set[str] = set()
        header_row = worksheet[HEADER_ROW]
        for cell in header_row:
            value = getattr(cell, "value", None)
            if isinstance(value, str):
                header = value.strip()
                if header in REQUIRED_HEADERS:
                    found.add(header)
        return all(header in found for header in REQUIRED_HEADERS)

    def _read_header_map(
        self, worksheet: Worksheet, source: Path
    ) -> dict[str, int]:
        """Map exact header text -> 1-based column index.

        Unmapped duplicate headers (e.g. ``مرکز پخش``) keep the first column.
        Mapped headers (code/name/barcode/prices) must appear exactly once.
        """
        header_map: dict[str, int] = {}
        mapped_columns: dict[str, list[str]] = defaultdict(list)

        for cell in worksheet[HEADER_ROW]:
            value = cell.value
            if value is None or not isinstance(value, str):
                continue
            header = value.strip()
            if not header:
                continue

            column_letter = get_column_letter(
                getattr(cell, "column", None) or 1
            )
            if header in MAPPED_HEADERS:
                mapped_columns[header].append(column_letter)

            # First wins for unmapped duplicates such as مرکز پخش.
            header_map.setdefault(header, cell.column)

        if not header_map:
            raise PioneersExcelError(
                "Pioneers header row is empty",
                filename=source,
                sheet=worksheet.title,
                row=HEADER_ROW,
                reason="no header cells found",
            )

        for header, columns in mapped_columns.items():
            if len(columns) > 1:
                raise AmbiguousMappedHeaderError(
                    f"Ambiguous mapped Pioneers column:\n{header}",
                    filename=source,
                    sheet=worksheet.title,
                    row=HEADER_ROW,
                    column=header,
                    reason=f"appears in columns {', '.join(columns)}",
                )

        return header_map

    def _validate_required_headers(
        self,
        header_map: Mapping[str, int],
        source: Path,
        sheet: str,
    ) -> None:
        for header in REQUIRED_HEADERS:
            if header not in header_map:
                raise MissingRequiredColumnError(
                    f"Missing required Pioneers column:\n{header}",
                    filename=source,
                    sheet=sheet,
                    column=header,
                    reason="required header not found in row 1",
                )

    def _parse_rows(
        self,
        worksheet: Worksheet,
        header_map: Mapping[str, int],
    ) -> tuple[list[PioneersProduct], list[InvalidProductRow]]:
        code_col = header_map[HEADER_CODE]
        name_col = header_map[HEADER_NAME]
        optional_cols = {
            field_name: header_map[header]
            for field_name, header in OPTIONAL_FIELD_HEADERS.items()
            if header in header_map
        }

        products: list[PioneersProduct] = []
        invalid_rows: list[InvalidProductRow] = []
        # enumerate() supplies Excel row numbers: read_only EmptyCell has no .row
        for row_number, row_cells in enumerate(
            worksheet.iter_rows(min_row=HEADER_ROW + 1),
            start=HEADER_ROW + 1,
        ):
            if not row_cells:
                continue
            cells_by_col, raw_values = self._extract_row_cells(row_cells)

            try:
                code = self._parse_code_or_barcode(
                    cells_by_col.get(code_col),
                    column=HEADER_CODE,
                )
                name = self._parse_name(cells_by_col.get(name_col))
            except (ValueError, InvalidProductRowError) as exc:
                column = (
                    getattr(exc, "column", None)
                    if isinstance(exc, InvalidProductRowError)
                    else None
                )
                reason = (
                    exc.reason
                    if isinstance(exc, InvalidProductRowError) and exc.reason
                    else str(exc)
                )
                invalid_rows.append(
                    InvalidProductRow(
                        row_number=row_number,
                        reason=reason,
                        column=column,
                        raw_values=raw_values,
                    )
                )
                continue

            if code is None and name is None:
                if self._row_is_entirely_empty(
                    (snap.value for snap in cells_by_col.values())
                ):
                    continue
                # Non-empty junk without identity fields — keep for review.
                invalid_rows.append(
                    InvalidProductRow(
                        row_number=row_number,
                        reason=(
                            "row has data but both کد کالا and "
                            "نـــام کــالا are empty"
                        ),
                        raw_values=raw_values,
                    )
                )
                continue

            if code is None:
                invalid_rows.append(
                    InvalidProductRow(
                        row_number=row_number,
                        reason="product-like row has empty کد کالا",
                        column=HEADER_CODE,
                        code=None,
                        name=name,
                        raw_values=raw_values,
                    )
                )
                continue
            if name is None:
                invalid_rows.append(
                    InvalidProductRow(
                        row_number=row_number,
                        reason="product-like row has empty نـــام کــالا",
                        column=HEADER_NAME,
                        code=code,
                        name=None,
                        raw_values=raw_values,
                    )
                )
                continue

            try:
                product = PioneersProduct(
                    row_number=row_number,
                    code=code,
                    name=name,
                    barcode=self._parse_code_or_barcode(
                        cells_by_col.get(optional_cols["barcode"])
                        if "barcode" in optional_cols
                        else None,
                        column=HEADER_BARCODE,
                        optional=True,
                    ),
                    unit=None,  # no dedicated unit column in this workbook
                    sale_price=self._parse_optional_price(
                        cells_by_col.get(optional_cols["sale_price"]).value
                        if "sale_price" in optional_cols
                        and optional_cols["sale_price"] in cells_by_col
                        else None,
                        column=HEADER_SALE_PRICE,
                    ),
                    last_purchase_price=self._parse_optional_price(
                        cells_by_col.get(
                            optional_cols["last_purchase_price"]
                        ).value
                        if "last_purchase_price" in optional_cols
                        and optional_cols["last_purchase_price"] in cells_by_col
                        else None,
                        column=HEADER_LAST_PURCHASE_PRICE,
                    ),
                    consumer_price=self._parse_optional_price(
                        cells_by_col.get(optional_cols["consumer_price"]).value
                        if "consumer_price" in optional_cols
                        and optional_cols["consumer_price"] in cells_by_col
                        else None,
                        column=HEADER_CONSUMER_PRICE,
                    ),
                    colleague_price=self._parse_optional_price(
                        cells_by_col.get(optional_cols["colleague_price"]).value
                        if "colleague_price" in optional_cols
                        and optional_cols["colleague_price"] in cells_by_col
                        else None,
                        column=HEADER_COLLEAGUE_PRICE,
                    ),
                    raw_values=raw_values,
                )
            except (ValueError, TypeError, InvalidProductRowError) as exc:
                column = (
                    getattr(exc, "column", None)
                    if isinstance(exc, InvalidProductRowError)
                    else None
                )
                reason = (
                    exc.reason
                    if isinstance(exc, InvalidProductRowError) and exc.reason
                    else str(exc)
                )
                invalid_rows.append(
                    InvalidProductRow(
                        row_number=row_number,
                        reason=reason,
                        column=column,
                        code=code,
                        name=name,
                        raw_values=raw_values,
                    )
                )
                continue

            products.append(product)

        return products, invalid_rows

    @staticmethod
    def _extract_row_cells(
        row_cells: Iterable[Any],
    ) -> tuple[dict[int, _CellSnapshot], dict[str, Any]]:
        """Build column-index snapshots and letter→value raw maps.

        ``read_only`` mode may yield ``EmptyCell`` objects without ``column`` /
        ``row`` attributes; fall back to 1-based position within the tuple.
        """
        cells_by_col: dict[int, _CellSnapshot] = {}
        raw_values: dict[str, Any] = {}
        for index, cell in enumerate(row_cells, start=1):
            column = getattr(cell, "column", None)
            if column is None:
                column = index
            value = cell.value
            data_type = getattr(cell, "data_type", "n")
            number_format = getattr(cell, "number_format", "General") or "General"
            cells_by_col[column] = _CellSnapshot(
                value=value,
                data_type=str(data_type),
                number_format=str(number_format),
            )
            raw_values[get_column_letter(column)] = value
        return cells_by_col, raw_values

    @staticmethod
    def _row_is_entirely_empty(values: Iterable[Any]) -> bool:
        for value in values:
            if value is None:
                continue
            if isinstance(value, str) and not value.strip():
                continue
            return False
        return True

    def _parse_code_or_barcode(
        self,
        cell: _CellSnapshot | None,
        *,
        column: str,
        optional: bool = False,
    ) -> str | None:
        """Parse code/barcode identifiers with leading-zero safety."""
        if cell is None or cell.value is None:
            return None
        if isinstance(cell.value, str) and not cell.value.strip():
            return None

        self._reject_formula_or_error_identifier(cell, column=column)

        if isinstance(cell.value, bool):
            raise InvalidProductRowError(
                "Invalid identifier",
                column=column,
                reason="boolean is not a valid identifier",
            )

        if isinstance(cell.value, str):
            return cell.value.strip()

        if isinstance(cell.value, (int, float, Decimal)):
            padded = self._format_numeric_identifier(cell)
            if padded is not None:
                return padded
            raise InvalidProductRowError(
                "Unsupported numeric identifier format",
                column=column,
                reason=(
                    "numeric cell has no zero-pad number format; "
                    "store as text to preserve leading zeros"
                ),
            )

        if optional and cell.value is None:
            return None
        raise InvalidProductRowError(
            "Invalid identifier",
            column=column,
            reason=f"unsupported identifier type: {type(cell.value).__name__}",
        )

    def _parse_name(self, cell: _CellSnapshot | None) -> str | None:
        """Parse product name; reject formulas/errors, allow plain text/numbers."""
        if cell is None or cell.value is None:
            return None
        if isinstance(cell.value, str) and not cell.value.strip():
            return None

        self._reject_formula_or_error_identifier(cell, column=HEADER_NAME)

        if isinstance(cell.value, bool):
            raise InvalidProductRowError(
                "Invalid product name",
                column=HEADER_NAME,
                reason="boolean is not a valid product name",
            )
        if isinstance(cell.value, str):
            return cell.value.strip()
        if isinstance(cell.value, int):
            return str(cell.value)
        if isinstance(cell.value, Decimal):
            if cell.value != cell.value.to_integral_value():
                raise InvalidProductRowError(
                    "Invalid product name",
                    column=HEADER_NAME,
                    reason=f"non-integer numeric name: {cell.value!r}",
                )
            return format(cell.value, "f").split(".")[0]
        if isinstance(cell.value, float):
            as_decimal = Decimal(str(cell.value))
            if as_decimal != as_decimal.to_integral_value():
                raise InvalidProductRowError(
                    "Invalid product name",
                    column=HEADER_NAME,
                    reason=f"non-integer numeric name: {cell.value!r}",
                )
            return format(int(as_decimal), "d")
        raise InvalidProductRowError(
            "Invalid product name",
            column=HEADER_NAME,
            reason=f"unsupported name type: {type(cell.value).__name__}",
        )

    @staticmethod
    def _reject_formula_or_error_identifier(
        cell: _CellSnapshot,
        *,
        column: str,
    ) -> None:
        if cell.data_type == "f" or (
            isinstance(cell.value, str) and cell.value.startswith("=")
        ):
            raise InvalidProductRowError(
                "Formula cell is not a valid identifier",
                column=column,
                reason="formula/error cells cannot be used as codes/barcodes/names",
            )
        if cell.data_type == "e":
            raise InvalidProductRowError(
                "Error cell is not a valid identifier",
                column=column,
                reason=f"Excel error value: {cell.value!r}",
            )
        if isinstance(cell.value, str) and cell.value.strip().upper() in {
            err.upper() for err in _EXCEL_ERROR_VALUES
        }:
            raise InvalidProductRowError(
                "Error cell is not a valid identifier",
                column=column,
                reason=f"Excel error value: {cell.value!r}",
            )

    @staticmethod
    def _format_numeric_identifier(cell: _CellSnapshot) -> str | None:
        """Return zero-padded text when number format safely encodes width.

        Returns ``None`` when the format cannot preserve leading zeros,
        including multi-section formats (unsupported).
        """
        fmt = cell.number_format.strip()
        # Reject multi-section formats — later sections are not interpreted.
        if ";" in fmt:
            return None
        match = _ZERO_PAD_FORMAT_RE.fullmatch(fmt)
        if match is None:
            return None

        width = len(match.group(0))
        value = cell.value
        if isinstance(value, bool):
            return None
        if isinstance(value, int):
            number = value
        elif isinstance(value, Decimal):
            if value != value.to_integral_value():
                return None
            number = int(value)
        elif isinstance(value, float):
            as_decimal = Decimal(str(value))
            if as_decimal != as_decimal.to_integral_value():
                return None
            number = int(as_decimal)
        else:
            return None

        if number < 0:
            return None
        text = format(number, "d")
        if len(text) > width:
            # Value wider than pad width — keep full digits, no truncation.
            return text
        return text.zfill(width)

    @staticmethod
    def _parse_optional_price(
        value: Any,
        *,
        column: str,
    ) -> Decimal | None:
        """Parse a price cell to Decimal, or None if empty.

        Raises ``InvalidProductRowError`` (caught by ``_parse_rows`` and stored
        as an ``InvalidProductRow`` for review) when the value is unusable.
        """
        if value is None:
            return None
        if isinstance(value, str) and not value.strip():
            return None
        if isinstance(value, bool):
            raise InvalidProductRowError(
                "Invalid price value",
                column=column,
                reason="boolean is not a valid price",
            )
        if isinstance(value, str) and value.startswith("="):
            raise InvalidProductRowError(
                "Invalid price value",
                column=column,
                reason="formula cells are not supported for prices",
            )
        if isinstance(value, Decimal):
            if not value.is_finite():
                raise InvalidProductRowError(
                    "Invalid price value",
                    column=column,
                    reason=f"non-finite decimal: {value!r}",
                )
            return value
        if isinstance(value, int):
            return Decimal(value)
        if isinstance(value, float):
            # Convert via str to avoid introducing extra binary noise beyond
            # what the float already lost.
            try:
                amount = Decimal(str(value))
            except InvalidOperation as exc:
                raise InvalidProductRowError(
                    "Invalid price value",
                    column=column,
                    reason=f"cannot parse float price: {value!r}",
                ) from exc
            if not amount.is_finite():
                raise InvalidProductRowError(
                    "Invalid price value",
                    column=column,
                    reason=f"non-finite float price: {value!r}",
                )
            return amount
        if isinstance(value, str):
            try:
                amount = Decimal(value.strip())
            except InvalidOperation as exc:
                raise InvalidProductRowError(
                    "Invalid price value",
                    column=column,
                    reason=f"cannot parse price text: {value!r}",
                ) from exc
            if not amount.is_finite():
                raise InvalidProductRowError(
                    "Invalid price value",
                    column=column,
                    reason=f"non-finite price text: {value!r}",
                )
            return amount
        raise InvalidProductRowError(
            "Invalid price value",
            column=column,
            reason=f"unsupported price type: {type(value).__name__}",
        )

    @staticmethod
    def _find_duplicate_codes(
        products: list[PioneersProduct],
    ) -> dict[str, tuple[int, ...]]:
        rows_by_code: dict[str, list[int]] = defaultdict(list)
        for product in products:
            rows_by_code[product.code].append(product.row_number)
        return {
            code: tuple(rows)
            for code, rows in rows_by_code.items()
            if len(rows) > 1
        }


__all__ = [
    "HEADER_BARCODE",
    "HEADER_CODE",
    "HEADER_COLLEAGUE_PRICE",
    "HEADER_CONSUMER_PRICE",
    "HEADER_LAST_PURCHASE_PRICE",
    "HEADER_NAME",
    "HEADER_ROW",
    "HEADER_SALE_PRICE",
    "InvalidProductRow",
    "InvalidProductRowError",
    "MAPPED_HEADERS",
    "AmbiguousMappedHeaderError",
    "MissingRequiredColumnError",
    "OPTIONAL_FIELD_HEADERS",
    "PRODUCT_SHEET_NAME",
    "PioneersExcelError",
    "PioneersExcelParser",
    "PioneersParseResult",
    "REQUIRED_HEADERS",
    "SheetNotFoundError",
]
