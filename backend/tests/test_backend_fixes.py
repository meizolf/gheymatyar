"""Regression tests for the pre-frontend backend fixes."""

from __future__ import annotations

import hashlib
import sqlite3
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, inspect, text

from gheymatyar.db import Database
from gheymatyar.db.migrate import LEGACY_MAPPINGS_TABLE, MigrationError, migrate_schema
from gheymatyar.domain.enums import MatchStatus, ReviewDecision
from gheymatyar.domain.product import PioneersProduct
from gheymatyar.domain.supplier_item import SupplierPriceItem
from gheymatyar.domain.update_candidate import UpdateCandidate
from gheymatyar.parsers.pioneers_excel_parser import (
    HEADER_CODE,
    HEADER_NAME,
    HEADER_SALE_PRICE,
    PioneersExcelParser,
)
from gheymatyar.repositories import MappingRepository, SettingsRepository, SupplierRepository
from gheymatyar.services.matching import match_supplier_items
from gheymatyar.services.price_validation import (
    IDENTITY_CONFLICT_WARNING,
    is_exportable,
    parse_suspicious_percent_threshold,
    resolve_suspicious_percent_threshold,
)
from gheymatyar.services.processing import ProcessingService
from gheymatyar.services.session_store import SessionStore
from gheymatyar.suppliers.azin_pipe import AzinPipeProcessor


def _write_xlsx(path: Path, rows: list[list[object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "لیست کالا"
    ws.append([HEADER_CODE, HEADER_NAME, HEADER_SALE_PRICE])
    for row in rows:
        ws.append(row)
    wb.save(path)


# --- 1. Stale export invalidation ---


def test_review_exclude_invalidates_previous_export(
    tmp_path: Path, pioneers_azin_excel: Path, azin_pdf: Path, db: Database
) -> None:
    store = SessionStore()
    svc = ProcessingService(store)
    session = svc.create_session()
    with pioneers_azin_excel.open("rb") as handle:
        svc.upload_pioneers_excel(session.session_id, "p.xlsx", handle)
    with azin_pdf.open("rb") as handle:
        svc.upload_supplier_pdf(
            session.session_id,
            supplier_key="azin_pipe",
            filename="a.pdf",
            stream=handle,
        )
    session = svc.get_session(session.session_id)
    product = session.products[0]
    session.supplier_items = [
        SupplierPriceItem(
            supplier_id="azin_pipe",
            supplier_name=product.name,
            supplier_code="819121020",
            price=product.sale_price or Decimal("1000"),
            unit="عدد",
            raw_values={"currency": "rial", "price_basis": "piece", "page": 1},
        )
    ]
    with db.session() as sess:
        SettingsRepository(sess).set("pioneers_sale_currency", "rial")
        SettingsRepository(sess).set("pioneers_sale_basis", "piece")
        svc.run_matching(session.session_id, sess)

    session = svc.get_session(session.session_id)
    matched = next(c for c in session.candidates if c.status == MatchStatus.MATCHED)
    svc.set_review_decision(
        session.session_id,
        candidate_id=matched.candidate_id,
        decision=ReviewDecision.APPROVE,
    )
    # Exclude all other matched so export is deterministic
    for c in svc.get_session(session.session_id).candidates:
        if c.candidate_id == matched.candidate_id:
            continue
        if c.status == MatchStatus.MATCHED:
            svc.set_review_decision(
                session.session_id,
                candidate_id=c.candidate_id,
                decision=ReviewDecision.EXCLUDE,
            )

    result = svc.export_excel(session.session_id)
    assert Path(result.output_path).is_file()
    assert svc.get_session(session.session_id).output_path is not None

    svc.set_review_decision(
        session.session_id,
        candidate_id=matched.candidate_id,
        decision=ReviewDecision.EXCLUDE,
    )
    session = svc.get_session(session.session_id)
    assert session.output_path is None
    assert not Path(result.output_path).exists()


# --- 2. Atomic replacement uploads ---


def test_failed_excel_reupload_keeps_previous_state(tmp_path: Path) -> None:
    store = SessionStore()
    svc = ProcessingService(store)
    session = svc.create_session()

    good = tmp_path / "good.xlsx"
    _write_xlsx(good, [["100", "Product A", "1000"]])
    with good.open("rb") as handle:
        svc.upload_pioneers_excel(session.session_id, "good.xlsx", handle)

    session = svc.get_session(session.session_id)
    assert len(session.products) == 1
    assert session.products[0].name == "Product A"
    good_bytes = session.pioneers_path.read_bytes()

    # Corrupt upload: not a valid xlsx
    with pytest.raises(Exception):
        svc.upload_pioneers_excel(
            session.session_id, "bad.xlsx", BytesIO(b"not-an-excel-file")
        )

    session = svc.get_session(session.session_id)
    assert len(session.products) == 1
    assert session.products[0].name == "Product A"
    assert session.pioneers_path.read_bytes() == good_bytes


# --- 3. Legacy SQLite migration ---


def test_legacy_code_mapping_schema_migrates_safely(tmp_path: Path) -> None:
    path = tmp_path / "legacy.db"
    engine = create_engine(f"sqlite:///{path}")
    with engine.begin() as conn:
        conn.execute(
            text(
                """
                CREATE TABLE suppliers (
                    id INTEGER PRIMARY KEY,
                    name VARCHAR(255) NOT NULL UNIQUE,
                    parser_type VARCHAR(128) NOT NULL,
                    created_at DATETIME NOT NULL
                )
                """
            )
        )
        conn.execute(
            text(
                """
                CREATE TABLE product_mappings (
                    id INTEGER PRIMARY KEY,
                    supplier_id INTEGER NOT NULL,
                    supplier_code VARCHAR(255) NOT NULL,
                    pioneers_code VARCHAR(255) NOT NULL,
                    mapping_type VARCHAR(32) NOT NULL,
                    created_at DATETIME NOT NULL,
                    updated_at DATETIME NOT NULL,
                    UNIQUE (supplier_id, supplier_code)
                )
                """
            )
        )
        conn.execute(
            text(
                "INSERT INTO suppliers (id, name, parser_type, created_at) "
                "VALUES (1, 'آذین لوله', 'azin_pipe', '2026-01-01')"
            )
        )
        conn.execute(
            text(
                "INSERT INTO product_mappings "
                "(id, supplier_id, supplier_code, pioneers_code, mapping_type, "
                "created_at, updated_at) "
                "VALUES (1, 1, 'ABC', '100', 'MANUAL', '2026-01-01', '2026-01-01')"
            )
        )

    notes = migrate_schema(engine)
    assert notes
    Database(f"sqlite:///{path}").create_tables()

    inspector = inspect(engine)
    tables = set(inspector.get_table_names())
    assert LEGACY_MAPPINGS_TABLE in tables
    assert "product_mappings" in tables
    new_cols = {c["name"] for c in inspector.get_columns("product_mappings")}
    assert "supplier_name_normalized" in new_cols
    assert "identity_key" in new_cols
    assert "legacy_unscoped" in new_cols
    # Live table stores the variant code. The archived code-key schema is separate.
    assert "supplier_code" in new_cols

    with engine.connect() as conn:
        legacy_count = conn.execute(
            text(f"SELECT COUNT(*) FROM {LEGACY_MAPPINGS_TABLE}")
        ).scalar()
        new_count = conn.execute(text("SELECT COUNT(*) FROM product_mappings")).scalar()
    assert legacy_count == 1
    assert new_count == 0  # names unknown → require reconfirmation

    # Matching against migrated DB must not crash
    db = Database(f"sqlite:///{path}")
    store = SessionStore()
    svc = ProcessingService(store)
    session = svc.create_session()
    excel = tmp_path / "p.xlsx"
    _write_xlsx(excel, [["100", "اذین بوشن 20", "204600"]])
    with excel.open("rb") as handle:
        svc.upload_pioneers_excel(session.session_id, "p.xlsx", handle)

    # Minimal supplier item without needing full PDF
    session = svc.get_session(session.session_id)
    session.supplier_key = "azin_pipe"
    session.supplier_display_name = "آذین لوله"
    session.supplier_confirmed = True
    session.supplier_items = [
        SupplierPriceItem(
            supplier_id="azin_pipe",
            supplier_code="855321101",
            supplier_name="اذین بوشن 20",
            price=Decimal("204600"),
        )
    ]
    with db.session() as sess:
        svc.run_matching(session.session_id, sess)
    assert svc.get_session(session.session_id).matched_count == 1
    engine.dispose()
    db.engine.dispose()


def test_name_only_mappings_migrate_as_legacy_unscoped(tmp_path: Path) -> None:
    path = tmp_path / "names.db"
    engine = create_engine(f"sqlite:///{path}")
    with engine.begin() as conn:
        conn.execute(
            text(
                """
                CREATE TABLE suppliers (
                    id INTEGER PRIMARY KEY,
                    name VARCHAR(255) NOT NULL UNIQUE,
                    parser_type VARCHAR(128) NOT NULL,
                    created_at DATETIME NOT NULL
                )
                """
            )
        )
        conn.execute(
            text(
                """
                CREATE TABLE product_mappings (
                    id INTEGER PRIMARY KEY,
                    supplier_id INTEGER NOT NULL,
                    supplier_name_raw TEXT NOT NULL,
                    supplier_name_normalized VARCHAR(512) NOT NULL,
                    pioneers_name_raw TEXT NOT NULL,
                    pioneers_name_normalized VARCHAR(512) NOT NULL,
                    pioneers_code VARCHAR(255),
                    mapping_type VARCHAR(32) NOT NULL,
                    created_at DATETIME NOT NULL,
                    updated_at DATETIME NOT NULL,
                    UNIQUE (supplier_id, supplier_name_normalized)
                )
                """
            )
        )
        conn.execute(
            text(
                "INSERT INTO suppliers (id, name, parser_type, created_at) "
                "VALUES (1, 'آذین لوله', 'azin_pipe', '2026-01-01')"
            )
        )
        conn.execute(
            text(
                "INSERT INTO product_mappings ("
                "id, supplier_id, supplier_name_raw, supplier_name_normalized, "
                "pioneers_name_raw, pioneers_name_normalized, pioneers_code, "
                "mapping_type, created_at, updated_at) "
                "VALUES (7, 1, 'کالای مشترک 20', 'کالای مشترک 20', 'هدف', 'هدف', "
                "'100', 'MANUAL', '2026-01-01', '2026-01-01')"
            )
        )

    notes = migrate_schema(engine)
    assert any("legacy_unscoped" in note for note in notes)
    Database(f"sqlite:///{path}").create_tables()

    with engine.connect() as conn:
        row = conn.execute(
            text(
                "SELECT id, pioneers_code, pioneers_name_raw, legacy_unscoped, "
                "supplier_code, identity_key FROM product_mappings"
            )
        ).one()
    assert row.id == 7
    assert row.pioneers_code == "100"
    assert row.pioneers_name_raw == "هدف"
    assert row.legacy_unscoped in (1, True)
    assert row.supplier_code is None
    assert row.identity_key == "کالای مشترک 20\x1f\x1f"
    assert migrate_schema(engine) == []
    engine.dispose()


def test_failed_mapping_copy_rolls_back_and_reopens(
    tmp_path: Path,
) -> None:
    from sqlalchemy import event
    from sqlalchemy.exc import OperationalError

    path = tmp_path / "legacy-upgrade.db"
    with sqlite3.connect(path) as db:
        db.executescript(
            """
            CREATE TABLE suppliers (
                id INTEGER PRIMARY KEY,
                name VARCHAR(255) NOT NULL UNIQUE,
                parser_type VARCHAR(128) NOT NULL,
                created_at DATETIME NOT NULL
            );
            CREATE TABLE product_mappings (
                id INTEGER PRIMARY KEY,
                supplier_id INTEGER NOT NULL,
                supplier_name_raw TEXT NOT NULL,
                supplier_name_normalized VARCHAR(512) NOT NULL,
                pioneers_name_raw TEXT NOT NULL,
                pioneers_name_normalized VARCHAR(512) NOT NULL,
                pioneers_code VARCHAR(255),
                mapping_type VARCHAR(32) NOT NULL,
                created_at DATETIME NOT NULL,
                updated_at DATETIME NOT NULL,
                UNIQUE (supplier_id, supplier_name_normalized)
            );
            INSERT INTO suppliers VALUES (
                1, 'dina', 'dina_polymer', '2026-01-01'
            );
            INSERT INTO product_mappings VALUES (
                7, 1, 'source', 'source', 'target', 'target', '100',
                'MANUAL', '2026-01-01', '2026-01-01'
            );
            """
        )
    database = Database(f"sqlite:///{path}")

    def fail_copy(conn, cursor, statement, parameters, context, executemany):
        if statement.lstrip().startswith("INSERT INTO product_mappings"):
            raise OperationalError(
                statement, parameters, Exception("injected copy failure")
            )

    event.listen(database.engine, "before_cursor_execute", fail_copy)
    with pytest.raises(OperationalError):
        database.create_tables()
    event.remove(database.engine, "before_cursor_execute", fail_copy)
    database.engine.dispose()

    reopened = Database(f"sqlite:///{path}")
    reopened.create_tables()
    with sqlite3.connect(path) as db:
        tables = {
            row[0]
            for row in db.execute("SELECT name FROM sqlite_master WHERE type='table'")
        }
        assert "product_mappings__pre_identity" not in tables
        active = db.execute(
            "SELECT pioneers_code, legacy_unscoped FROM product_mappings"
        ).fetchone()
    assert active[0] == "100"
    assert active[1] in (1, True)
    reopened.engine.dispose()


def test_stranded_mapping_upgrade_is_recovered(tmp_path: Path) -> None:
    path = tmp_path / "stranded.db"
    with sqlite3.connect(path) as db:
        db.executescript(
            """
            CREATE TABLE suppliers (
                id INTEGER PRIMARY KEY,
                name VARCHAR(255) NOT NULL UNIQUE,
                parser_type VARCHAR(128) NOT NULL,
                created_at DATETIME NOT NULL
            );
            CREATE TABLE product_mappings__pre_identity (
                id INTEGER PRIMARY KEY,
                supplier_id INTEGER NOT NULL,
                supplier_name_raw TEXT NOT NULL,
                supplier_name_normalized VARCHAR(512) NOT NULL,
                pioneers_name_raw TEXT NOT NULL,
                pioneers_name_normalized VARCHAR(512) NOT NULL,
                pioneers_code VARCHAR(255),
                mapping_type VARCHAR(32) NOT NULL,
                created_at DATETIME NOT NULL,
                updated_at DATETIME NOT NULL
            );
            CREATE TABLE product_mappings (
                id INTEGER PRIMARY KEY,
                supplier_id INTEGER NOT NULL,
                identity_key TEXT NOT NULL,
                supplier_name_raw TEXT NOT NULL,
                supplier_name_normalized VARCHAR(512) NOT NULL,
                supplier_code VARCHAR(255),
                unit_normalized VARCHAR(255),
                legacy_unscoped BOOLEAN NOT NULL DEFAULT 0,
                pioneers_name_raw TEXT NOT NULL,
                pioneers_name_normalized VARCHAR(512) NOT NULL,
                pioneers_code VARCHAR(255),
                mapping_type VARCHAR(32) NOT NULL,
                created_at DATETIME NOT NULL,
                updated_at DATETIME NOT NULL
            );
            INSERT INTO suppliers VALUES (1, 'dina', 'dina_polymer', '2026-01-01');
            INSERT INTO product_mappings__pre_identity VALUES (
                7, 1, 'source', 'source', 'target', 'target', '100',
                'MANUAL', '2026-01-01', '2026-01-01'
            );
            """
        )
    database = Database(f"sqlite:///{path}")
    database.create_tables()
    with sqlite3.connect(path) as db:
        tables = {
            row[0]
            for row in db.execute("SELECT name FROM sqlite_master WHERE type='table'")
        }
        assert "product_mappings__pre_identity" not in tables
        row = db.execute(
            "SELECT id, pioneers_code, legacy_unscoped FROM product_mappings"
        ).fetchone()
    assert row[0] == 7
    assert row[1] == "100"
    assert row[2] in (1, True)
    database.engine.dispose()


def test_unrecognized_mapping_schema_is_not_dropped(tmp_path: Path) -> None:
    path = tmp_path / "weird.db"
    engine = create_engine(f"sqlite:///{path}")
    with engine.begin() as conn:
        conn.execute(
            text(
                """
                CREATE TABLE product_mappings (
                    id INTEGER PRIMARY KEY,
                    mystery_column TEXT NOT NULL,
                    another_col INTEGER
                )
                """
            )
        )
        conn.execute(
            text(
                "INSERT INTO product_mappings (id, mystery_column, another_col) "
                "VALUES (1, 'keep-me', 42)"
            )
        )

    with pytest.raises(MigrationError, match="unrecognized schema"):
        migrate_schema(engine)

    with engine.connect() as conn:
        row = conn.execute(
            text("SELECT mystery_column, another_col FROM product_mappings WHERE id = 1")
        ).one()
    assert row == ("keep-me", 42)
    assert "product_mappings" in inspect(engine).get_table_names()
    engine.dispose()


def test_existing_legacy_archive_survives_second_migration(tmp_path: Path) -> None:
    path = tmp_path / "archive.db"
    engine = create_engine(f"sqlite:///{path}")
    with engine.begin() as conn:
        conn.execute(
            text(
                f"""
                CREATE TABLE {LEGACY_MAPPINGS_TABLE} (
                    id INTEGER PRIMARY KEY,
                    supplier_id INTEGER NOT NULL,
                    supplier_code VARCHAR(255) NOT NULL,
                    pioneers_code VARCHAR(255) NOT NULL,
                    mapping_type VARCHAR(32) NOT NULL,
                    created_at DATETIME NOT NULL,
                    updated_at DATETIME NOT NULL
                )
                """
            )
        )
        conn.execute(
            text(
                f"INSERT INTO {LEGACY_MAPPINGS_TABLE} "
                "(id, supplier_id, supplier_code, pioneers_code, mapping_type, "
                "created_at, updated_at) "
                "VALUES (1, 1, 'OLD', '1', 'MANUAL', '2026-01-01', '2026-01-01')"
            )
        )
        conn.execute(
            text(
                """
                CREATE TABLE product_mappings (
                    id INTEGER PRIMARY KEY,
                    supplier_id INTEGER NOT NULL,
                    supplier_code VARCHAR(255) NOT NULL,
                    pioneers_code VARCHAR(255) NOT NULL,
                    mapping_type VARCHAR(32) NOT NULL,
                    created_at DATETIME NOT NULL,
                    updated_at DATETIME NOT NULL
                )
                """
            )
        )
        conn.execute(
            text(
                "INSERT INTO product_mappings "
                "(id, supplier_id, supplier_code, pioneers_code, mapping_type, "
                "created_at, updated_at) "
                "VALUES (2, 1, 'NEW', '2', 'MANUAL', '2026-01-02', '2026-01-02')"
            )
        )

    notes = migrate_schema(engine)
    assert any("legacy_code_2" in n or "preserved existing archive" in n for n in notes)

    tables = set(inspect(engine).get_table_names())
    assert LEGACY_MAPPINGS_TABLE in tables
    assert f"{LEGACY_MAPPINGS_TABLE}_2" in tables
    assert "product_mappings" not in tables  # renamed; create_all recreates later

    with engine.connect() as conn:
        old = conn.execute(
            text(f"SELECT supplier_code FROM {LEGACY_MAPPINGS_TABLE} WHERE id = 1")
        ).scalar()
        archived = conn.execute(
            text(
                f"SELECT supplier_code FROM {LEGACY_MAPPINGS_TABLE}_2 WHERE id = 2"
            )
        ).scalar()
    assert old == "OLD"
    assert archived == "NEW"
    engine.dispose()


# --- 4. Duplicate code export blocked ---


def _assert_duplicate_blocks_export(result, *, code: str, product_row: int) -> None:
    from gheymatyar.services.price_validation import apply_price_warnings

    assert code in result.duplicate_codes
    product = next(p for p in result.products if p.row_number == product_row)
    candidate = UpdateCandidate(
        candidate_id="c1",
        row_number=product.row_number,
        pioneers_code=product.code,
        pioneers_name=product.name,
        supplier_name=product.name,
        old_price=product.sale_price or Decimal("1000"),
        new_price=Decimal("1500"),
        absolute_change=Decimal("500"),
        percentage_change=Decimal("50"),
        match_type=None,
        status=MatchStatus.MATCHED,
        review_decision=ReviewDecision.APPROVE,
        warnings=[],
    )
    conflicted = frozenset(
        row for rows in result.duplicate_codes.values() for row in rows
    )
    enriched = apply_price_warnings([candidate], conflicted_row_numbers=conflicted)[0]
    assert IDENTITY_CONFLICT_WARNING in enriched.warnings
    assert not is_exportable(enriched)


def test_duplicate_code_with_invalid_nameless_row_blocks_export(tmp_path: Path) -> None:
    """Valid product sharing a code with a nameless invalid row must be blocked."""
    path = tmp_path / "dup_invalid.xlsx"
    _write_xlsx(
        path,
        [
            ["DUP", "Valid Named Product", "1000"],
            ["DUP", None, "2000"],  # invalid (no name) but still carries the code
        ],
    )
    result = PioneersExcelParser().parse(path)
    assert len(result.products) == 1
    assert result.invalid_rows[0].code == "DUP"
    assert set(result.duplicate_codes["DUP"]) == {2, 3}
    _assert_duplicate_blocks_export(result, code="DUP", product_row=2)


def test_duplicate_code_with_formula_name_blocks_export(tmp_path: Path) -> None:
    """Code recorded before name formula rejection must still block export."""
    path = tmp_path / "dup_formula_name.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "لیست کالا"
    ws.append([HEADER_CODE, HEADER_NAME, HEADER_SALE_PRICE])
    ws.append(["FORMULA_DUP", "Valid Named Product", "1000"])
    ws.append(["FORMULA_DUP", "=B2", "2000"])
    wb.save(path)

    result = PioneersExcelParser().parse(path)
    assert len(result.products) == 1
    assert result.invalid_rows[0].code == "FORMULA_DUP"
    assert "formula" in result.invalid_rows[0].reason.lower()
    assert set(result.duplicate_codes["FORMULA_DUP"]) == {2, 3}
    _assert_duplicate_blocks_export(result, code="FORMULA_DUP", product_row=2)


def test_duplicate_code_with_excel_error_name_blocks_export(tmp_path: Path) -> None:
    """Code recorded before Excel-error name rejection must still block export."""
    path = tmp_path / "dup_error_name.xlsx"
    _write_xlsx(
        path,
        [
            ["ERR_DUP", "Valid Named Product", "1000"],
            ["ERR_DUP", "#N/A", "2000"],
        ],
    )
    result = PioneersExcelParser().parse(path)
    assert len(result.products) == 1
    assert result.invalid_rows[0].code == "ERR_DUP"
    assert "error" in result.invalid_rows[0].reason.lower()
    assert set(result.duplicate_codes["ERR_DUP"]) == {2, 3}
    _assert_duplicate_blocks_export(result, code="ERR_DUP", product_row=2)


def test_duplicate_pioneers_code_cannot_export_even_if_approved() -> None:
    store = SessionStore()
    svc = ProcessingService(store)
    session = svc.create_session()
    session.products = [
        PioneersProduct(row_number=2, code="DUP", name="One", sale_price=Decimal("10")),
        PioneersProduct(row_number=3, code="DUP", name="Two", sale_price=Decimal("20")),
    ]
    session.duplicate_codes = {"DUP": (2, 3)}
    session.supplier_key = "azin_pipe"
    session.supplier_items = [
        SupplierPriceItem(
            supplier_id="azin_pipe",
            supplier_name="One",
            price=Decimal("15"),
        )
    ]
    result = match_supplier_items(
        session.supplier_items, session.products, processor=AzinPipeProcessor()
    )
    from gheymatyar.services.price_validation import apply_price_warnings

    enriched = apply_price_warnings(
        result.candidates, conflicted_row_numbers=session.conflicted_row_numbers
    )
    matched = next(c for c in enriched if c.status == MatchStatus.MATCHED)
    assert IDENTITY_CONFLICT_WARNING in matched.warnings
    assert not is_exportable(matched)
    approved = matched.model_copy(update={"review_decision": ReviewDecision.APPROVE})
    assert not is_exportable(approved)


# --- 5. Threshold resolution ---


def test_threshold_request_overrides_saved_overrides_default(db: Database) -> None:
    with db.session() as sess:
        SettingsRepository(sess).set("suspicious_percent_threshold", "15")
        saved = SettingsRepository(sess).get("suspicious_percent_threshold")
        assert saved is not None
        resolved = resolve_suspicious_percent_threshold(
            request_value="40",
            saved_value=saved.value,
        )
        assert resolved == Decimal("40")
        resolved_saved = resolve_suspicious_percent_threshold(
            request_value=None,
            saved_value=saved.value,
        )
        assert resolved_saved == Decimal("15")
        resolved_default = resolve_suspicious_percent_threshold(
            request_value=None,
            saved_value=None,
        )
        assert resolved_default == Decimal("30")


@pytest.mark.parametrize("bad", ["0", "-5", "NaN", "Infinity", "-Infinity", "abc", ""])
def test_threshold_rejects_invalid_values(bad: str) -> None:
    with pytest.raises(ValueError):
        parse_suspicious_percent_threshold(bad)


def test_match_uses_saved_threshold(client, pioneers_azin_excel: Path, azin_pdf: Path) -> None:
    client.put(
        "/api/settings",
        json={"settings": {"suspicious_percent_threshold": "5"}},
    )
    sid = client.post("/api/sessions").json()["session_id"]
    with pioneers_azin_excel.open("rb") as handle:
        client.post(
            f"/api/sessions/{sid}/pioneers-excel",
            files={"file": ("p.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    with azin_pdf.open("rb") as handle:
        client.post(
            f"/api/sessions/{sid}/supplier-pdf",
            data={"supplier_key": "azin_pipe"},
            files={"file": ("a.pdf", handle, "application/pdf")},
        )
    # No request override — saved 5% should apply
    match = client.post(f"/api/sessions/{sid}/match", json={})
    assert match.status_code == 200, match.text
    matched = [c for c in match.json()["candidates"] if c["status"] == "MATCHED"]
    assert matched
    # Large price moves should be warned under 5%
    warned = [c for c in matched if any("exceeds 5%" in w for w in c["warnings"])]
    assert warned


# --- 6. Optional codes ---


def test_name_only_product_matches_and_exports(tmp_path: Path, db: Database) -> None:
    excel = tmp_path / "name_only.xlsx"
    _write_xlsx(excel, [[None, "اذین بوشن 20", "100000"]])
    result = PioneersExcelParser().parse(excel)
    assert len(result.products) == 1
    assert result.products[0].code is None

    store = SessionStore()
    svc = ProcessingService(store)
    session = svc.create_session()
    with excel.open("rb") as handle:
        svc.upload_pioneers_excel(session.session_id, "n.xlsx", handle)
    session = svc.get_session(session.session_id)
    session.supplier_key = "azin_pipe"
    session.supplier_confirmed = True
    session.supplier_items = [
        SupplierPriceItem(
            supplier_id="azin_pipe",
            supplier_name="اذین بوشن 20",
            price=Decimal("204600"),
            unit="عدد",
            raw_values={"currency": "rial", "price_basis": "piece", "page": 1},
        )
    ]
    with db.session() as sess:
        SettingsRepository(sess).set("pioneers_sale_currency", "rial")
        SettingsRepository(sess).set("pioneers_sale_basis", "piece")
        svc.run_matching(session.session_id, sess)
    session = svc.get_session(session.session_id)
    matched = next(c for c in session.candidates if c.status == MatchStatus.MATCHED)
    assert matched.pioneers_code is None
    svc.set_review_decision(
        session.session_id,
        candidate_id=matched.candidate_id,
        decision=ReviewDecision.APPROVE,
    )
    exported = svc.export_excel(session.session_id)
    assert exported.updated_count == 1
    wb = load_workbook(exported.output_path, read_only=True, data_only=True)
    ws = wb.active
    row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert str(row[2]) == "204600"
    wb.close()


# --- 7. Invalid row diagnostics in API ---


def test_api_exposes_invalid_row_diagnostics(client, tmp_path: Path) -> None:
    excel = tmp_path / "invalid.xlsx"
    _write_xlsx(
        excel,
        [
            ["90099", None, "100"],  # missing name → invalid
            ["90100", "Valid", "150"],
        ],
    )
    sid = client.post("/api/sessions").json()["session_id"]
    with excel.open("rb") as handle:
        resp = client.post(
            f"/api/sessions/{sid}/pioneers-excel",
            files={"file": ("i.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["invalid_row_count"] == 1
    assert len(body["invalid_rows"]) == 1
    row = body["invalid_rows"][0]
    assert row["row_number"] == 2
    assert row["column"] == HEADER_NAME
    assert "نـــام کــالا" in row["reason"]
    assert row["code"] == "90099"


# --- 8. Reject xlsm ---


def test_api_rejects_xlsm_upload(client) -> None:
    sid = client.post("/api/sessions").json()["session_id"]
    resp = client.post(
        f"/api/sessions/{sid}/pioneers-excel",
        files={"file": ("macro.xlsm", BytesIO(b"PK fake"), "application/vnd.ms-excel.sheet.macroEnabled.12")},
    )
    assert resp.status_code == 400
    assert "xlsm" in resp.json()["detail"].lower()
