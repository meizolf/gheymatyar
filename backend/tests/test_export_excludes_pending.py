"""Regression: excluded candidates must not appear in the exported workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from tests.confirmed_metre import prepare_two_exportable_metre_rows


def test_excluded_row_not_updated_in_download(
    client, pioneers_azin_excel: Path, dina_pdf: Path, tmp_path: Path
) -> None:
    session_id, keep, drop = prepare_two_exportable_metre_rows(
        client, pioneers_azin_excel, dina_pdf, tmp_path
    )
    client.post(
        f"/api/sessions/{session_id}/candidates/{drop['candidate_id']}/review",
        json={"decision": "EXCLUDE"},
    )

    export = client.post(f"/api/sessions/{session_id}/export")
    assert export.status_code == 200
    assert export.json()["updated_count"] == 1

    download = client.get(f"/api/sessions/{session_id}/download?format=xlsx")
    assert download.status_code == 200
    out = tmp_path / "out.xlsx"
    out.write_bytes(download.content)

    wb = load_workbook(out, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    sale_idx = headers.index("نرخ فروش اصلی")
    code_idx = headers.index("کد کالا")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    assert len(rows) == 1
    assert str(rows[0][sale_idx]).replace(",", "") == str(keep["new_price"]).replace(
        ",", ""
    )
    assert str(rows[0][code_idx]) == str(keep["pioneers_code"])
    assert str(drop["pioneers_code"]) != str(rows[0][code_idx])

    # Original fixture untouched
    assert pioneers_azin_excel.stat().st_size > 0
