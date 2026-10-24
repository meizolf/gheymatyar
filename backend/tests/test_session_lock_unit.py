"""Unit-level proof that session locks serialize review vs export."""

from __future__ import annotations

import threading
import time
from decimal import Decimal
from pathlib import Path

from gheymatyar.domain.enums import MatchStatus, ReviewDecision
from gheymatyar.domain.update_candidate import UpdateCandidate
from gheymatyar.services.processing import ProcessingService
from gheymatyar.services.session_store import SessionStore


def test_session_lock_serializes_exclude_before_export_snapshot(
    tmp_path: Path, pioneers_azin_excel: Path
) -> None:
    store = SessionStore()
    service = ProcessingService(store)
    session = service.create_session()

    # Minimal matched candidates — skip full PDF parse for lock timing.
    with pioneers_azin_excel.open("rb") as handle:
        service.upload_pioneers_excel(session.session_id, "p.xlsx", handle)

    session = service.get_session(session.session_id)
    session.supplier_key = "azin_pipe"
    session.supplier_confirmed = True
    # Seed two matched exportable candidates pointing at real product rows.
    products = session.products[:2]
    assert len(products) == 2
    candidates = []
    for product in products:
        candidates.append(
            UpdateCandidate(
                candidate_id=f"c-{product.row_number}",
                row_number=product.row_number,
                pioneers_code=product.code,
                pioneers_name=product.name,
                supplier_code=None,
                supplier_name=product.name,
                old_price=product.sale_price or Decimal("1000"),
                new_price=(product.sale_price or Decimal("1000")) + Decimal("100"),
                absolute_change=Decimal("100"),
                percentage_change=Decimal("10"),
                match_type=None,
                status=MatchStatus.MATCHED,
                review_decision=ReviewDecision.INCLUDE,
                warnings=[],
                suggestions=[],
                source_currency="rial",
                price_basis="piece",
                source_unit="عدد",
            )
        )
    session.candidates = candidates
    session.matched_count = 2
    keep_id, drop_id = candidates[0].candidate_id, candidates[1].candidate_id

    order: list[str] = []
    errors: list[BaseException] = []
    barrier = threading.Barrier(2)

    def exclude_slow() -> None:
        try:
            barrier.wait(timeout=5)
            lock = store.session_lock(session.session_id)
            with lock:
                order.append("exclude-enter")
                time.sleep(0.15)  # hold lock so export cannot snapshot early
                service.set_review_decision(
                    session.session_id,
                    candidate_id=drop_id,
                    decision=ReviewDecision.EXCLUDE,
                )
                order.append("exclude-done")
        except BaseException as exc:  # noqa: BLE001
            errors.append(exc)

    def export_after() -> None:
        try:
            barrier.wait(timeout=5)
            # Small delay so exclude acquires the lock first.
            time.sleep(0.02)
            order.append("export-enter")
            result = service.export_excel(session.session_id)
            order.append("export-done")
            assert result.updated_count == 1, result.updated_count
        except BaseException as exc:  # noqa: BLE001
            errors.append(exc)

    t1 = threading.Thread(target=exclude_slow)
    t2 = threading.Thread(target=export_after)
    t1.start()
    t2.start()
    t1.join(timeout=30)
    t2.join(timeout=30)
    assert not errors, errors
    assert order.index("exclude-enter") < order.index("export-done")
    assert order.index("exclude-done") < order.index("export-done")

    from openpyxl import load_workbook

    out = service.get_session(session.session_id).output_path
    assert out is not None and out.is_file()
    wb = load_workbook(out, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    sale_idx = headers.index("نرخ فروش اصلی")
    code_idx = headers.index("کد کالا")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    assert len(rows) == 1
    assert str(rows[0][code_idx]) == str(candidates[0].pioneers_code)
    assert str(rows[0][sale_idx]).replace(",", "") == str(candidates[0].new_price)
    assert str(candidates[1].pioneers_code) not in {str(row[code_idx]) for row in rows}
    assert keep_id  # silence unused if linted
