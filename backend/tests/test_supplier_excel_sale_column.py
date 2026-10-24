"""Printed Rial suppliers export only نرخ فروش اصلی; originals stay untouched."""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook

from tests.confirmed_metre import metre_items, workbook_with_supplier_names

from gheymatyar.parsers.pioneers_excel_parser import HEADER_NAME, HEADER_SALE_PRICE
from gheymatyar.suppliers.abrahan_pipe import AbrahanPipeProcessor


def _snapshot(path: Path) -> dict[tuple[int, int], object]:
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook.active
        return {
            (cell.row, cell.column): cell.value
            for row in sheet.iter_rows()
            for cell in row
        }
    finally:
        workbook.close()


def test_abrahan_export_changes_only_final_sale_price(
    client, pioneers_azin_excel: Path, abrahan_pdf: Path, tmp_path: Path
) -> None:
    item = next(
        row
        for row in AbrahanPipeProcessor().parse_pdf(abrahan_pdf).items
        if str(row.raw_values.get("source_row_number")) == "52"
        and row.raw_values.get("page") == 1
    )
    edited = tmp_path / "pioneers.xlsx"
    shutil.copy(pioneers_azin_excel, edited)
    workbook = load_workbook(edited)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    name_col = headers[HEADER_NAME]
    sale_col = headers[HEADER_SALE_PRICE]
    sheet.cell(2, name_col).value = item.supplier_name
    workbook.save(edited)
    workbook.close()
    before = _snapshot(edited)
    source_hash = pioneers_azin_excel.read_bytes()

    session_id = client.post("/api/sessions").json()["session_id"]
    with edited.open("rb") as handle:
        excel_resp = client.post(
            f"/api/sessions/{session_id}/pioneers-excel",
            files={
                "file": (
                    "pioneers.xlsx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert excel_resp.status_code == 200, excel_resp.text
    with abrahan_pdf.open("rb") as handle:
        pdf_resp = client.post(
            f"/api/sessions/{session_id}/supplier-pdf",
            data={"supplier_key": "abrahan_pipe"},
            files={"file": ("abrahan.pdf", handle, "application/pdf")},
        )
    assert pdf_resp.status_code == 200, pdf_resp.text
    assert pdf_resp.json()["supplier_item_count"] == 175

    match_resp = client.post(f"/api/sessions/{session_id}/match", json={})
    assert match_resp.status_code == 200, match_resp.text
    candidates = match_resp.json()["candidates"]
    keep = next(row for row in candidates if row["supplier_name"] == item.supplier_name)
    assert keep["status"] == "MATCHED"
    assert keep["new_price"] == str(item.price)
    assert keep["pricing_blockers"] == []
    if keep["warnings"]:
        review = client.post(
            f"/api/sessions/{session_id}/candidates/{keep['candidate_id']}/review",
            json={"decision": "APPROVE"},
        )
        assert review.status_code == 200, review.text
        assert review.json()["exportable"] is True
    for candidate in candidates:
        if candidate["candidate_id"] == keep["candidate_id"]:
            continue
        excluded = client.post(
            f"/api/sessions/{session_id}/candidates/{candidate['candidate_id']}/review",
            json={"decision": "EXCLUDE"},
        )
        assert excluded.status_code == 200, excluded.text

    export_resp = client.post(f"/api/sessions/{session_id}/export")
    assert export_resp.status_code == 200, export_resp.text
    assert export_resp.json()["updated_count"] == 1
    download = client.get(f"/api/sessions/{session_id}/download")
    assert download.status_code == 200
    output = tmp_path / "out.xlsx"
    output.write_bytes(download.content)
    after = _snapshot(output)
    assert max(row for row, _column in before) > 2
    assert {row for row, _column in after} == {1, 2}
    for key, value in before.items():
        if key[0] > 2:
            continue
        if key == (2, sale_col):
            assert str(after[key]).replace(",", "") == str(item.price)
        else:
            assert after[key] == value
    assert pioneers_azin_excel.read_bytes() == source_hash


def test_confirmed_metre_export_changes_only_final_sale_price(
    client, pioneers_azin_excel: Path, dina_pdf: Path, tmp_path: Path
) -> None:
    item = metre_items(dina_pdf)[0]
    assert item.unit == "متر"
    edited = tmp_path / "pioneers.xlsx"
    workbook_with_supplier_names(pioneers_azin_excel, edited, [item.supplier_name])
    workbook = load_workbook(edited)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    sale_col = headers[HEADER_SALE_PRICE]
    workbook.close()
    before = _snapshot(edited)
    source_hash = pioneers_azin_excel.read_bytes()

    session_id = client.post("/api/sessions").json()["session_id"]
    with edited.open("rb") as handle:
        excel_resp = client.post(
            f"/api/sessions/{session_id}/pioneers-excel",
            files={
                "file": (
                    "pioneers.xlsx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert excel_resp.status_code == 200, excel_resp.text
    with dina_pdf.open("rb") as handle:
        pdf_resp = client.post(
            f"/api/sessions/{session_id}/supplier-pdf",
            data={"supplier_key": "dina_polymer"},
            files={"file": ("dina.pdf", handle, "application/pdf")},
        )
    assert pdf_resp.status_code == 200, pdf_resp.text
    match_resp = client.post(f"/api/sessions/{session_id}/match", json={})
    assert match_resp.status_code == 200, match_resp.text
    candidates = match_resp.json()["candidates"]
    keep = next(row for row in candidates if row["supplier_code"] == item.supplier_code)
    assert keep["new_price"] == str(item.price)
    assert keep["new_price"] != str(item.price * 6)
    assert keep["pricing_blockers"] == []
    if keep["warnings"]:
        approved = client.post(
            f"/api/sessions/{session_id}/candidates/{keep['candidate_id']}/review",
            json={"decision": "APPROVE"},
        )
        assert approved.status_code == 200, approved.text
        assert approved.json()["exportable"] is True
    for candidate in candidates:
        if candidate["candidate_id"] == keep["candidate_id"]:
            continue
        excluded = client.post(
            f"/api/sessions/{session_id}/candidates/{candidate['candidate_id']}/review",
            json={"decision": "EXCLUDE"},
        )
        assert excluded.status_code == 200, excluded.text

    export_resp = client.post(f"/api/sessions/{session_id}/export")
    assert export_resp.status_code == 200, export_resp.text
    assert export_resp.json()["updated_count"] == 1
    download = client.get(f"/api/sessions/{session_id}/download")
    assert download.status_code == 200
    output = tmp_path / "out.xlsx"
    output.write_bytes(download.content)
    after = _snapshot(output)
    assert max(row for row, _column in before) > 2
    assert {row for row, _column in after} == {1, 2}
    for key, value in before.items():
        if key[0] > 2:
            continue
        if key == (2, sale_col):
            assert str(after[key]).replace(",", "") == str(item.price)
        else:
            assert after[key] == value
    assert pioneers_azin_excel.read_bytes() == source_hash
