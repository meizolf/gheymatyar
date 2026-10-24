"""Kian hose price list: names, Toman→Rial, and existing workflow integration."""

from __future__ import annotations

import shutil
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from gheymatyar.domain.enums import ReviewDecision
from gheymatyar.domain.product import PioneersProduct
from gheymatyar.repositories import (
    MappingRepository,
    ProductDecisionRepository,
    SupplierRepository,
)
from gheymatyar.parsers.pioneers_excel_parser import HEADER_NAME, HEADER_SALE_PRICE
from gheymatyar.services.matching import match_supplier_items
from gheymatyar.services.price_validation import apply_price_warnings, is_exportable
from gheymatyar.suppliers.detection import Confidence
from gheymatyar.suppliers.errors import SupplierPdfError
from gheymatyar.suppliers.kian import KianProcessor
from gheymatyar.suppliers.registry import get_default_registry
from tests.test_persistent_decisions_api import _restart

# Page-5 construction text. These literals are independent of the processor.
_ECONOMY = "بافت آلومینیوم دنباله برنجی اتصالات استیل بدون لیبل و کارتن"
_ALUMINUM = "بافت آلومینیوم قطعات برنجی با لیبل و کارتن"
_STEEL = "بافت استیل قطعات برنجی با لیبل و کارتن"
_GRADE_A = "تمامی قطعات برنجی زره کروم دوسربوش بلند دسته آسان شور قلاب گرد"
_GRADE_B = "کونیک زینک زره پولیش مهره استیل نری برنجی دوسربوش کوتاه قلاب ساده"
_MATERIAL_TEXT = {"اکونومی": _ECONOMY, "آلومینیوم": _ALUMINUM, "استیل": _STEEL}
_GRADE_TEXT = {"A": _GRADE_A, "B": _GRADE_B}


def _kian(body: str, material: str | None = None, grade: str | None = None) -> str:
    parts = ["شیلنگ کیان", body]
    if material:
        parts.extend((material, _MATERIAL_TEXT[material]))
    if grade:
        parts.extend((f"گرید {grade}", _GRADE_TEXT[grade]))
    return " ".join(parts)


# code, printed Toman, full Persian name
_EXAMPLES: dict[int, tuple[tuple[str, str, str], ...]] = {
    1: (
        ("101", "260800", _kian("پکیجی 1/2 30 سانت", "اکونومی")),
        ("127", "321100", _kian("پکیجی 1/2 30 سانت", "آلومینیوم")),
        ("153", "360600", _kian("پکیجی 1/2 30 سانت", "استیل")),
        ("116", "497900", _kian("آبگرمکن 1/2 دوسر مهره 100 سانت", "اکونومی")),
        ("126", "563300", _kian("آبگرمکن 1/2 دوسر مهره (یکسرخم) 70 سانت", "اکونومی")),
        ("178", "715700", _kian("آبگرمکن 1/2 دوسر مهره (یکسرخم) 70 سانت", "استیل")),
    ),
    2: (
        ("201", "204300", _kian("پیسواری 1/2 دوسر مهره 30 سانت", "اکونومی")),
        ("249", "275600", _kian("پیسواری 1/2 دوسر مهره 30 سانت", "استیل")),
        ("206", "177700", _kian("تست آچاری حصیری دوسر مهره 50 سانت", "اکونومی")),
        ("207", "188100", _kian("تست آچاری حصیری دوسر مهره 60 سانت", "اکونومی")),
        ("208", "335900", _kian("تکپایه 1/2 * M10 (جفت) 40 سانت", "اکونومی")),
        ("254", "468000", _kian("تکپایه 1/2 * M10 (جفت) 40 سانت", "استیل")),
        ("217", "386300", _kian("تکپایه تنسو 3/8 (160) 60 سانت", "اکونومی")),
        ("222", "143600", _kian("رابط 3/8 10 سانت", "اکونومی")),
    ),
    3: (
        ("301", "1592000", _kian("یک اینچ یکسر مهره * یکسر تبدیل 50 سانت", "استیل")),
        ("306", "2254400", _kian("یک اینچ یکسر مهره * یکسر تبدیل 100 سانت", "استیل")),
        ("310", "2253900", _kian("یک اینچ یکسر مهره (سرخم) * یکسر تبدیل 80 سانت", "استیل")),
        ("313", "1641200", _kian("یک اینچ دوسر مهره 50 سانت", "استیل")),
        ("319", "1907300", _kian("یک اینچ دوسر مهره (یکسرخم) 50 سانت", "استیل")),
        ("324", "2569700", _kian("یک اینچ دوسر مهره (یکسرخم) 100 سانت", "استیل")),
    ),
    4: (
        ("401", "412200", _kian("توالت ABS ریزبافت پولیش", grade="B")),
        ("402", "575900", _kian("توالت برنجی ریزبافت پولیش", grade="B")),
        ("403", "461400", _kian("دوش برنجی ریزبافت پولیش", grade="B")),
        ("404", "829900", _kian("توالت کروم", grade="B")),
        ("416", "965600", _kian("توالت کروم", grade="A")),
        ("407", "596800", _kian("دوش طلا براق", grade="B")),
        ("419", "1032700", _kian("دوش طلا براق", grade="A")),
        ("428", "226700", _kian("رابط دوش 50 سانت پولیش")),
        ("431", "355500", _kian("رابط دوش 50 سانت سفید طلا")),
        ("434", "295300", _kian("رابط دوش 50 سانت مشکی کروم")),
    ),
}


def _by_code(result):
    return {item.supplier_code: item for item in result.items}


def test_kian_parses_priced_pages_and_skips_notes(kian_pdf: Path) -> None:
    result = KianProcessor().parse_pdf(kian_pdf)
    counts = {
        page: sum(1 for item in result.items if item.raw_values["page"] == page)
        for page in range(1, 6)
    }
    assert counts == {1: 78, 2: 71, 3: 24, 4: 34, 5: 0}
    assert len(result.items) == 207
    assert result.source_metadata["pricing_rule"] == "toman_times_10_to_rial"
    assert result.source_metadata["package_multiplier_applied"] is False
    assert result.source_metadata["discount_applied"] is False
    assert result.warnings
    assert all("page 5" in warning for warning in result.warnings)
    assert len({item.supplier_code for item in result.items}) == 207
    processor = KianProcessor()
    assert len(
        {processor.normalize_supplier_name(item.supplier_name) for item in result.items}
    ) == 207
    assert all("کیان" in item.supplier_name for item in result.items)
    assert all(item.raw_values["pricing_blockers"] == [] for item in result.items)
    for item in result.items:
        material = item.raw_values["material"]
        grade = item.raw_values["grade"]
        if material:
            assert _MATERIAL_TEXT[material] in item.supplier_name
        if grade:
            assert _GRADE_TEXT[grade] in item.supplier_name
        if material is None and grade is None:
            assert "بافت" not in item.supplier_name
            assert "گرید" not in item.supplier_name
        assert item.raw_values["legacy_name"]
        assert item.raw_values["legacy_name"] != item.supplier_name
        assert "کیان" not in item.raw_values["legacy_name"]


def test_kian_examples_keep_distinguishing_names_and_toman_prices(kian_pdf: Path) -> None:
    result = KianProcessor().parse_pdf(kian_pdf)
    rows = _by_code(result)
    processor = KianProcessor()
    for page, examples in _EXAMPLES.items():
        assert len(examples) >= 3
        for code, toman, name in examples:
            item = rows[code]
            assert item.raw_values["page"] == page
            assert item.supplier_name == name
            assert item.price == Decimal(toman)
            assert item.raw_values["original_currency"] == "toman"
            assert item.raw_values["original_price"] == toman
            assert item.raw_values["currency"] == "rial"
            assert item.raw_values["price_basis"] == "printed"
            assert item.raw_values["pricing_blockers"] == []
            final = processor.calculate_final_sale_price(item)
            assert final == Decimal(toman) * Decimal("10")
            assert final != Decimal(toman) * Decimal("25")
            assert final != Decimal(toman) * Decimal("50")
            assert final != Decimal(toman) * Decimal("56")
            assert final != Decimal(toman) * Decimal("10") * Decimal("25")
            assert final != (Decimal(toman) - Decimal("50000")) * Decimal("10")

    economy = rows["101"]
    aluminum = rows["127"]
    steel = rows["153"]
    assert economy.supplier_name != aluminum.supplier_name != steel.supplier_name
    assert {economy.price, aluminum.price, steel.price} == {
        Decimal("260800"),
        Decimal("321100"),
        Decimal("360600"),
    }

    woven = [item for item in result.items if "تست آچاری" in item.supplier_name]
    assert {item.supplier_code for item in woven} == {"206", "207"}
    assert {item.raw_values["material"] for item in woven} == {"اکونومی"}

    abs_rows = [item for item in result.items if "ABS" in item.supplier_name]
    assert len(abs_rows) == 1
    assert abs_rows[0].supplier_code == "401"
    assert "گرید A" not in abs_rows[0].supplier_name
    assert rows["416"].supplier_name != rows["404"].supplier_name
    assert "طلا براق" in rows["419"].supplier_name
    assert "سفید طلا" in rows["431"].supplier_name
    assert rows["428"].supplier_name != rows["434"].supplier_name


def test_kian_rial_price_is_exportable_without_package_math(kian_pdf: Path) -> None:
    item = _by_code(KianProcessor().parse_pdf(kian_pdf))["101"]
    processor = KianProcessor()
    assert processor.calculate_final_sale_price(item) == Decimal("2608000")
    product = PioneersProduct(
        row_number=2,
        name=item.supplier_name,
        code="K-101",
        sale_price=Decimal("2608000"),
        unit="عدد",
    )
    matched = match_supplier_items(
        [item],
        [product],
        processor=processor,
    )
    enriched = apply_price_warnings(matched.candidates)
    assert enriched[0].status == "MATCHED"
    assert enriched[0].new_price == Decimal("2608000")
    assert enriched[0].source_currency == "rial"
    assert enriched[0].pricing_blockers == []
    assert is_exportable(enriched[0]) is False
    unchanged = enriched[0].model_copy(
        update={"review_decision": ReviewDecision.APPROVE}
    )
    assert is_exportable(unchanged) is False
    changed = enriched[0].model_copy(
        update={
            "review_decision": ReviewDecision.APPROVE,
            "old_price": Decimal("1"),
        }
    )
    assert is_exportable(changed) is True

    gold = _by_code(KianProcessor().parse_pdf(kian_pdf))["416"]
    assert processor.calculate_final_sale_price(gold) == Decimal("9656000")


def test_kian_is_registered_and_detected_from_content(kian_pdf: Path) -> None:
    registry = get_default_registry()
    assert "kian" in registry.keys()
    info = next(row for row in registry.list_suppliers() if row.key == "kian")
    assert info.display_name == "گروه تولیدی صنعتی کیان"
    detected = registry.detect(kian_pdf)
    assert detected.detected_supplier_id == "kian"
    assert detected.detected_supplier_name == "گروه تولیدی صنعتی کیان"
    assert detected.confidence is Confidence.HIGH
    assert detected.requires_confirmation is True
    assert detected.scores["top_industrial"] == 0
    registry.validate_pdf_for_supplier(kian_pdf, "kian")
    with pytest.raises(SupplierPdfError, match="does not identify"):
        registry.validate_pdf_for_supplier(kian_pdf, "top_industrial")


def test_kian_detection_waits_for_confirmation(client, kian_pdf: Path) -> None:
    session_id = client.post("/api/sessions").json()["session_id"]
    with kian_pdf.open("rb") as handle:
        detected = client.post(
            f"/api/sessions/{session_id}/supplier-pdf/detect",
            files={"file": ("price-list.pdf", handle, "application/pdf")},
        )
    assert detected.status_code == 200, detected.text
    body = detected.json()
    assert body["supplier_confirmed"] is False
    assert body["detection"]["detected_supplier_id"] == "kian"
    assert body["detection"]["requires_confirmation"] is True
    assert body["summary"]["supplier_item_count"] == 0

    confirmed = client.post(
        f"/api/sessions/{session_id}/supplier-pdf/confirm",
        json={"supplier_key": "kian"},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["supplier_confirmed"] is True
    assert confirmed.json()["supplier_item_count"] == 207
    listed = client.get("/api/suppliers").json()
    assert any(row["key"] == "kian" for row in listed)


def test_kian_mapping_and_ignore_persist(client, kian_pdf: Path, pioneers_azin_excel: Path) -> None:
    session_id = client.post("/api/sessions").json()["session_id"]
    with pioneers_azin_excel.open("rb") as handle:
        excel = client.post(
            f"/api/sessions/{session_id}/pioneers-excel",
            files={
                "file": (
                    "pioneers.xlsx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert excel.status_code == 200, excel.text
    with kian_pdf.open("rb") as handle:
        parsed = client.post(
            f"/api/sessions/{session_id}/supplier-pdf",
            data={"supplier_key": "kian"},
            files={"file": ("kian.pdf", handle, "application/pdf")},
        )
    assert parsed.status_code == 200, parsed.text
    matched = client.post(f"/api/sessions/{session_id}/match", json={})
    assert matched.status_code == 200, matched.text
    rows = matched.json()["candidates"]
    keep = next(row for row in rows if row["supplier_code"] == "101")
    drop = next(row for row in rows if row["supplier_code"] == "127")
    assert keep["new_price"] == "2608000"
    assert drop["new_price"] == "3211000"

    ignored = client.post(
        f"/api/sessions/{session_id}/candidates/{drop['candidate_id']}/ignore"
    )
    assert ignored.status_code == 200, ignored.text
    assert ignored.json()["status"] == "IGNORED"

    products = client.app.state.processing.get_session(session_id).products
    target = products[0]
    mapped = client.post(
        f"/api/sessions/{session_id}/candidates/{keep['candidate_id']}/confirm-mapping",
        json={"pioneers_row_number": target.row_number, "save": True},
    )
    assert mapped.status_code == 200, mapped.text
    assert mapped.json()["match_type"] == "MANUAL"
    assert mapped.json()["pioneers_code"] == target.code

    _restart(client)
    reloaded = client.post("/api/sessions").json()["session_id"]
    with pioneers_azin_excel.open("rb") as handle:
        client.post(
            f"/api/sessions/{reloaded}/pioneers-excel",
            files={
                "file": (
                    "pioneers.xlsx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    with kian_pdf.open("rb") as handle:
        client.post(
            f"/api/sessions/{reloaded}/supplier-pdf",
            data={"supplier_key": "kian"},
            files={"file": ("kian.pdf", handle, "application/pdf")},
        )
    again = client.post(f"/api/sessions/{reloaded}/match", json={})
    assert again.status_code == 200, again.text
    fresh = {row["supplier_code"]: row for row in again.json()["candidates"]}
    assert fresh["127"]["status"] == "IGNORED"
    assert fresh["101"]["match_type"] == "SAVED_MAPPING"
    assert fresh["101"]["pioneers_code"] == target.code
    assert fresh["101"]["new_price"] == "2608000"
    assert fresh["153"]["status"] != "IGNORED"


def _snapshot(path: Path) -> dict[tuple[int, int], object]:
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook.active
        return {
            (cell.row, cell.column): cell.value
            for row in sheet.iter_rows()
            for cell in row
        }
    finally:
        workbook.close()


def test_kian_export_changes_only_sale_column_in_rial(
    client, pioneers_azin_excel: Path, kian_pdf: Path, tmp_path: Path
) -> None:
    item = _by_code(KianProcessor().parse_pdf(kian_pdf))["416"]
    assert item.price == Decimal("965600")
    edited = tmp_path / "pioneers.xlsx"
    shutil.copy(pioneers_azin_excel, edited)
    workbook = load_workbook(edited)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    name_col = headers[HEADER_NAME]
    sale_col = headers[HEADER_SALE_PRICE]
    sheet.cell(2, name_col).value = item.supplier_name
    workbook.save(edited)
    workbook.close()
    before = _snapshot(edited)
    source_hash = pioneers_azin_excel.read_bytes()

    session_id = client.post("/api/sessions").json()["session_id"]
    with edited.open("rb") as handle:
        excel_resp = client.post(
            f"/api/sessions/{session_id}/pioneers-excel",
            files={
                "file": (
                    "pioneers.xlsx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert excel_resp.status_code == 200, excel_resp.text
    with kian_pdf.open("rb") as handle:
        pdf_resp = client.post(
            f"/api/sessions/{session_id}/supplier-pdf",
            data={"supplier_key": "kian"},
            files={"file": ("kian.pdf", handle, "application/pdf")},
        )
    assert pdf_resp.status_code == 200, pdf_resp.text
    match_resp = client.post(f"/api/sessions/{session_id}/match", json={})
    assert match_resp.status_code == 200, match_resp.text
    candidates = match_resp.json()["candidates"]
    keep = next(row for row in candidates if row["supplier_code"] == "416")
    assert keep["status"] == "MATCHED"
    assert keep["new_price"] == "9656000"
    assert keep["pricing_blockers"] == []
    if keep["warnings"]:
        review = client.post(
            f"/api/sessions/{session_id}/candidates/{keep['candidate_id']}/review",
            json={"decision": "APPROVE"},
        )
        assert review.status_code == 200, review.text
        assert review.json()["exportable"] is True
    else:
        assert keep["exportable"] is True

    export_resp = client.post(f"/api/sessions/{session_id}/export")
    assert export_resp.status_code == 200, export_resp.text
    assert export_resp.json()["updated_count"] == 1
    download = client.get(f"/api/sessions/{session_id}/download")
    assert download.status_code == 200
    output = tmp_path / "out.xlsx"
    output.write_bytes(download.content)
    after = _snapshot(output)
    assert max(row for row, _column in before) > 2
    assert {row for row, _column in after} == {1, 2}
    for key, value in before.items():
        if key[0] > 2:
            continue
        if key == (2, sale_col):
            assert str(after[key]).replace(",", "") == "9656000"
        else:
            assert after[key] == value
    assert pioneers_azin_excel.read_bytes() == source_hash


def _mutate_pdf(source: Path, dest: Path, edit) -> Path:
    import pymupdf

    with pymupdf.open(source) as doc:
        edit(doc[0])
        doc.save(dest)
    return dest


def _insert_price_token(page, token: str, x: float) -> None:
    page.insert_text((x, 170), token, fontsize=5)


def _upload_kian(client, excel: Path, pdf: Path) -> str:
    session_id = client.post("/api/sessions").json()["session_id"]
    with excel.open("rb") as handle:
        uploaded = client.post(
            f"/api/sessions/{session_id}/pioneers-excel",
            files={
                "file": (
                    "pioneers.xlsx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert uploaded.status_code == 200, uploaded.text
    with pdf.open("rb") as handle:
        parsed = client.post(
            f"/api/sessions/{session_id}/supplier-pdf",
            data={"supplier_key": "kian"},
            files={"file": ("kian.pdf", handle, "application/pdf")},
        )
    assert parsed.status_code == 200, parsed.text
    return session_id


def _excel_named(source: Path, dest: Path, product_name: str) -> Path:
    shutil.copy(source, dest)
    workbook = load_workbook(dest)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers[HEADER_NAME]).value = product_name
    workbook.save(dest)
    workbook.close()
    return dest


def test_kian_invalid_price_suffix_blocks_export(
    client, kian_pdf: Path, pioneers_azin_excel: Path, tmp_path: Path
) -> None:
    mutated = _mutate_pdf(
        kian_pdf,
        tmp_path / "invalid.pdf",
        lambda page: _insert_price_token(page, "#N/A", 329),
    )
    item = _by_code(KianProcessor().parse_pdf(mutated))["101"]
    assert item.price == Decimal("260800")
    assert "#N/A" in item.raw_values["original_price_text"]
    assert item.raw_values["pricing_blockers"]
    assert "not a single amount" in item.raw_values["pricing_blockers"][0]

    excel = _excel_named(pioneers_azin_excel, tmp_path / "pioneers.xlsx", item.supplier_name)
    session_id = _upload_kian(client, excel, mutated)
    matched = client.post(f"/api/sessions/{session_id}/match", json={})
    assert matched.status_code == 200, matched.text
    rows = matched.json()["candidates"]
    blocked = next(row for row in rows if row["supplier_code"] == "101")
    assert blocked["status"] == "MATCHED"
    assert blocked["exportable"] is False
    assert blocked["pricing_blockers"]
    approved = client.post(
        f"/api/sessions/{session_id}/candidates/{blocked['candidate_id']}/review",
        json={"decision": "APPROVE"},
    )
    assert approved.status_code == 400, approved.text
    assert "unresolved" in approved.text
    for row in rows:
        if row["candidate_id"] == blocked["candidate_id"]:
            continue
        excluded = client.post(
            f"/api/sessions/{session_id}/candidates/{row['candidate_id']}/review",
            json={"decision": "EXCLUDE"},
        )
        assert excluded.status_code == 200, excluded.text
    exported = client.post(f"/api/sessions/{session_id}/export")
    assert exported.status_code == 400, exported.text


def _price_word(text: str, x: float) -> tuple:
    width = max(len(text) * 4, 6)
    return (x, 160.0, x + width, 172.0, text, 0, 0, 0)


def test_kian_dash_only_cell_is_missing_and_mixed_dash_is_blocked() -> None:
    from gheymatyar.suppliers.kian import _read_price_cell

    dashed = _read_price_cell(
        [_price_word("----------", 100), _price_word("-", 180), _price_word("−", 200)]
    )
    assert dashed.kind == "dashed"
    assert dashed.amount is None
    assert dashed.blocker is None
    assert dashed.text == "---------- - −"
    assert _read_price_cell([]).kind == "absent"

    orders = (
        [_price_word("-", 100), _price_word("260,800", 140)],
        [_price_word("260,800", 100), _price_word("-", 180)],
        [_price_word("−", 90), _price_word("260,800", 140)],
    )
    for words in orders:
        cell = _read_price_cell(words)
        assert cell.kind == "blocked"
        assert cell.amount == Decimal("260800")
        assert "260,800" in cell.text
        assert cell.text != "260,800"
        assert cell.blocker is not None
        assert cell.text in cell.blocker
        assert "not a single amount" in cell.blocker


@pytest.mark.parametrize("x", [283, 329])
def test_kian_mixed_dash_blocks_approval_and_manual_mapping(
    client, kian_pdf: Path, pioneers_azin_excel: Path, tmp_path: Path, x: float
) -> None:
    source_bytes = pioneers_azin_excel.read_bytes()
    mutated = _mutate_pdf(
        kian_pdf,
        tmp_path / f"dash-{x}.pdf",
        lambda page: _insert_price_token(page, "-", x),
    )
    item = _by_code(KianProcessor().parse_pdf(mutated))["101"]
    assert item.price == Decimal("260800")
    assert item.raw_values["original_price_text"] != "260,800"
    assert "-" in item.raw_values["original_price_text"]
    assert "260,800" in item.raw_values["original_price_text"]
    assert item.raw_values["pricing_blockers"]
    assert "not a single amount" in item.raw_values["pricing_blockers"][0]

    excel = _excel_named(
        pioneers_azin_excel, tmp_path / f"pioneers-{x}.xlsx", item.supplier_name
    )
    session_id = _upload_kian(client, excel, mutated)
    matched = client.post(f"/api/sessions/{session_id}/match", json={})
    assert matched.status_code == 200, matched.text
    rows = matched.json()["candidates"]
    blocked = next(row for row in rows if row["supplier_code"] == "101")
    assert blocked["status"] == "MATCHED"
    assert blocked["exportable"] is False
    assert blocked["pricing_blockers"]

    mapped = client.post(
        f"/api/sessions/{session_id}/candidates/{blocked['candidate_id']}/confirm-mapping",
        json={"pioneers_row_number": 2, "save": True},
    )
    assert mapped.status_code == 200, mapped.text
    assert mapped.json()["match_type"] == "MANUAL"
    assert mapped.json()["exportable"] is False
    assert mapped.json()["pricing_blockers"]

    approved = client.post(
        f"/api/sessions/{session_id}/candidates/{blocked['candidate_id']}/review",
        json={"decision": "APPROVE"},
    )
    assert approved.status_code == 400, approved.text
    assert "unresolved" in approved.text
    for row in rows:
        if row["candidate_id"] == blocked["candidate_id"]:
            continue
        excluded = client.post(
            f"/api/sessions/{session_id}/candidates/{row['candidate_id']}/review",
            json={"decision": "EXCLUDE"},
        )
        assert excluded.status_code == 200, excluded.text
    exported = client.post(f"/api/sessions/{session_id}/export")
    assert exported.status_code == 400, exported.text
    assert pioneers_azin_excel.read_bytes() == source_bytes


def test_kian_split_decimal_becomes_rial(
    client, kian_pdf: Path, pioneers_azin_excel: Path, tmp_path: Path
) -> None:
    mutated = _mutate_pdf(
        kian_pdf,
        tmp_path / "decimal.pdf",
        lambda page: _insert_price_token(page, ".5", 332),
    )
    processor = KianProcessor()
    item = _by_code(processor.parse_pdf(mutated))["101"]
    assert item.price == Decimal("260800.5")
    assert item.raw_values["pricing_blockers"] == []
    assert processor.calculate_final_sale_price(item) == Decimal("2608005")

    excel = _excel_named(pioneers_azin_excel, tmp_path / "pioneers.xlsx", item.supplier_name)
    before = _snapshot(excel)
    session_id = _upload_kian(client, excel, mutated)
    matched = client.post(f"/api/sessions/{session_id}/match", json={})
    assert matched.status_code == 200, matched.text
    rows = matched.json()["candidates"]
    keep = next(row for row in rows if row["supplier_code"] == "101")
    assert keep["status"] == "MATCHED"
    assert keep["new_price"] == "2608005"
    assert keep["pricing_blockers"] == []
    if not keep["exportable"]:
        approved = client.post(
            f"/api/sessions/{session_id}/candidates/{keep['candidate_id']}/review",
            json={"decision": "APPROVE"},
        )
        assert approved.status_code == 200, approved.text
        assert approved.json()["exportable"] is True
    for row in rows:
        if row["candidate_id"] == keep["candidate_id"]:
            continue
        excluded = client.post(
            f"/api/sessions/{session_id}/candidates/{row['candidate_id']}/review",
            json={"decision": "EXCLUDE"},
        )
        assert excluded.status_code == 200, excluded.text
    exported = client.post(f"/api/sessions/{session_id}/export")
    assert exported.status_code == 200, exported.text
    assert exported.json()["updated_count"] == 1
    output = tmp_path / "out.xlsx"
    output.write_bytes(client.get(f"/api/sessions/{session_id}/download").content)
    after = _snapshot(output)
    sale_col = next(
        column for (row, column), value in before.items() if row == 1 and value == HEADER_SALE_PRICE
    )
    assert max(row for row, _column in before) > 2
    assert {row for row, _column in after} == {1, 2}
    for key, value in before.items():
        if key[0] > 2:
            continue
        if key == (2, sale_col):
            assert str(after[key]).replace(",", "") == "2608005"
        else:
            assert after[key] == value


def test_kian_missing_price_creates_no_product(kian_pdf: Path, tmp_path: Path) -> None:
    import pymupdf

    def redact(page) -> None:
        hits = [word for word in page.get_text("words") if word[4] == "260,800"]
        assert hits
        word = min(hits, key=lambda item: abs(((item[1] + item[3]) / 2) - 165))
        page.add_redact_annot(
            pymupdf.Rect(word[0] - 0.5, word[1] - 0.5, word[2] + 0.5, word[3] + 0.5),
            fill=(1, 1, 1),
        )
        page.apply_redactions(images=0, graphics=0)

    mutated = _mutate_pdf(kian_pdf, tmp_path / "missing.pdf", redact)
    result = KianProcessor().parse_pdf(mutated)
    assert "101" not in {item.supplier_code for item in result.items}
    assert any("101" in warning for warning in result.warnings)
    assert len(result.items) == 206
    counts = {
        page: sum(1 for item in result.items if item.raw_values["page"] == page)
        for page in range(1, 6)
    }
    assert counts[1] == 77
    assert counts[5] == 0


def test_kian_prior_name_keeps_mapping_and_ignore(
    client, kian_pdf: Path, pioneers_azin_excel: Path
) -> None:
    processor = KianProcessor()
    parsed = processor.parse_pdf(kian_pdf)
    by_code = _by_code(parsed)
    session_id = _upload_kian(client, pioneers_azin_excel, kian_pdf)
    prepared = client.post(f"/api/sessions/{session_id}/match", json={})
    assert prepared.status_code == 200, prepared.text
    products = client.app.state.processing.get_session(session_id).products
    assert len(products) >= 2
    mapped_product = products[0]
    decoy_product = products[1]

    def legacy(code: str) -> str:
        return processor.normalize_supplier_name(by_code[code].raw_values["legacy_name"])

    with client.app.state.database.session() as db:
        supplier = next(
            row for row in SupplierRepository(db).list_all() if row.parser_type == "kian"
        )
        ProductDecisionRepository(db).ignore(
            supplier_id=supplier.id,
            supplier_name_raw=by_code["153"].raw_values["legacy_name"],
            supplier_name_normalized=legacy("153"),
            supplier_code="153",
        )
        ProductDecisionRepository(db).ignore(
            supplier_id=supplier.id,
            supplier_name_raw=by_code["101"].raw_values["legacy_name"],
            supplier_name_normalized=legacy("101"),
            supplier_code="999",
        )
        MappingRepository(db).upsert(
            supplier_id=supplier.id,
            supplier_name_raw=by_code["101"].raw_values["legacy_name"],
            supplier_name_normalized=legacy("101"),
            supplier_code="101",
            pioneers_name_raw=mapped_product.name,
            pioneers_code=mapped_product.code,
        )
        MappingRepository(db).upsert(
            supplier_id=supplier.id,
            supplier_name_raw="کالای دیگر",
            supplier_name_normalized=processor.normalize_supplier_name("کالای دیگر"),
            supplier_code="101",
            pioneers_name_raw=decoy_product.name,
            pioneers_code=decoy_product.code,
        )

    _restart(client)
    reloaded = _upload_kian(client, pioneers_azin_excel, kian_pdf)
    again = client.post(f"/api/sessions/{reloaded}/match", json={})
    assert again.status_code == 200, again.text
    fresh = {row["supplier_code"]: row for row in again.json()["candidates"]}
    assert fresh["153"]["status"] == "IGNORED"
    restored = client.post(
        f"/api/sessions/{reloaded}/candidates/{fresh['153']['candidate_id']}/restore"
    )
    assert restored.status_code == 200, restored.text
    assert restored.json()["status"] != "IGNORED"
    assert fresh["101"]["status"] == "MATCHED"
    assert fresh["101"]["match_type"] == "SAVED_MAPPING"
    assert fresh["101"]["pioneers_code"] == mapped_product.code
    assert fresh["101"]["pioneers_code"] != decoy_product.code
    assert fresh["127"]["status"] != "IGNORED"
    assert fresh["127"]["match_type"] != "SAVED_MAPPING"
