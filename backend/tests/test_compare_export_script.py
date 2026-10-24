"""Smoke helpers for packaged export verification."""

from __future__ import annotations

import hashlib
import importlib.util
from pathlib import Path

import pytest
from openpyxl import Workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
COMPARE_PATH = REPO_ROOT / "scripts" / "windows" / "compare_export.py"


def _load_compare():
    spec = importlib.util.spec_from_file_location("compare_export", COMPARE_PATH)
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def compare_mod():
    return _load_compare()


def _write_basic(path: Path, *, sale: str = "100", name: str = "الف") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(["کد", "نرخ فروش اصلی", "نام"])
    ws.append(["1", sale, name])
    wb.create_sheet("Meta").append(["note", "unchanged"])
    wb.save(path)
    wb.close()


def test_compare_export_accepts_sale_only_change(tmp_path: Path, compare_mod) -> None:
    source = tmp_path / "src.xlsx"
    exported = tmp_path / "out.xlsx"
    _write_basic(source, sale="100")
    _write_basic(exported, sale="150")
    compare_mod.compare(source, exported, _sha(source))


def test_compare_export_accepts_filtered_product_rows(
    tmp_path: Path, compare_mod
) -> None:
    source = tmp_path / "src.xlsx"
    exported = tmp_path / "out.xlsx"
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(["کد", "نرخ فروش اصلی", "نام"])
    ws.append(["1", "100", "الف"])
    ws.append(["2", "200", "ب"])
    ws.append(["3", "300", "پ"])
    wb.create_sheet("Meta").append(["note", "unchanged"])
    wb.save(source)
    wb.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(["کد", "نرخ فروش اصلی", "نام"])
    ws.append(["2", "250", "ب"])
    wb.create_sheet("Meta").append(["note", "unchanged"])
    wb.save(exported)
    wb.close()
    compare_mod.compare(source, exported, _sha(source))


def test_compare_export_rejects_other_column_change(
    tmp_path: Path, compare_mod
) -> None:
    source = tmp_path / "src.xlsx"
    exported = tmp_path / "out.xlsx"
    _write_basic(source, name="الف")
    _write_basic(exported, name="دیگر")
    with pytest.raises(SystemExit, match="cell changed"):
        compare_mod.compare(source, exported, _sha(source))


def test_compare_export_rejects_formula_change(tmp_path: Path, compare_mod) -> None:
    source = tmp_path / "src.xlsx"
    exported = tmp_path / "out.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(["کد", "نرخ فروش اصلی", "نام"])
    ws.append(["1", "100", "الف"])
    ws["C2"] = "=A2"
    meta = wb.create_sheet("Meta")
    meta["A1"] = "=1+1"
    wb.save(source)
    wb.close()

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "Products"
    ws2.append(["کد", "نرخ فروش اصلی", "نام"])
    ws2.append(["1", "150", "الف"])
    ws2["C2"] = "=A2"
    meta2 = wb2.create_sheet("Meta")
    meta2["A1"] = "=2+2"  # formula changed on non-product sheet
    wb2.save(exported)
    wb2.close()

    with pytest.raises(SystemExit, match="cell changed"):
        compare_mod.compare(source, exported, _sha(source))


def test_compare_export_rejects_extra_sheet_change(
    tmp_path: Path, compare_mod
) -> None:
    source = tmp_path / "src.xlsx"
    exported = tmp_path / "out.xlsx"
    _write_basic(source, sale="100")
    _write_basic(exported, sale="150")
    from openpyxl import load_workbook

    wb = load_workbook(exported)
    wb["Meta"]["B1"] = "tampered"
    wb.save(exported)
    wb.close()

    with pytest.raises(SystemExit, match="cell changed"):
        compare_mod.compare(source, exported, _sha(source))
