"""Tests for the safe Pioneers final-sale-price writer."""

from __future__ import annotations

import hashlib
import shutil
import xml.etree.ElementTree as ET
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from gheymatyar.domain import FinalSalePriceUpdate
from gheymatyar.parsers.pioneers_excel_parser import (
    HEADER_CODE,
    HEADER_NAME,
    HEADER_SALE_PRICE,
    PRODUCT_SHEET_NAME,
)
from gheymatyar.services.errors import (
    AmbiguousProductSheetError,
    DuplicateUpdateError,
    InvalidPriceError,
    OutputAlreadyExistsError,
    PioneersExcelWriteError,
    ProductIdentityMismatchError,
    ProductSheetNotFoundError,
    SourceOverwriteError,
    UnsafePriceCellError,
)
from gheymatyar.services.pioneers_excel_writer import update_final_sale_prices

REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURE_PATH = REPO_ROOT / "fixtures" / "pioneers" / "pioneers_sample.xlsx"

NEIGHBOR_PRICE_HEADERS = (
    "نرخ آخرین خرید",
    "نرخ بهای تمام شده",
    "نرخ خرید روز",
    "نرخ مصرف کننده",
    "نرخ همکاری",
    "نرخ لوله کش",
)

# Row 8 in the anonymized fixture.
BARCODE_ROW = 8
BARCODE_NAME = "Test Barcode Zero"
BARCODE_CODE = "90004"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _copy_fixture(directory: Path, name: str = "source.xlsx") -> Path:
    destination = directory / name
    shutil.copy2(FIXTURE_PATH, destination)
    return destination


def _price(
    row_number: int,
    pioneers_name: str,
    final_sale_price: object,
    pioneers_code: str | None = None,
    *,
    include_code: bool = True,
) -> dict[str, object]:
    payload: dict[str, object] = {
        "row_number": row_number,
        "pioneers_name": pioneers_name,
        "final_sale_price": final_sale_price,
    }
    if include_code:
        payload["pioneers_code"] = pioneers_code
    return payload


def _barcode_update(
    price: object = Decimal("4950000"),
    *,
    include_code: bool = True,
    pioneers_code: str | None = BARCODE_CODE,
    pioneers_name: str = BARCODE_NAME,
    row_number: int = BARCODE_ROW,
) -> dict[str, object]:
    return _price(
        row_number,
        pioneers_name,
        price,
        pioneers_code,
        include_code=include_code,
    )


def _cell_values(path: Path) -> dict[tuple[str, str], object]:
    workbook = load_workbook(path, data_only=False)
    try:
        values: dict[tuple[str, str], object] = {}
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    values[(worksheet.title, cell.coordinate)] = cell.value
        return values
    finally:
        workbook.close()


def _changed_cells(
    source: Path, output: Path
) -> dict[tuple[str, str], tuple[object, object]]:
    before = _cell_values(source)
    after = _cell_values(output)
    keys = set(before) | set(after)
    return {
        key: (before.get(key), after.get(key))
        for key in keys
        if before.get(key) != after.get(key)
    }


def _header_map(path: Path, sheet_name: str | None = None) -> dict[str, str]:
    workbook = load_workbook(path, read_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        mapping: dict[str, str] = {}
        for cell in worksheet[1]:
            if isinstance(cell.value, str) and cell.value.strip():
                letters = "".join(ch for ch in cell.coordinate if ch.isalpha())
                mapping[cell.value.strip()] = letters
        return mapping
    finally:
        workbook.close()


@pytest.mark.parametrize("sheet_name", [PRODUCT_SHEET_NAME, "نمونه کالا"])
def test_product_sheet_detected_by_headers_not_title(
    tmp_path: Path, sheet_name: str
) -> None:
    source = _copy_fixture(tmp_path, f"{sheet_name}.xlsx")
    workbook = load_workbook(source)
    workbook.worksheets[0].title = sheet_name
    workbook.save(source)
    workbook.close()

    output = tmp_path / "out.xlsx"
    update_final_sale_prices(source, output, [_barcode_update()])
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == [sheet_name]
        assert workbook[sheet_name]["AK1"].value == HEADER_SALE_PRICE
        assert workbook[sheet_name].max_row == 2
        assert workbook[sheet_name]["D2"].value == BARCODE_CODE
        assert workbook[sheet_name]["AK2"].value == "4950000"
    finally:
        workbook.close()


def test_no_matching_product_sheet_fails(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    workbook[PRODUCT_SHEET_NAME]["AK1"] = "نرخ فروش بازنشسته"
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    with pytest.raises(ProductSheetNotFoundError) as exc_info:
        update_final_sale_prices(source, output, [_barcode_update()])
    assert HEADER_SALE_PRICE in str(exc_info.value)
    assert HEADER_CODE in str(exc_info.value)
    assert HEADER_NAME in str(exc_info.value)
    assert not output.exists()
    assert _sha256(source) == before


def test_ambiguous_product_sheets_fail(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    copy = workbook.copy_worksheet(workbook[PRODUCT_SHEET_NAME])
    copy.title = "نمونه کالا"
    copy.sheet_state = "hidden"
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    with pytest.raises(AmbiguousProductSheetError) as exc_info:
        update_final_sale_prices(source, output, [_barcode_update()])
    message = str(exc_info.value)
    assert PRODUCT_SHEET_NAME in message
    assert "نمونه کالا" in message
    assert not output.exists()
    assert _sha256(source) == before


def test_single_final_price_update(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(source, output, [_barcode_update()])
    assert result.updated_count == 1
    assert result.exported_row_count == 1
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 2
        assert sheet["D2"].value == BARCODE_CODE
        assert sheet["E2"].value == BARCODE_NAME
        assert sheet["AK2"].value == "4950000"
    finally:
        workbook.close()


def test_multiple_final_price_updates(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    update_final_sale_prices(
        source,
        output,
        [
            _price(3, "Test Fitting Zero", Decimal("920000"), "001234"),
            _barcode_update(Decimal("525000")),
        ],
    )
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 3
        assert sheet["D2"].value == "001234"
        assert sheet["AK2"].value == "920000"
        assert sheet["D3"].value == BARCODE_CODE
        assert sheet["AK3"].value == "525000"
    finally:
        workbook.close()


def test_wrong_name_on_correct_row_fails(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    with pytest.raises(ProductIdentityMismatchError) as exc_info:
        update_final_sale_prices(
            source,
            output,
            [_barcode_update(pioneers_name="زانو 25")],
        )
    assert "زانو 25" in str(exc_info.value)
    assert BARCODE_NAME in str(exc_info.value)
    assert not output.exists()
    assert _sha256(source) == before


def test_repeated_spaces_in_name_still_match(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    workbook[PRODUCT_SHEET_NAME]["E8"] = "Test  Barcode   Zero"
    workbook.save(source)
    workbook.close()
    output = tmp_path / "out.xlsx"
    update_final_sale_prices(source, output, [_barcode_update()])
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 2
        assert sheet["E2"].value == "Test  Barcode   Zero"
        assert sheet["AK2"].value == "4950000"
    finally:
        workbook.close()


def test_code_mismatch_fails_when_name_matches(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    with pytest.raises(ProductIdentityMismatchError) as exc_info:
        update_final_sale_prices(
            source,
            output,
            [_price(3, "Test Fitting Zero", Decimal("920000"), "1234")],
        )
    message = str(exc_info.value)
    assert "1234" in message
    assert "001234" in message
    assert not output.exists()


def test_update_without_product_code(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(
        source,
        output,
        [_barcode_update(include_code=False)],
    )
    assert result.changes[0].pioneers_code is None
    assert result.changes[0].pioneers_name == BARCODE_NAME
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 2
        assert sheet["AK2"].value == "4950000"
        assert sheet["D2"].value == BARCODE_CODE
    finally:
        workbook.close()


def test_row_order_is_taken_from_the_current_workbook(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    row_three = [sheet.cell(3, column).value for column in range(1, 41)]
    row_eight = [sheet.cell(8, column).value for column in range(1, 41)]
    for column, value in enumerate(row_eight, start=1):
        sheet.cell(3, column).value = value
    for column, value in enumerate(row_three, start=1):
        sheet.cell(8, column).value = value
    workbook.save(source)
    workbook.close()

    output = tmp_path / "out.xlsx"
    update_final_sale_prices(
        source,
        output,
        [_price(3, BARCODE_NAME, Decimal("4950000"), BARCODE_CODE)],
    )
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 2
        assert sheet["D2"].value == BARCODE_CODE
        assert sheet["E2"].value == BARCODE_NAME
        assert sheet["AK2"].value == "4950000"
        codes = [sheet.cell(row, 4).value for row in range(2, sheet.max_row + 1)]
        assert "001234" not in codes
    finally:
        workbook.close()


def test_row_far_below_the_sample_fixture_is_updated(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    for column in range(1, 41):
        sheet.cell(500, column).value = sheet.cell(8, column).value
    sheet["E500"] = "Far Row Product"
    sheet["D500"] = "98000"
    sheet["AK500"] = "1000"
    workbook.save(source)
    workbook.close()

    output = tmp_path / "out.xlsx"
    update_final_sale_prices(
        source,
        output,
        [_price(500, "Far Row Product", Decimal("880000"), "98000")],
    )
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 2
        assert sheet["D2"].value == "98000"
        assert sheet["E2"].value == "Far Row Product"
        assert sheet["AK2"].value == "880000"
    finally:
        workbook.close()


def test_other_price_columns_unchanged(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    update_final_sale_prices(source, output, [_barcode_update()])
    before = load_workbook(source)
    after = load_workbook(output)
    try:
        headers = _header_map(source)
        assert after[PRODUCT_SHEET_NAME].max_row == 2
        for header in NEIGHBOR_PRICE_HEADERS:
            letter = headers[header]
            assert after[PRODUCT_SHEET_NAME][f"{letter}1"].value == (
                before[PRODUCT_SHEET_NAME][f"{letter}1"].value
            )
            assert after[PRODUCT_SHEET_NAME][f"{letter}2"].value == (
                before[PRODUCT_SHEET_NAME][f"{letter}8"].value
            )
        assert after[PRODUCT_SHEET_NAME]["AH2"].value == "5000"
        assert after[PRODUCT_SHEET_NAME]["AI2"].value == "0"
        assert after[PRODUCT_SHEET_NAME]["AL2"].value == "8000"
        assert after[PRODUCT_SHEET_NAME]["AM2"].value == "7000"
    finally:
        before.close()
        after.close()


def test_product_metadata_unchanged(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    update_final_sale_prices(
        source,
        output,
        [_price(3, "Test Fitting Zero", Decimal("920000"), "001234")],
    )
    before = load_workbook(source)
    after = load_workbook(output)
    try:
        assert after[PRODUCT_SHEET_NAME].max_row == 2
        for letter in ("D", "E", "F", "H", "AF", "C"):
            assert after[PRODUCT_SHEET_NAME][f"{letter}2"].value == (
                before[PRODUCT_SHEET_NAME][f"{letter}3"].value
            )
        assert after[PRODUCT_SHEET_NAME]["D2"].value == "001234"
        assert after[PRODUCT_SHEET_NAME]["E2"].value == "Test Fitting Zero"
        assert after[PRODUCT_SHEET_NAME]["F2"].value == "001122334455"
    finally:
        before.close()
        after.close()


def test_unrelated_rows_unchanged(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    update_final_sale_prices(source, output, [_barcode_update()])
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 2
        assert sheet["D2"].value == BARCODE_CODE
        assert sheet["AK2"].value == "4950000"
        codes = [sheet.cell(row, 4).value for row in range(2, sheet.max_row + 1)]
        assert codes == [BARCODE_CODE]
    finally:
        workbook.close()


def test_cell_diff_only_requested_sale_prices(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    update_final_sale_prices(
        source,
        output,
        [
            _price(3, "Test Fitting Zero", Decimal("920000"), "001234"),
            _price(6, "Test Precision Price", Decimal("990000"), "90003"),
        ],
    )
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 3
        assert sheet["D2"].value == "001234"
        assert sheet["AK2"].value == "920000"
        assert sheet["D3"].value == "90003"
        assert sheet["AK3"].value == "990000"
    finally:
        workbook.close()


def test_source_hash_unchanged(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    before = _sha256(source)
    update_final_sale_prices(source, output, [_barcode_update()])
    assert _sha256(source) == before
    assert _sha256(FIXTURE_PATH) == before


def test_same_source_and_output_rejected(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    before = _sha256(source)
    with pytest.raises(SourceOverwriteError):
        update_final_sale_prices(source, source, [_barcode_update()])
    assert _sha256(source) == before


def test_equivalent_source_and_output_paths_rejected(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    source = data_dir / "pioneers.xlsx"
    shutil.copy2(FIXTURE_PATH, source)
    output = data_dir / "folder" / ".." / "pioneers.xlsx"
    before = _sha256(source)
    with pytest.raises(SourceOverwriteError):
        update_final_sale_prices(source, output, [_barcode_update()])
    assert _sha256(source) == before


def test_existing_output_file_rejected(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    output.write_bytes(b"already here")
    with pytest.raises(OutputAlreadyExistsError):
        update_final_sale_prices(source, output, [_barcode_update()])
    assert output.read_bytes() == b"already here"


def test_duplicate_target_row_with_the_same_price_is_one_row(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(
        source,
        output,
        [_barcode_update(Decimal("4500000")), _barcode_update(Decimal("4500000"))],
    )
    assert result.updated_count == 1
    assert result.exported_row_count == 1
    workbook = load_workbook(output)
    try:
        assert workbook[PRODUCT_SHEET_NAME].max_row == 2
        assert workbook[PRODUCT_SHEET_NAME]["AK2"].value == "4500000"
    finally:
        workbook.close()


def test_conflicting_prices_for_one_product_do_not_export(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    with pytest.raises(DuplicateUpdateError) as exc_info:
        update_final_sale_prices(
            source,
            output,
            [_barcode_update(Decimal("4500000")), _barcode_update(Decimal("4600000"))],
        )
    assert "conflicting final sale prices" in str(exc_info.value)
    assert not output.exists()
    assert _sha256(source) == before


@pytest.mark.parametrize(
    "price",
    [0, Decimal("0"), "0", -1, Decimal("-5"), "not-a-price", None, "NaN", "Infinity"],
)
def test_invalid_prices_fail(tmp_path: Path, price: object) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    with pytest.raises(InvalidPriceError):
        update_final_sale_prices(source, output, [_barcode_update(price)])
    assert not output.exists()


def test_float_price_is_rejected(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    with pytest.raises(InvalidPriceError):
        update_final_sale_prices(source, output, [_barcode_update(4950000.0)])
    assert not output.exists()


def test_fractional_decimal_is_rejected_without_rounding(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    with pytest.raises(InvalidPriceError):
        update_final_sale_prices(
            source, output, [_barcode_update(Decimal("4950000.50"))]
        )
    assert not output.exists()


def test_decimal_whole_prices_are_written_as_integer_text(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(
        source, output, [_barcode_update(Decimal("4950000.00"))]
    )
    workbook = load_workbook(output)
    try:
        cell = workbook[PRODUCT_SHEET_NAME]["AK2"]
        assert cell.value == "4950000"
        assert isinstance(cell.value, str)
        assert cell.data_type == "s"
    finally:
        workbook.close()
    assert result.changes[0].new_final_sale_price == Decimal("4950000.00")
    assert result.changes[0].old_final_sale_price == Decimal("7500")


def test_missing_sale_price_header_does_not_guess_ak(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    workbook[PRODUCT_SHEET_NAME]["AK1"] = "ستون دیگر"
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    with pytest.raises(ProductSheetNotFoundError):
        update_final_sale_prices(source, output, [_barcode_update()])
    assert not output.exists()
    assert _sha256(source) == before
    workbook = load_workbook(source)
    try:
        assert workbook[PRODUCT_SHEET_NAME]["AK8"].value == "7500"
    finally:
        workbook.close()


def test_sale_price_header_is_used_when_column_is_not_ak(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    moved = get_column_letter(41)
    sheet[f"{moved}1"] = sheet["AK1"].value
    for row in range(2, (sheet.max_row or 1) + 1):
        sheet.cell(row, 41).value = sheet.cell(row, 37).value
    sheet["AK1"] = "ستون دیگر"
    workbook.save(source)
    workbook.close()

    output = tmp_path / "out.xlsx"
    update_final_sale_prices(source, output, [_barcode_update()])
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 2
        assert sheet["D2"].value == BARCODE_CODE
        assert sheet["AK2"].value == "7500"
        assert sheet[f"{moved}1"].value == HEADER_SALE_PRICE
        assert sheet[f"{moved}2"].value == "4950000"
    finally:
        workbook.close()


def test_current_fixture_sale_price_header_is_column_ak() -> None:
    headers = _header_map(FIXTURE_PATH, PRODUCT_SHEET_NAME)
    assert headers[HEADER_SALE_PRICE] == "AK"
    assert headers[HEADER_CODE] == "D"
    assert headers[HEADER_NAME] == "E"


def test_updated_cell_style_is_preserved(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    cell = workbook[PRODUCT_SHEET_NAME]["AK8"]
    cell.font = Font(name="Carlito", size=14, bold=True, color="FF112233")
    cell.fill = PatternFill(fill_type="solid", fgColor="FFABCDEF")
    cell.border = Border(left=Side(style="thin", color="FF010203"))
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.protection = Protection(locked=False)
    cell.number_format = "#,##0"
    workbook.save(source)
    workbook.close()

    output = tmp_path / "out.xlsx"
    update_final_sale_prices(source, output, [_barcode_update()])
    workbook = load_workbook(output)
    try:
        cell = workbook[PRODUCT_SHEET_NAME]["AK2"]
        assert cell.value == "4950000"
        assert cell.font.name == "Carlito"
        assert cell.font.size == 14
        assert cell.font.bold is True
        assert cell.font.color is not None and cell.font.color.rgb == "FF112233"
        assert cell.fill.fgColor is not None and cell.fill.fgColor.rgb == "FFABCDEF"
        assert cell.border.left.style == "thin"
        assert cell.alignment.horizontal == "center"
        assert cell.protection.locked is False
        assert cell.number_format == "#,##0"
    finally:
        workbook.close()


def test_unselected_rows_and_stale_ranges_are_removed(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    sheet["AG3"] = "=AH3"
    sheet.merge_cells("A2:A3")
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = "A1:AN8"
    sheet.row_dimensions[8].hidden = True
    sheet.column_dimensions["B"].hidden = True
    note = workbook.create_sheet("یادداشت")
    note.sheet_state = "hidden"
    note["A1"] = "=1+1"
    workbook.defined_names.add(
        DefinedName(name="SaleAnchor", attr_text="'لیست کالا'!$D$2:$AN$8")
    )
    validation = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    validation.add("C2:C8")
    sheet.add_data_validation(validation)
    workbook.save(source)
    workbook.close()

    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(
        source,
        output,
        [_price(3, "Test Fitting Zero", Decimal("920000"), "001234")],
    )
    assert result.exported_row_count == 1
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 2
        assert sheet["D2"].value == "001234"
        assert sheet["AG2"].value == "=AH2"
        assert sheet["AK2"].value == "920000"
        assert "A2:A3" not in {str(rng) for rng in sheet.merged_cells.ranges}
        assert sheet.freeze_panes == "D2"
        assert sheet.auto_filter.ref is None
        assert sheet.row_dimensions[2].hidden is not True
        assert sheet.column_dimensions["B"].hidden is True
        assert workbook.sheetnames == [PRODUCT_SHEET_NAME, "یادداشت"]
        assert workbook["یادداشت"].sheet_state == "hidden"
        assert workbook["یادداشت"]["A1"].value == "=1+1"
        assert "SaleAnchor" not in workbook.defined_names
        assert list(sheet.tables) == []
        assert len(sheet.data_validations.dataValidation) == 0
        codes = [sheet.cell(row, 4).value for row in range(2, (sheet.max_row or 1) + 1)]
        assert BARCODE_CODE not in codes
    finally:
        workbook.close()


def test_retained_formula_merge_and_height_move_with_the_row(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    name = sheet["E4"].value
    code = sheet["D4"].value
    sheet["G4"] = "=F4*2"
    sheet.merge_cells("A4:B4")
    sheet.row_dimensions[4].height = 44
    workbook.save(source)
    workbook.close()

    output = tmp_path / "out.xlsx"
    update_final_sale_prices(
        source,
        output,
        [_price(4, name, Decimal("880000"), code)],
    )
    workbook = load_workbook(output)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row == 2
        assert sheet["D2"].value == code
        assert sheet["E2"].value == name
        assert sheet["G2"].value == "=F2*2"
        assert sheet.row_dimensions[2].height == 44
        assert "A2:B2" in {str(rng) for rng in sheet.merged_cells.ranges}
    finally:
        workbook.close()


def test_formula_that_depends_on_a_deleted_row_is_rejected(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    name = sheet["E2"].value
    code = sheet["D2"].value
    sheet["G2"] = "=F8*2"
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    with pytest.raises(PioneersExcelWriteError, match="not in the export"):
        update_final_sale_prices(
            source,
            output,
            [_price(2, name, Decimal("880000"), code)],
        )
    assert not output.exists()
    assert _sha256(source) == before


def test_chart_or_pivot_cache_rejects_the_workbook(tmp_path: Path) -> None:
    from openpyxl.chart import BarChart, Reference

    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=37, min_row=1, max_row=3), titles_from_data=True)
    sheet.add_chart(chart, "A12")
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    with pytest.raises(PioneersExcelWriteError, match="chart"):
        update_final_sale_prices(source, output, [_barcode_update()])
    assert not output.exists()
    assert _sha256(source) == before


def test_formula_in_target_sale_price_cell_is_rejected(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    workbook[PRODUCT_SHEET_NAME]["AK8"] = "=AH8"
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    with pytest.raises(UnsafePriceCellError):
        update_final_sale_prices(source, output, [_barcode_update()])
    assert not output.exists()
    assert _sha256(source) == before


def test_all_or_nothing_when_one_identity_is_wrong(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    with pytest.raises(ProductIdentityMismatchError):
        update_final_sale_prices(
            source,
            output,
            [
                _barcode_update(),
                _price(2, "نام اشتباه", Decimal("1000"), "90001"),
            ],
        )
    assert not output.exists()
    assert _sha256(source) == before
    assert list(tmp_path.glob(".gheymatyar-sale-price-*")) == []


def test_successful_result_object(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(
        source,
        output,
        [
            FinalSalePriceUpdate(
                row_number=3,
                pioneers_name="Test Fitting Zero",
                pioneers_code="001234",
                final_sale_price=Decimal("920000"),
            ),
            FinalSalePriceUpdate(
                row_number=8,
                pioneers_name=BARCODE_NAME,
                final_sale_price=Decimal("4950000"),
            ),
        ],
    )
    assert result.updated_count == 2
    assert Path(result.source_path) == source.resolve()
    assert Path(result.output_path) == output.resolve()
    assert result.changes[0].row_number == 3
    assert result.changes[0].pioneers_name == "Test Fitting Zero"
    assert result.changes[0].pioneers_code == "001234"
    assert result.changes[0].old_final_sale_price == Decimal("30000")
    assert result.changes[0].new_final_sale_price == Decimal("920000")
    assert result.changes[1].row_number == 8
    assert result.changes[1].pioneers_name == BARCODE_NAME
    assert result.changes[1].pioneers_code is None
    assert result.changes[1].old_final_sale_price == Decimal("7500")
    assert result.changes[1].new_final_sale_price == Decimal("4950000")


def _approved_valve(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        return _price(2, sheet["E2"].value, Decimal("880000"), str(sheet["D2"].value))
    finally:
        workbook.close()


def _add_hidden_catalog(path: Path, *, header_row: int, title: str) -> str:
    """Copy every product onto a veryHidden sheet whose headers are not on row 1."""
    workbook = load_workbook(path)
    source = workbook[PRODUCT_SHEET_NAME]
    hidden = workbook.create_sheet(title)
    hidden.sheet_state = "veryHidden"
    for row in range(1, header_row):
        hidden.cell(row, 1).value = f"یادداشت {row}"
    hidden.cell(header_row, 1).value = HEADER_CODE
    hidden.cell(header_row, 2).value = HEADER_NAME
    hidden.cell(header_row, 3).value = HEADER_SALE_PRICE
    for offset, src_row in enumerate(range(2, (source.max_row or 1) + 1), start=1):
        dest = header_row + offset
        hidden.cell(dest, 1).value = source.cell(src_row, 4).value
        hidden.cell(dest, 2).value = source.cell(src_row, 5).value
        hidden.cell(dest, 3).value = source.cell(src_row, 37).value
    excluded = str(source.cell(3, 4).value)
    workbook.save(path)
    workbook.close()
    return excluded


def _xml_product_codes(path: Path) -> list[str]:
    """Product codes from worksheet XML identity columns, not every cell value."""
    with zipfile.ZipFile(path) as archive:
        assert not any(
            name.startswith(("xl/pivot", "xl/charts/", "xl/tables/"))
            for name in archive.namelist()
        )
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for element in root:
                if not element.tag.endswith("si"):
                    continue
                shared.append(
                    "".join(
                        text.text or ""
                        for text in element.iter()
                        if text.tag.endswith("t")
                    )
                )
        codes: list[str] = []
        for name in archive.namelist():
            if "/worksheets/" not in name or not name.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(name))
            rows: dict[int, dict[str, str]] = {}
            for element in root.iter():
                if not element.tag.endswith("c"):
                    continue
                ref = element.get("r") or ""
                column = "".join(char for char in ref if char.isalpha())
                row_text = "".join(char for char in ref if char.isdigit())
                if not column or not row_text:
                    continue
                kind = element.get("t")
                if kind == "inlineStr":
                    value = "".join(
                        text.text or ""
                        for text in element.iter()
                        if text.tag.endswith("t")
                    )
                else:
                    raw = next(
                        (
                            child.text or ""
                            for child in element
                            if child.tag.endswith("v")
                        ),
                        "",
                    )
                    value = shared[int(raw)] if kind == "s" and raw else raw
                rows.setdefault(int(row_text), {})[column] = value
            header_row = None
            code_column = None
            for row_number, values in rows.items():
                positions = {text.strip(): column for column, text in values.items()}
                if HEADER_CODE in positions and HEADER_NAME in positions:
                    header_row = row_number
                    code_column = positions[HEADER_CODE]
                    break
            if header_row is None or code_column is None:
                continue
            for row_number in sorted(row for row in rows if row > header_row):
                code = rows[row_number].get(code_column, "").strip()
                if code:
                    codes.append(code)
        return codes


def test_hidden_catalog_copy_does_not_double_one_approval(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    excluded = _add_hidden_catalog(source, header_row=2, title="کاتالوگ تکراری")
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(source, output, [_approved_valve(source)])
    assert result.exported_row_count == 1
    saved = load_workbook(output)
    try:
        assert saved[PRODUCT_SHEET_NAME].max_row == 2
        hidden = saved["کاتالوگ تکراری"]
        assert hidden.sheet_state == "veryHidden"
        assert hidden["A1"].value == "یادداشت 1"
        assert hidden.max_row == 2
        assert hidden["A3"].value is None
    finally:
        saved.close()
    assert _xml_product_codes(output) == ["90001"]
    assert excluded not in _xml_product_codes(output)
    assert _sha256(source) == before


def test_catalog_headers_past_row_15_do_not_keep_unapproved_products(
    tmp_path: Path,
) -> None:
    source = _copy_fixture(tmp_path)
    excluded = _add_hidden_catalog(source, header_row=16, title="کاتالوگ ردیف ۱۶")
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(source, output, [_approved_valve(source)])
    assert result.exported_row_count == 1
    saved = load_workbook(output)
    try:
        hidden = saved["کاتالوگ ردیف ۱۶"]
        assert hidden.sheet_state == "veryHidden"
        assert hidden["A16"].value == HEADER_CODE
        assert hidden["A1"].value == "یادداشت 1"
        assert hidden.max_row == 16
        assert hidden["A17"].value is None
    finally:
        saved.close()
    codes = _xml_product_codes(output)
    assert codes == ["90001"]
    assert excluded not in codes
    with zipfile.ZipFile(output) as archive:
        payload = "\n".join(
            archive.read(name).decode("utf-8", errors="replace")
            for name in archive.namelist()
            if name.endswith(".xml") and "/worksheets/" in name
        )
    assert f">{excluded}<" not in payload
    assert _sha256(source) == before


def test_sale_price_may_equal_an_excluded_product_code(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    sheet["D3"].value = "123456"
    sheet["E3"].value = "Excluded Product 123456"
    approved_name = sheet["E2"].value
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(
        source,
        output,
        [_price(2, approved_name, Decimal("123456"), "90001")],
    )
    assert result.exported_row_count == 1
    saved = load_workbook(output)
    try:
        exported = saved[PRODUCT_SHEET_NAME]
        assert exported.max_row == 2
        assert exported["D2"].value == "90001"
        assert exported["AK2"].value == "123456"
        assert exported["E2"].value == approved_name
    finally:
        saved.close()
    assert _xml_product_codes(output) == ["90001"]
    with zipfile.ZipFile(output) as archive:
        payload = "\n".join(
            archive.read(name).decode("utf-8", errors="replace")
            for name in archive.namelist()
            if name.endswith(".xml")
        )
    assert ">123456<" in payload
    assert "Excluded Product 123456" not in payload
    assert _sha256(source) == before


def test_model_rejects_float_money() -> None:
    with pytest.raises(ValidationError):
        FinalSalePriceUpdate(
            row_number=8,
            pioneers_name=BARCODE_NAME,
            final_sale_price=1.5,
        )


def test_save_failure_leaves_no_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _copy_fixture(tmp_path)
    output = tmp_path / "out.xlsx"
    before = _sha256(source)

    def _boom(self, filename: object) -> None:  # noqa: ANN001
        raise OSError("disk full")

    monkeypatch.setattr(
        "gheymatyar.services.pioneers_excel_writer.Workbook.save",
        _boom,
    )
    with pytest.raises(OSError, match="disk full"):
        update_final_sale_prices(source, output, [_barcode_update()])
    assert not output.exists()
    assert _sha256(source) == before
    assert list(tmp_path.glob(".gheymatyar-sale-price-*")) == []


def test_partial_header_hidden_catalog_does_not_keep_unapproved_products(
    tmp_path: Path,
) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    hidden = workbook.create_sheet("هویت ناقص")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = HEADER_CODE
    hidden["B1"] = HEADER_NAME
    hidden["C1"] = "قیمت"
    hidden["A2"] = "ARCH-0001"
    hidden["B2"] = "Arch product"
    hidden["C2"] = "10"
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(source, output, [_approved_valve(source)])
    assert result.exported_row_count == 1
    saved = load_workbook(output)
    try:
        assert saved[PRODUCT_SHEET_NAME].max_row == 2
        hidden = saved["هویت ناقص"]
        assert hidden.sheet_state == "veryHidden"
        assert hidden["A1"].value == HEADER_CODE
        assert hidden["C1"].value == "قیمت"
        assert hidden.max_row == 1
        for sheet in saved.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.value != "ARCH-0001"
    finally:
        saved.close()
    assert _xml_product_codes(output) == ["90001"]
    with zipfile.ZipFile(output) as archive:
        payload = "\n".join(
            archive.read(name).decode("utf-8", errors="replace")
            for name in archive.namelist()
            if name.endswith(".xml") and "/worksheets/" in name
        )
    assert "ARCH-0001" not in payload
    assert _sha256(source) == before


def test_spaced_code_and_name_headers_are_preserved(tmp_path: Path) -> None:
    source = _copy_fixture(tmp_path)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    code_header = f" {sheet['D1'].value} "
    name_header = f" {sheet['E1'].value} "
    sheet["D1"] = code_header
    sheet["E1"] = name_header
    approved_name = sheet["E2"].value
    approved_code = str(sheet["D2"].value)
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(
        source,
        output,
        [_price(2, approved_name, Decimal("880000"), approved_code)],
    )
    assert result.exported_row_count == 1
    saved = load_workbook(output)
    try:
        exported = saved[PRODUCT_SHEET_NAME]
        assert exported["D1"].value == code_header
        assert exported["E1"].value == name_header
        assert exported["D2"].value == approved_code
        assert exported.max_row == 2
    finally:
        saved.close()
    assert _xml_product_codes(output) == [approved_code]
    assert _sha256(source) == before
