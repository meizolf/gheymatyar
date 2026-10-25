"""Tests for Pioneers HTML-based .xls detection, parsing, and conversion."""

from __future__ import annotations

import hashlib
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from gheymatyar.parsers.errors import PioneersExcelError
from gheymatyar.parsers.pioneers_excel_parser import (
    HEADER_CODE,
    HEADER_NAME,
    HEADER_SALE_PRICE,
    PioneersExcelParser,
)
from gheymatyar.parsers.pioneers_html_xls import (
    convert_pioneers_html_xls_to_xlsx,
    html_xls_preservation_notes,
    read_pioneers_html_xls_rows,
)
from gheymatyar.parsers.workbook_format import WorkbookFormat, detect_workbook_format
from gheymatyar.services.pioneers_excel_writer import update_final_sale_prices
from gheymatyar.services.processing import ProcessingService
from gheymatyar.services.session_store import SessionStore

REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURES = REPO_ROOT / "fixtures" / "pioneers"
HTML_SAMPLE = FIXTURES / "pioneers_html_sample.xls"
HTML_AZIN = FIXTURES / "pioneers_azin_subset.xls"
XLSX_AZIN = FIXTURES / "pioneers_azin_subset.xlsx"


def _resolve_optional_path(env_var: str) -> Path | None:
    """Large fixtures are opt-in via env vars (no hard-coded Downloads paths)."""
    raw = os.environ.get(env_var, "").strip()
    if not raw:
        return None
    path = Path(raw).expanduser()
    return path if path.is_file() else None


ORIGINAL_HTML_XLS = _resolve_optional_path("GHEYMATYAR_PIONEERS_HTML_XLS")
CONVERTED_XLSX = _resolve_optional_path("GHEYMATYAR_PIONEERS_CONVERTED_XLSX")

_HEADERS_40 = [
    "انتخاب",
    "ردیف",
    "فعال",
    HEADER_CODE,
    HEADER_NAME,
    "بارکد",
    "کد رده",
    "نام رده",
    "نقطه سفارش",
    "حداکثر نقطه سفارش",
    "درصد تخفیف",
    "درصد تخفیف پنل فروش",
    "درصد تخفیف نقدی",
    "درصد مالیات",
    "درصد پورسانت",
    "تعداد آمپر",
    "نرخ آمپر",
    "مدت تسویه (روز)",
    "سایر مشخصات 2 : ",
    "سایر مشخصات 1 : ",
    "مرکز پخش",
    "مشخصات فنی : ",
    " کد تامین کننده :",
    "نام تامین کننده :",
    "مرکز پخش",
    "نام مرکز پخش",
    "نرخ جذب",
    "بهای استاندارد",
    "وزن",
    "حجم",
    "عرض",
    "مانده",
    "FormulaResult",
    "نرخ آخرین خرید",
    "نرخ بهای تمام شده",
    "نرخ خرید روز",
    HEADER_SALE_PRICE,
    "نرخ مصرف کننده",
    "نرخ همکاری",
    "نرخ لوله کش",
]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _td(value: str, *, colspan: str = "1", rowspan: str | None = None) -> str:
    attrs = [f'colspan="{colspan}"']
    if rowspan is not None:
        attrs.append(f'rowspan="{rowspan}"')
    return f"<td {' '.join(attrs)}>{value}</td>"


def _write_html_xls(path: Path, rows: list[list[str]]) -> Path:
    lines = ["<Html>", "<Body>", "<Table>"]
    for row in rows:
        lines.append("<tr>")
        for cell in row:
            lines.append(_td(cell))
        lines.append("</tr>")
    lines += ["</Table>", "</Body>", "</Html>", ""]
    path.write_text("\n".join(lines), encoding="utf-8-sig")
    return path


def _product_row(
    code: str,
    name: str,
    sale: str,
    *,
    barcode: str = "",
    extra: dict[int, str] | None = None,
) -> list[str]:
    row = [""] * 40
    row[1] = "1"
    row[2] = "True"
    row[3] = code
    row[4] = name
    row[5] = barcode
    row[36] = sale
    if extra:
        for index, value in extra.items():
            row[index] = value
    return row


def test_detect_xlsx_zip_magic() -> None:
    data = XLSX_AZIN.read_bytes()[:64]
    assert detect_workbook_format(data) == WorkbookFormat.XLSX


def test_detect_pioneers_html_xls() -> None:
    data = HTML_SAMPLE.read_bytes()
    assert detect_workbook_format(data) == WorkbookFormat.PIONEERS_HTML_XLS


def test_detect_ole_binary_xls_unsupported() -> None:
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32
    assert detect_workbook_format(ole) == WorkbookFormat.UNSUPPORTED


def test_detect_random_html_without_pioneers_headers_unsupported() -> None:
    html = b"\xef\xbb\xbf<Html><Body><Table><tr><td>hello</td></tr></Table></Body></Html>"
    assert detect_workbook_format(html) == WorkbookFormat.UNSUPPORTED


def test_read_html_sample_extracts_all_columns_and_rows() -> None:
    rows = read_pioneers_html_xls_rows(HTML_SAMPLE)
    assert len(rows[0]) == 40
    assert rows[0][3].strip() == HEADER_CODE
    assert rows[0][4].strip() == HEADER_NAME
    assert rows[0][36].strip() == HEADER_SALE_PRICE
    assert len(rows) == 7  # header + 6 products
    assert rows[1][3] == "10001"
    assert "      " in rows[1][4]
    assert rows[1][36] == "4610000"
    leading = next(r for r in rows[1:] if r[5].startswith("0"))
    assert leading[5] == "015113"
    empty_name = next(r for r in rows[1:] if r[3] == "14867")
    assert empty_name[4] == ""


def test_preserves_leading_and_trailing_cell_whitespace(tmp_path: Path) -> None:
    headers = list(_HEADERS_40)
    product = _product_row("10001", " name  ", "1000")
    product[18] = " trail "
    path = _write_html_xls(tmp_path / "spaces.xls", [headers, product])
    rows = read_pioneers_html_xls_rows(path)
    assert rows[0][18] == "سایر مشخصات 2 : "
    assert rows[0][22] == " کد تامین کننده :"
    assert rows[1][4] == " name  "
    assert rows[1][18] == " trail "

    out = tmp_path / "spaces.xlsx"
    convert_pioneers_html_xls_to_xlsx(path, out)
    wb = load_workbook(out, read_only=True, data_only=False)
    try:
        sheet = wb.active
        assert sheet["E2"].value == " name  "
        assert sheet["S1"].value == "سایر مشخصات 2 : "
        assert sheet["W1"].value == " کد تامین کننده :"
    finally:
        wb.close()


def test_rejects_short_row_instead_of_padding(tmp_path: Path) -> None:
    headers = list(_HEADERS_40)
    short = _product_row("10001", "A", "1000")[:39]
    path = _write_html_xls(tmp_path / "short.xls", [headers, short])
    with pytest.raises(PioneersExcelError, match="column count") as excinfo:
        read_pioneers_html_xls_rows(path)
    assert excinfo.value.row == 2


def test_rejects_colspan_greater_than_one(tmp_path: Path) -> None:
    headers = list(_HEADERS_40)
    product = _product_row("10001", "A", "1000")
    path = tmp_path / "colspan.xls"
    lines = ["<Html>", "<Body>", "<Table>", "<tr>"]
    for header in headers:
        lines.append(_td(header, colspan="1"))
    lines.append("</tr><tr>")
    lines.append(_td(product[0], colspan="2"))
    for cell in product[1:]:
        lines.append(_td(cell, colspan="1"))
    lines += ["</tr>", "</Table>", "</Body>", "</Html>"]
    path.write_text("\n".join(lines), encoding="utf-8-sig")
    with pytest.raises(PioneersExcelError, match="colspan"):
        read_pioneers_html_xls_rows(path)


def test_rejects_nested_table(tmp_path: Path) -> None:
    headers = list(_HEADERS_40)
    product = _product_row("10001", "A", "1000")
    path = tmp_path / "nested.xls"
    inner = "<Table><tr><td>x</td></tr></Table>"
    product = list(product)
    product[4] = f"outer{inner}"
    # Put a real nested table element between cells.
    lines = ["<Html>", "<Body>", "<Table>", "<tr>"]
    for header in headers:
        lines.append(_td(header))
    lines.append("</tr><tr>")
    for index, cell in enumerate(product):
        if index == 4:
            lines.append(f"<td colspan=\"1\">name{inner}</td>")
        else:
            lines.append(_td(cell))
    lines += ["</tr>", "</Table>", "</Body>", "</Html>"]
    path.write_text("\n".join(lines), encoding="utf-8-sig")
    with pytest.raises(PioneersExcelError, match="nested"):
        read_pioneers_html_xls_rows(path)


def test_writes_formula_like_and_error_literals_as_text(tmp_path: Path) -> None:
    headers = list(_HEADERS_40)
    product = _product_row("10001", "=1+1", "1000", barcode="#N/A")
    path = _write_html_xls(tmp_path / "lit.xls", [headers, product])
    out = tmp_path / "lit.xlsx"
    convert_pioneers_html_xls_to_xlsx(path, out)
    wb = load_workbook(out, data_only=False)
    try:
        sheet = wb.active
        assert sheet["E2"].value == "=1+1"
        assert sheet["E2"].data_type == "s"
        assert sheet["F2"].value == "#N/A"
        assert sheet["F2"].data_type == "s"
    finally:
        wb.close()


def test_rejects_overlong_cell_with_row_and_column(tmp_path: Path) -> None:
    headers = list(_HEADERS_40)
    product = _product_row("10001", "A" * 33000, "1000")
    path = _write_html_xls(tmp_path / "long.xls", [headers, product])
    with pytest.raises(PioneersExcelError, match="32767") as excinfo:
        read_pioneers_html_xls_rows(path)
    assert excinfo.value.row == 2
    assert excinfo.value.column == "E"


def test_convert_html_sample_round_trips_through_parser(tmp_path: Path) -> None:
    out = tmp_path / "from_html.xlsx"
    convert_pioneers_html_xls_to_xlsx(HTML_SAMPLE, out)
    assert out.is_file()

    wb = load_workbook(out, read_only=True, data_only=False)
    try:
        assert wb.sheetnames == ["لیست کالا"]
        sheet = wb["لیست کالا"]
        assert sheet["D1"].value == "کد کالا"
        assert sheet["E2"].value == "پارس پلاست دریچه سیخ زن      160"
        assert sheet["AK2"].value == "4610000"
        assert sheet["F5"].value == "015113"
    finally:
        wb.close()

    result = PioneersExcelParser().parse(out)
    codes = {p.code for p in result.products if p.code}
    assert "10001" in codes
    assert "015113" in {p.barcode for p in result.products if p.barcode}
    assert any(row.code == "14867" for row in result.invalid_rows)


def test_convert_refuses_to_overwrite_source(tmp_path: Path) -> None:
    target = tmp_path / "same.xls"
    target.write_bytes(HTML_SAMPLE.read_bytes())
    with pytest.raises(PioneersExcelError, match="overwrite"):
        convert_pioneers_html_xls_to_xlsx(target, target)


def test_azin_html_matches_azin_xlsx_values(tmp_path: Path) -> None:
    html_rows = read_pioneers_html_xls_rows(HTML_AZIN)
    converted = tmp_path / "azin.xlsx"
    convert_pioneers_html_xls_to_xlsx(HTML_AZIN, converted)

    wb = load_workbook(XLSX_AZIN, read_only=True, data_only=True)
    try:
        xlsx_rows = [
            [("" if v is None else str(v)) for v in row]
            for row in wb.active.iter_rows(values_only=True)
        ]
    finally:
        wb.close()

    assert html_rows[0] == xlsx_rows[0]
    assert len(html_rows) == len(xlsx_rows)
    for html_row, xlsx_row in zip(html_rows, xlsx_rows, strict=True):
        assert html_row == xlsx_row

    result_html = PioneersExcelParser().parse(converted)
    result_xlsx = PioneersExcelParser().parse(XLSX_AZIN)
    assert len(result_html.products) == len(result_xlsx.products)
    by_code_h = {p.code: p for p in result_html.products if p.code}
    by_code_x = {p.code: p for p in result_xlsx.products if p.code}
    assert by_code_h.keys() == by_code_x.keys()
    for code, product in by_code_h.items():
        other = by_code_x[code]
        assert product.name == other.name
        assert product.sale_price == other.sale_price
        assert product.barcode == other.barcode


def test_upload_html_xls_via_processing_service() -> None:
    store = SessionStore()
    svc = ProcessingService(store)
    session = svc.create_session()
    before = _sha256(HTML_SAMPLE)
    with HTML_SAMPLE.open("rb") as handle:
        svc.upload_pioneers_excel(session.session_id, "گزارش.xls", handle)
    after = _sha256(HTML_SAMPLE)
    assert before == after

    session = svc.get_session(session.session_id)
    assert session.pioneers_path is not None
    assert session.pioneers_path.suffix == ".xlsx"
    assert session.pioneers_path.is_file()
    assert session.pioneers_filename.endswith(".xls")
    assert len(session.products) >= 5


def _sheet_matrix(path: Path) -> list[list[object | None]]:
    wb = load_workbook(path, data_only=False)
    try:
        sheet = wb.active
        return [
            [cell.value for cell in row]
            for row in sheet.iter_rows(
                min_row=1,
                max_row=sheet.max_row,
                max_col=sheet.max_column,
            )
        ]
    finally:
        wb.close()


def test_full_grid_only_selected_sale_prices_change(tmp_path: Path) -> None:
    source = tmp_path / "pioneers.xlsx"
    convert_pioneers_html_xls_to_xlsx(HTML_SAMPLE, source)
    before_hash = _sha256(source)
    before = _sheet_matrix(source)
    output = tmp_path / "out.xlsx"

    updates = [
        {
            "row_number": 2,
            "pioneers_name": "پارس پلاست دریچه سیخ زن      160",
            "pioneers_code": "10001",
            "final_sale_price": "9999999",
        },
        {
            "row_number": 3,
            "pioneers_name": before[2][4],
            "pioneers_code": before[2][3],
            "final_sale_price": "8888888",
        },
    ]
    result = update_final_sale_prices(source, output, updates)
    assert result.updated_count == 2
    assert _sha256(source) == before_hash

    after = _sheet_matrix(output)
    assert len(before) > 3
    assert len(after) == 3
    sale_col = 36  # 0-based AK
    assert after[0] == before[0]
    for out_index, src_index in ((1, 1), (2, 2)):
        brow, arow = before[src_index], after[out_index]
        assert len(arow) == len(brow)
        for c_index, (bval, aval) in enumerate(zip(brow, arow, strict=True)):
            if c_index == sale_col:
                assert aval != bval
                continue
            assert aval == bval, (
                f"unexpected change at "
                f"{get_column_letter(c_index + 1)} source-row {src_index + 1}: "
                f"{bval!r} -> {aval!r}"
            )
    assert after[1][sale_col] == "9999999"
    assert after[2][sale_col] == "8888888"


def test_html_azin_xls_plus_pdf_export_workflow(client, tmp_path: Path) -> None:
    """Complete legacy HTML .xls + supplier PDF → updated .xlsx path."""
    pdf = REPO_ROOT / "fixtures" / "suppliers" / "458_124550609784457.pdf"
    if not pdf.is_file():
        pytest.skip("azin PDF fixture missing")

    before = _sha256(HTML_AZIN)
    html_rows = read_pioneers_html_xls_rows(HTML_AZIN)
    source_xlsx = tmp_path / "source_from_html.xlsx"
    convert_pioneers_html_xls_to_xlsx(HTML_AZIN, source_xlsx)
    source_matrix = _sheet_matrix(source_xlsx)

    sid = client.post("/api/sessions").json()["session_id"]
    with HTML_AZIN.open("rb") as handle:
        excel = client.post(
            f"/api/sessions/{sid}/pioneers-excel",
            files={"file": ("azin.xls", handle, "application/vnd.ms-excel")},
        )
    assert excel.status_code == 200, excel.text

    with pdf.open("rb") as handle:
        pdf_resp = client.post(
            f"/api/sessions/{sid}/supplier-pdf",
            data={"supplier_key": "azin_pipe"},
            files={"file": ("azin.pdf", handle, "application/pdf")},
        )
    assert pdf_resp.status_code == 200, pdf_resp.text

    match = client.post(f"/api/sessions/{sid}/match")
    assert match.status_code == 200, match.text
    preview = client.get(f"/api/sessions/{sid}/preview").json()["candidates"]
    keep = next(c for c in preview if c["status"] == "MATCHED")
    assert keep["pricing_blockers"] == []
    if keep["warnings"] or not keep["exportable"]:
        approved = client.post(
            f"/api/sessions/{sid}/candidates/{keep['candidate_id']}/review",
            json={"decision": "APPROVE"},
        )
        assert approved.status_code == 200, approved.text
    for row in preview:
        if row["candidate_id"] == keep["candidate_id"]:
            continue
        excluded = client.post(
            f"/api/sessions/{sid}/candidates/{row['candidate_id']}/review",
            json={"decision": "EXCLUDE"},
        )
        assert excluded.status_code == 200, excluded.text
    export = client.post(f"/api/sessions/{sid}/export")
    assert export.status_code == 200, export.text
    assert export.json()["updated_count"] == 1
    assert export.json()["exported_row_count"] == 1
    download = client.get(f"/api/sessions/{sid}/download?format=xlsx")
    assert download.status_code == 200
    exported_path = tmp_path / "azin_out.xlsx"
    exported_path.write_bytes(download.content)
    assert len(_sheet_matrix(exported_path)) == 2
    assert _sha256(HTML_AZIN) == before
    assert source_matrix
    assert html_rows[0][36].strip() == HEADER_SALE_PRICE


def test_api_accepts_html_xls_upload(client) -> None:
    sid = client.post("/api/sessions").json()["session_id"]
    with HTML_SAMPLE.open("rb") as handle:
        resp = client.post(
            f"/api/sessions/{sid}/pioneers-excel",
            files={
                "file": (
                    "pioneers.xls",
                    handle,
                    "application/vnd.ms-excel",
                )
            },
        )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["product_count"] >= 5
    assert body["pioneers_filename"].endswith(".xls")


def test_api_rejects_binary_ole_xls(client) -> None:
    sid = client.post("/api/sessions").json()["session_id"]
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
    resp = client.post(
        f"/api/sessions/{sid}/pioneers-excel",
        files={"file": ("legacy.xls", ole, "application/vnd.ms-excel")},
    )
    assert resp.status_code == 400
    assert "binary" in resp.json()["detail"].lower()


def test_api_rejects_short_row_html_xls(client, tmp_path: Path) -> None:
    headers = list(_HEADERS_40)
    short = _product_row("10001", "A", "1000")[:39]
    path = _write_html_xls(tmp_path / "short.xls", [headers, short])
    sid = client.post("/api/sessions").json()["session_id"]
    with path.open("rb") as handle:
        resp = client.post(
            f"/api/sessions/{sid}/pioneers-excel",
            files={"file": ("short.xls", handle, "application/vnd.ms-excel")},
        )
    assert resp.status_code == 400
    assert "column count" in resp.json()["detail"].lower()


def test_preservation_notes_mention_styles() -> None:
    notes = " ".join(html_xls_preservation_notes())
    assert "font" in notes.lower() or "Yekan" in notes or "styles" in notes.lower()


@pytest.mark.skipif(
    ORIGINAL_HTML_XLS is None,
    reason="set GHEYMATYAR_PIONEERS_HTML_XLS to the original HTML .xls path",
)
def test_original_full_html_xls_parses_all_rows(tmp_path: Path) -> None:
    assert ORIGINAL_HTML_XLS is not None
    before = _sha256(ORIGINAL_HTML_XLS)
    rows = read_pioneers_html_xls_rows(ORIGINAL_HTML_XLS)
    assert len(rows[0]) == 40
    assert len(rows) >= 6700
    assert all(len(row) == 40 for row in rows)
    # Exact whitespace from source must survive (regression for strip()).
    assert any(cell != cell.strip() for cell in rows[0])
    assert any(row[4] != row[4].strip() for row in rows[1:] if row[4])

    out = tmp_path / "full.xlsx"
    convert_pioneers_html_xls_to_xlsx(ORIGINAL_HTML_XLS, out)
    result = PioneersExcelParser().parse(out)
    assert len(result.products) + len(result.invalid_rows) == len(rows) - 1
    assert any(row.code == "14867" for row in result.invalid_rows)

    if CONVERTED_XLSX is not None:
        converted = PioneersExcelParser().parse(CONVERTED_XLSX)
        by_new = {p.code: p for p in result.products if p.code}
        by_old = {p.code: p for p in converted.products if p.code}
        common = set(by_new) & set(by_old)
        assert len(common) >= 6000
        for code in ("10001",):
            if code in by_new and code in by_old:
                assert by_new[code].sale_price == by_old[code].sale_price

    assert _sha256(ORIGINAL_HTML_XLS) == before


@pytest.mark.skipif(
    ORIGINAL_HTML_XLS is None,
    reason="set GHEYMATYAR_PIONEERS_HTML_XLS to the original HTML .xls path",
)
def test_original_html_xls_plus_azin_pdf_export_workflow(client, tmp_path: Path) -> None:
    """Legacy HTML .xls upload → match/export path using the azin PDF fixture."""
    assert ORIGINAL_HTML_XLS is not None
    pdf = REPO_ROOT / "fixtures" / "suppliers" / "458_124550609784457.pdf"
    if not pdf.is_file():
        pytest.skip("azin PDF fixture missing")

    before = _sha256(ORIGINAL_HTML_XLS)
    sid = client.post("/api/sessions").json()["session_id"]
    with ORIGINAL_HTML_XLS.open("rb") as handle:
        excel = client.post(
            f"/api/sessions/{sid}/pioneers-excel",
            files={"file": ("full.xls", handle, "application/vnd.ms-excel")},
        )
    assert excel.status_code == 200, excel.text
    assert excel.json()["product_count"] >= 6700

    with pdf.open("rb") as handle:
        pdf_resp = client.post(
            f"/api/sessions/{sid}/supplier-pdf",
            data={"supplier_key": "azin_pipe"},
            files={"file": ("azin.pdf", handle, "application/pdf")},
        )
    assert pdf_resp.status_code == 200, pdf_resp.text

    match = client.post(f"/api/sessions/{sid}/match")
    assert match.status_code == 200, match.text
    preview = client.get(f"/api/sessions/{sid}/preview").json()["candidates"]
    matched = [c for c in preview if c["status"] == "MATCHED"]
    assert matched, "expected at least one MATCHED candidate against full catalog"
    assert all(not row["pricing_blockers"] for row in matched)
    keep = matched[0]
    if keep["warnings"] or not keep["exportable"]:
        approved = client.post(
            f"/api/sessions/{sid}/candidates/{keep['candidate_id']}/review",
            json={"decision": "APPROVE"},
        )
        assert approved.status_code == 200, approved.text
    for row in preview:
        if row["candidate_id"] == keep["candidate_id"]:
            continue
        excluded = client.post(
            f"/api/sessions/{sid}/candidates/{row['candidate_id']}/review",
            json={"decision": "EXCLUDE"},
        )
        assert excluded.status_code == 200, excluded.text
    export = client.post(f"/api/sessions/{sid}/export")
    assert export.status_code == 200, export.text
    assert export.json()["updated_count"] == 1
    assert _sha256(ORIGINAL_HTML_XLS) == before
