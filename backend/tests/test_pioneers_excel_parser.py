"""Tests for the read-only Pioneers Excel parser."""

from __future__ import annotations

import hashlib
import shutil
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from gheymatyar.parsers import (
    AmbiguousMappedHeaderError,
    InvalidProductRow,
    PioneersExcelError,
    PioneersExcelParser,
    SheetNotFoundError,
)
from gheymatyar.parsers.pioneers_excel_parser import (
    HEADER_BARCODE,
    HEADER_CODE,
    HEADER_NAME,
    HEADER_SALE_PRICE,
    PRODUCT_SHEET_NAME,
)

REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURE_PATH = REPO_ROOT / "fixtures" / "pioneers" / "pioneers_sample.xlsx"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_workbook(path: Path, headers: list[str], rows: list[list[object]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = PRODUCT_SHEET_NAME
    for col, header in enumerate(headers, start=1):
        ws.cell(1, col, header)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(r_idx, c_idx, value)
    wb.save(path)
    return path


@pytest.fixture()
def sample_path() -> Path:
    assert FIXTURE_PATH.is_file(), f"missing fixture: {FIXTURE_PATH}"
    return FIXTURE_PATH


def test_valid_workbook_parses_correctly(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    assert len(result.products) == 6
    first = result.products[0]
    assert first.row_number == 2
    assert first.code == "90001"
    assert first.name == "Test Valve 20mm"
    assert first.barcode is None
    assert first.unit is None
    assert first.sale_price == Decimal("150000")
    assert first.last_purchase_price == Decimal("100000")
    assert first.consumer_price == Decimal("160000")
    assert first.colleague_price == Decimal("0")


def test_expected_product_count(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    assert len(result) == 6
    assert len(result.products) == 6


def test_excel_row_numbers_preserved(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    assert [p.row_number for p in result.products] == [2, 3, 4, 6, 7, 8]


def test_product_codes_and_names(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    assert [p.code for p in result.products] == [
        "90001",
        "001234",
        "90002",
        "90003",
        "90001",
        "90004",
    ]
    assert result.products[1].name == "Test Fitting Zero"


def test_optional_fields_can_be_none(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    empty_prices = result.products[2]
    assert empty_prices.code == "90002"
    assert empty_prices.barcode is None
    assert empty_prices.unit is None
    assert empty_prices.sale_price is None
    assert empty_prices.last_purchase_price is None
    assert empty_prices.consumer_price is None
    assert empty_prices.colleague_price is None


def test_decimal_price_precision(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    precise = next(p for p in result.products if p.code == "90003")
    assert precise.sale_price == Decimal("0.1")
    assert precise.last_purchase_price == Decimal("10.1")
    # Must not become a binary-float artifact such as 0.10000000000000001
    assert str(precise.sale_price) == "0.1"
    assert precise.sale_price * 10 == Decimal("1.0")


def test_raw_values_preserved(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    first = result.products[0]
    assert first.raw_values["D"] == "90001"
    assert first.raw_values["E"] == "Test Valve 20mm"
    assert first.raw_values["AK"] == "150000"
    assert "A" in first.raw_values


def test_missing_required_header_produces_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "missing_header.xlsx"
    _write_workbook(
        path,
        headers=["کد کالا", "نام دیگر", "نرخ فروش اصلی"],
        rows=[["90001", "X", "100"]],
    )
    with pytest.raises(SheetNotFoundError) as exc_info:
        PioneersExcelParser().parse(path)
    message = str(exc_info.value)
    assert "No sheet matches" in message
    assert HEADER_NAME in message


def test_missing_product_code_collected_for_review(tmp_path: Path) -> None:
    path = tmp_path / "missing_code.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_CODE, HEADER_NAME, HEADER_SALE_PRICE],
        rows=[
            [None, "Has Name Only", "100"],
            ["91002", "Valid After Invalid", "200"],
        ],
    )
    result = PioneersExcelParser().parse(path)
    assert len(result.products) == 1
    assert result.products[0].code == "91002"
    assert len(result.invalid_rows) == 1
    invalid = result.invalid_rows[0]
    assert isinstance(invalid, InvalidProductRow)
    assert invalid.row_number == 2
    assert invalid.column == HEADER_CODE
    assert invalid.name == "Has Name Only"
    assert "کد کالا" in invalid.reason


def test_missing_product_name_collected_for_review(tmp_path: Path) -> None:
    path = tmp_path / "missing_name.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_CODE, HEADER_NAME, HEADER_SALE_PRICE],
        rows=[
            ["90099", None, "100"],
            ["90100", "Valid Neighbor", "150"],
        ],
    )
    result = PioneersExcelParser().parse(path)
    assert [p.code for p in result.products] == ["90100"]
    assert len(result.invalid_rows) == 1
    invalid = result.invalid_rows[0]
    assert invalid.row_number == 2
    assert invalid.column == HEADER_NAME
    assert invalid.code == "90099"
    assert "نـــام کــالا" in invalid.reason


def test_duplicate_product_code_not_silently_hidden(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    assert "90001" in result.duplicate_codes
    assert result.duplicate_codes["90001"] == (2, 7)
    # Both rows are still returned
    assert sum(1 for p in result.products if p.code == "90001") == 2


def test_leading_zero_product_code_preserved(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    zero_code = next(p for p in result.products if p.row_number == 3)
    assert zero_code.code == "001234"
    assert zero_code.code != "1234"


def test_leading_zero_barcode_preserved(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    zero_barcode = next(p for p in result.products if p.code == "90004")
    assert zero_barcode.barcode == "015113"
    fitting = next(p for p in result.products if p.code == "001234")
    assert fitting.barcode == "001122334455"


def test_blank_rows_skipped(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    # Fixture row 5 is blank; products jump from Excel row 4 to 6
    assert 5 not in {p.row_number for p in result.products}
    assert 5 not in {row.row_number for row in result.invalid_rows}


def test_known_non_product_footer_rows_skipped(tmp_path: Path) -> None:
    """Entirely blank trailing rows are ignored."""
    path = tmp_path / "with_blank_footer.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_CODE, HEADER_NAME, HEADER_SALE_PRICE],
        rows=[
            ["91001", "Footer Test Product", "500"],
            [None, None, None],
            ["", "", ""],
        ],
    )
    result = PioneersExcelParser().parse(path)
    assert len(result.products) == 1
    assert result.products[0].code == "91001"
    assert result.products[0].row_number == 2
    assert result.invalid_rows == []


def test_nonempty_row_without_code_and_name_is_invalid(tmp_path: Path) -> None:
    """Point 2: do not silently skip rows that have data but no identity."""
    path = tmp_path / "orphan_data.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_CODE, HEADER_NAME, HEADER_SALE_PRICE],
        rows=[
            ["93001", "Valid Product", "100"],
            [None, None, "999"],  # price present, no code/name
            ["93002", "Another Valid", "200"],
        ],
    )
    result = PioneersExcelParser().parse(path)
    assert [p.code for p in result.products] == ["93001", "93002"]
    assert len(result.invalid_rows) == 1
    invalid = result.invalid_rows[0]
    assert invalid.row_number == 3
    assert "both" in invalid.reason
    assert invalid.raw_values.get("C") == "999"


def test_numeric_code_with_zero_pad_format_preserves_leading_zeros(
    tmp_path: Path,
) -> None:
    path = tmp_path / "zero_pad_code.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = PRODUCT_SHEET_NAME
    ws["A1"] = HEADER_CODE
    ws["B1"] = HEADER_NAME
    ws["A2"] = 1234
    ws["A2"].number_format = "000000"
    ws["B2"] = "Padded Code Product"
    wb.save(path)

    result = PioneersExcelParser().parse(path)
    assert len(result.products) == 1
    assert result.products[0].code == "001234"
    assert result.invalid_rows == []


def test_multi_section_zero_pad_format_is_rejected(tmp_path: Path) -> None:
    """Multi-section formats must not be treated as plain zero-pad."""
    path = tmp_path / "multi_section_format.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = PRODUCT_SHEET_NAME
    ws["A1"] = HEADER_CODE
    ws["B1"] = HEADER_NAME
    ws["A2"] = 0
    ws["A2"].number_format = '000000;000000;"ZERO"'
    ws["B2"] = "Multi Section Format Product"
    wb.save(path)

    result = PioneersExcelParser().parse(path)
    assert result.products == []
    assert len(result.invalid_rows) == 1
    assert result.invalid_rows[0].column == HEADER_CODE
    # Must not silently become the first-section string "000000".
    assert result.invalid_rows[0].raw_values.get("A") == 0


def test_unsupported_numeric_code_format_is_invalid(tmp_path: Path) -> None:
    path = tmp_path / "bare_numeric_code.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = PRODUCT_SHEET_NAME
    ws["A1"] = HEADER_CODE
    ws["B1"] = HEADER_NAME
    ws["A2"] = 1234
    ws["A2"].number_format = "General"
    ws["B2"] = "Bare Numeric Code"
    wb.save(path)

    result = PioneersExcelParser().parse(path)
    assert result.products == []
    assert len(result.invalid_rows) == 1
    assert result.invalid_rows[0].column == HEADER_CODE
    assert "numeric" in result.invalid_rows[0].reason.lower()


def test_formula_code_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "formula_code.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = PRODUCT_SHEET_NAME
    ws["A1"] = HEADER_CODE
    ws["B1"] = HEADER_NAME
    ws["A2"] = "=A1"
    ws["B2"] = "Formula Code Product"
    wb.save(path)

    result = PioneersExcelParser().parse(path)
    assert result.products == []
    assert len(result.invalid_rows) == 1
    assert "formula" in result.invalid_rows[0].reason.lower()


def test_excel_error_barcode_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "error_barcode.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_CODE, HEADER_NAME, HEADER_BARCODE],
        rows=[["94001", "Error Barcode Product", "#REF!"]],
    )
    result = PioneersExcelParser().parse(path)
    assert result.products == []
    assert len(result.invalid_rows) == 1
    assert result.invalid_rows[0].column == HEADER_BARCODE
    assert "error" in result.invalid_rows[0].reason.lower()


def test_duplicate_mapped_header_rejected(tmp_path: Path) -> None:
    path = tmp_path / "dup_mapped_header.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_CODE, HEADER_NAME, HEADER_CODE, HEADER_SALE_PRICE],
        rows=[["95001", "Dup Header Product", "95001-b", "100"]],
    )
    with pytest.raises(AmbiguousMappedHeaderError) as exc_info:
        PioneersExcelParser().parse(path)
    assert HEADER_CODE in str(exc_info.value)


def test_unmapped_duplicate_header_still_allowed(tmp_path: Path) -> None:
    path = tmp_path / "dup_unmapped_header.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_CODE, HEADER_NAME, "مرکز پخش", "مرکز پخش"],
        rows=[["96001", "Unmapped Dup OK", "A", "B"]],
    )
    result = PioneersExcelParser().parse(path)
    assert len(result.products) == 1
    assert result.products[0].code == "96001"


def test_corrupt_workbook_raises_pioneers_excel_error(tmp_path: Path) -> None:
    path = tmp_path / "corrupt.xlsx"
    path.write_bytes(b"this is not a real workbook")
    with pytest.raises(PioneersExcelError) as exc_info:
        PioneersExcelParser().parse(path)
    assert "Failed to open" in str(exc_info.value)
    assert exc_info.value.filename == str(path)
    assert exc_info.value.reason


def test_deferred_worksheet_parse_error_is_wrapped(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Malformed sheet XML during row reading must not leak raw ParseError."""
    from xml.etree.ElementTree import ParseError

    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    path = tmp_path / "deferred_corrupt.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_CODE, HEADER_NAME],
        rows=[["97001", "Will Fail On Read"]],
    )

    def _boom(self, *args: object, **kwargs: object):  # noqa: ANN001
        raise ParseError("malformed worksheet xml")

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", _boom)

    with pytest.raises(PioneersExcelError) as exc_info:
        PioneersExcelParser().parse(path)
    assert "Failed to read" in str(exc_info.value)
    assert exc_info.value.filename == str(path)
    assert exc_info.value.sheet == PRODUCT_SHEET_NAME
    assert "malformed worksheet xml" in (exc_info.value.reason or "")


def test_source_workbook_hash_unchanged(sample_path: Path, tmp_path: Path) -> None:
    working = tmp_path / "immutable_check.xlsx"
    shutil.copy2(sample_path, working)
    before = _sha256(working)
    PioneersExcelParser().parse(working)
    after = _sha256(working)
    assert before == after

    # Also ensure openpyxl still sees the same sheet/title (no accidental save).
    wb = load_workbook(working, read_only=True)
    try:
        assert wb.sheetnames == [PRODUCT_SHEET_NAME]
    finally:
        wb.close()


def test_missing_sheet_raises(tmp_path: Path) -> None:
    path = tmp_path / "wrong_sheet.xlsx"
    wb = Workbook()
    wb.active.title = "Other"
    wb.active["A1"] = "یادداشت"
    wb.save(path)
    with pytest.raises(SheetNotFoundError) as exc_info:
        PioneersExcelParser().parse(path)
    assert HEADER_CODE in str(exc_info.value)
    assert "Other" in str(exc_info.value)


def test_renamed_product_sheet_still_parses(tmp_path: Path) -> None:
    path = tmp_path / "renamed_sheet.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_CODE, HEADER_NAME, HEADER_SALE_PRICE],
        rows=[["91001", "Renamed Sheet Product", "100"]],
    )
    workbook = load_workbook(path)
    workbook.worksheets[0].title = "نمونه کالا"
    workbook.save(path)
    workbook.close()
    result = PioneersExcelParser().parse(path)
    assert len(result.products) == 1
    assert result.products[0].code == "91001"
    assert result.products[0].name == "Renamed Sheet Product"


def test_missing_code_header_message(tmp_path: Path) -> None:
    path = tmp_path / "missing_code_header.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_NAME, HEADER_SALE_PRICE],
        rows=[["Only Name", "100"]],
    )
    with pytest.raises(SheetNotFoundError) as exc_info:
        PioneersExcelParser().parse(path)
    assert HEADER_CODE in str(exc_info.value)


def test_invalid_price_text_collected_for_review(tmp_path: Path) -> None:
    path = tmp_path / "bad_price.xlsx"
    _write_workbook(
        path,
        headers=[HEADER_CODE, HEADER_NAME, HEADER_SALE_PRICE],
        rows=[
            ["92001", "Bad Price Product", "not-a-price"],
            ["92002", "Good Price Product", "1000"],
        ],
    )
    result = PioneersExcelParser().parse(path)
    assert len(result.products) == 1
    assert result.products[0].code == "92002"
    assert len(result.invalid_rows) == 1
    invalid = result.invalid_rows[0]
    assert invalid.row_number == 2
    assert invalid.code == "92001"
    assert invalid.column == HEADER_SALE_PRICE
    assert "price" in invalid.reason.lower()


def test_sample_fixture_has_no_invalid_rows(sample_path: Path) -> None:
    result = PioneersExcelParser().parse(sample_path)
    assert result.invalid_rows == []
