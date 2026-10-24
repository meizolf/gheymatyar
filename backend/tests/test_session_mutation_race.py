"""Per-session locks: explicit ordered concurrency scenarios for review vs export."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from tests.confirmed_metre import prepare_two_exportable_metre_rows


def _product_codes(path: Path) -> set[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    code_idx = headers.index("کد کالا")
    codes = {
        str(row[code_idx])
        for row in ws.iter_rows(min_row=2, values_only=True)
        if row[code_idx] not in (None, "")
    }
    wb.close()
    return codes


def _sale_price_at(path: Path, row_number: int) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    sale_idx = headers.index("نرخ فروش اصلی")
    vals = [
        c.value
        for c in next(ws.iter_rows(min_row=row_number, max_row=row_number))
    ]
    wb.close()
    return str(vals[sale_idx]).replace(",", "")


def test_order_exclude_then_export_omits_excluded_row(
    client, pioneers_azin_excel: Path, dina_pdf: Path, tmp_path: Path
) -> None:
    """Intended order: exclusion completes before export snapshots."""
    session_id, keep, drop = prepare_two_exportable_metre_rows(
        client, pioneers_azin_excel, dina_pdf, tmp_path
    )

    exclude = client.post(
        f"/api/sessions/{session_id}/candidates/{drop['candidate_id']}/review",
        json={"decision": "EXCLUDE"},
    )
    assert exclude.status_code == 200
    assert exclude.json()["exportable"] is False

    export = client.post(f"/api/sessions/{session_id}/export")
    assert export.status_code == 200
    assert export.json()["updated_count"] == 1

    download = client.get(f"/api/sessions/{session_id}/download")
    assert download.status_code == 200
    out = tmp_path / "exclude_then_export.xlsx"
    out.write_bytes(download.content)

    assert _sale_price_at(out, 2) == str(keep["new_price"]).replace(",", "")
    assert _product_codes(out) == {str(keep["pioneers_code"])}


def test_order_export_then_exclude_invalidates_stale_download(
    client, pioneers_azin_excel: Path, dina_pdf: Path, tmp_path: Path
) -> None:
    """Opposite order: export first, then exclude must invalidate that export."""
    session_id, keep, drop = prepare_two_exportable_metre_rows(
        client, pioneers_azin_excel, dina_pdf, tmp_path
    )

    export = client.post(f"/api/sessions/{session_id}/export")
    assert export.status_code == 200
    assert export.json()["updated_count"] == 2

    download = client.get(f"/api/sessions/{session_id}/download")
    assert download.status_code == 200
    first = tmp_path / "export_first.xlsx"
    first.write_bytes(download.content)
    # Both rows were exportable at snapshot time.
    assert _sale_price_at(first, drop["row_number"]) == str(drop["new_price"]).replace(
        ",", ""
    )

    exclude = client.post(
        f"/api/sessions/{session_id}/candidates/{drop['candidate_id']}/review",
        json={"decision": "EXCLUDE"},
    )
    assert exclude.status_code == 200
    # Exclusion invalidates prior export artifact.
    stale = client.get(f"/api/sessions/{session_id}/download")
    assert stale.status_code == 404

    re_export = client.post(f"/api/sessions/{session_id}/export")
    assert re_export.status_code == 200
    assert re_export.json()["updated_count"] == 1

    download2 = client.get(f"/api/sessions/{session_id}/download")
    assert download2.status_code == 200
    second = tmp_path / "export_after_exclude.xlsx"
    second.write_bytes(download2.content)
    assert _sale_price_at(second, 2) == str(keep["new_price"]).replace(",", "")
    assert _product_codes(second) == {str(keep["pioneers_code"])}
