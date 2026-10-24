"""Regressions for stale downloads, commit failures, and variant mappings.

These lock in the corrected behavior of the persistence defects reproduced by
the external review script.
"""

from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pymupdf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.exc import OperationalError

from gheymatyar.domain.supplier_item import SupplierPriceItem
from gheymatyar.parsers.pioneers_excel_parser import HEADER_CODE
from gheymatyar.repositories import MappingRepository, SupplierRepository
from tests.confirmed_metre import prepare_two_exportable_metre_rows
from tests.test_persistent_decisions_api import _restart, _upload_and_match

ROOT = Path(__file__).resolve().parents[2]
EXCEL = ROOT / "fixtures/pioneers/pioneers_azin_subset.xlsx"
PDF = ROOT / "fixtures/suppliers/dina_polymer_1405_07_04.pdf"
BACKEND = Path(__file__).resolve().parents[1]

_REOPEN = """
import json, sqlite3, sys
from pathlib import Path
from fastapi.testclient import TestClient
from gheymatyar.main import create_app
from gheymatyar.services.session_store import SessionStore
from gheymatyar.services.processing import ProcessingService
from tests.test_persistent_decisions_api import _upload_and_match
cfg = json.loads(sys.argv[1])
app = create_app(database_url=cfg["url"], mount_spa=False)
store = SessionStore()
app.state.sessions = store
app.state.processing = ProcessingService(store)
with TestClient(app) as client:
    sid = _upload_and_match(client, Path(cfg["excel"]), Path(cfg["pdf"]))
    rows = client.get(f"/api/sessions/{sid}/candidates").json()["candidates"]
    by_code = {c["supplier_code"]: c for c in rows if c["supplier_code"]}
    for code in cfg["deleted"]:
        assert by_code[code]["status"] == "IGNORED"
    mapped = by_code[cfg["mapped"]]
    assert mapped["match_type"] == "SAVED_MAPPING", mapped
    assert mapped["pioneers_code"] == cfg["target_code"]
    assert mapped["row_number"] != cfg["old_target_row"]
    restored = client.post(
        f"/api/sessions/{sid}/candidates/{by_code[cfg['deleted'][0]]['candidate_id']}/restore"
    )
    assert restored.status_code == 200, restored.text
    sid2 = _upload_and_match(client, Path(cfg["excel"]), Path(cfg["pdf"]))
    rows2 = client.get(f"/api/sessions/{sid2}/candidates").json()["candidates"]
    by_code2 = {c["supplier_code"]: c for c in rows2 if c["supplier_code"]}
    assert by_code2[cfg["deleted"][0]]["status"] != "IGNORED"
    assert by_code2[cfg["deleted"][1]]["status"] == "IGNORED"
    with sqlite3.connect(app.state.database.engine.url.database) as db:
        states = [r[0] for r in db.execute("SELECT decision FROM supplier_product_decisions")]
        assert sorted(states) == ["ACTIVE", "IGNORED"], states
    print("PASS")
"""


def setup_rows(client, tmp_path):
    return prepare_two_exportable_metre_rows(client, EXCEL, PDF, tmp_path)


def _cell(value: object) -> str:
    return str(value).replace(",", "").strip()


def _sale_for_code(sheet, code: str | None):
    """Sale cell of one product in a filtered workbook, or None if it is absent."""
    if not code:
        return None
    headers = {cell.value: cell.column for cell in sheet[1]}
    wanted = str(code)
    for row in range(2, (sheet.max_row or 1) + 1):
        value = sheet.cell(row, headers[HEADER_CODE]).value
        if value is not None and str(value) == wanted:
            return sheet.cell(row, headers["نرخ فروش اصلی"]).value
    return None


def _db_path(client) -> str:
    return client.app.state.database.engine.url.database


def test_stale_download_does_not_return_deleted_price(client, tmp_path: Path) -> None:
    sid, _keep, drop = setup_rows(client, tmp_path)
    exported = client.post(f"/api/sessions/{sid}/export")
    assert exported.status_code == 200, exported.text
    original_count = exported.json()["exported_row_count"]
    assert original_count == 2
    assert exported.json()["content_sha256"]
    other = _upload_and_match(client, tmp_path / "named.xlsx", PDF)
    rows = client.get(f"/api/sessions/{other}/candidates").json()["candidates"]
    target = next(row for row in rows if row["supplier_code"] == drop["supplier_code"])
    ignored = client.post(
        f"/api/sessions/{other}/candidates/{target['candidate_id']}/ignore"
    )
    assert ignored.status_code == 200, ignored.text
    with sqlite3.connect(_db_path(client)) as db:
        assert (
            db.execute(
                "SELECT count(*) FROM supplier_product_decisions WHERE decision='IGNORED'"
            ).fetchone()[0]
            == 1
        )

    download = client.get(f"/api/sessions/{sid}/download")
    assert download.status_code == 200, download.text
    assert download.headers["x-exported-row-count"] == "1"
    assert download.headers["x-exported-row-count"] != str(original_count)
    assert download.headers["x-export-content-sha256"] == hashlib.sha256(
        download.content
    ).hexdigest()
    assert download.headers["x-export-content-sha256"] != exported.json()["content_sha256"]
    workbook = load_workbook(BytesIO(download.content))
    sheet = workbook.active
    assert sheet.max_row == 2
    headers = {cell.value: cell.column for cell in sheet[1]}
    kept_price = _cell(sheet.cell(2, headers["نرخ فروش اصلی"]).value)
    kept_code = sheet.cell(2, headers["کد کالا"]).value
    codes = {
        sheet.cell(row, headers["کد کالا"]).value
        for row in range(2, sheet.max_row + 1)
    }
    workbook.close()
    assert drop["pioneers_code"] not in codes
    assert kept_code == _keep["pioneers_code"]
    assert kept_price == _cell(_keep["new_price"])
    assert kept_price != _cell(drop["new_price"])
    exported = client.post(f"/api/sessions/{sid}/export")
    assert exported.status_code == 200, exported.text
    assert exported.json()["updated_count"] == 1


def test_failed_commit_does_not_report_ignore_success(client, tmp_path: Path) -> None:
    sid, _keep, drop = setup_rows(client, tmp_path)
    error = OperationalError("COMMIT", {}, Exception("injected commit failure"))
    with TestClient(client.app, raise_server_exceptions=False) as no_raise:
        with patch("sqlalchemy.orm.Session.commit", side_effect=error):
            response = no_raise.post(
                f"/api/sessions/{sid}/candidates/{drop['candidate_id']}/ignore"
            )
    with sqlite3.connect(_db_path(client)) as db:
        count = db.execute(
            "SELECT count(*) FROM supplier_product_decisions"
        ).fetchone()[0]
    state = (
        client.app.state.processing.get_session(sid)
        .get_candidate(drop["candidate_id"])
        .status.value
    )
    assert response.status_code >= 400
    assert count == 0
    assert state != "IGNORED"


def test_variant_mappings_do_not_overwrite(client, tmp_path: Path) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    session = client.app.state.processing.get_session(sid)

    def variant(code: str, unit: str = "متر") -> SupplierPriceItem:
        return SupplierPriceItem(
            supplier_id="dina_polymer",
            supplier_code=code,
            supplier_name="کالای مشترک 20",
            unit=unit,
            price="123000",
            raw_values={
                "currency": "rial",
                "price_basis": "per_meter",
                "tax": "included",
            },
        )

    def match_variant(code: str, unit: str = "متر") -> dict:
        session.supplier_items = [variant(code, unit)]
        response = client.post(f"/api/sessions/{sid}/match")
        assert response.status_code == 200, response.text
        return response.json()["candidates"][0]

    first = match_variant("VARIANT-A")
    saved = client.post(
        f"/api/sessions/{sid}/candidates/{first['candidate_id']}/confirm-mapping",
        json={"pioneers_row_number": keep["row_number"], "save": True},
    )
    assert saved.status_code == 200, saved.text
    second = match_variant("VARIANT-B")
    assert second["match_type"] != "SAVED_MAPPING"
    assert second["row_number"] != keep["row_number"]
    saved_b = client.post(
        f"/api/sessions/{sid}/candidates/{second['candidate_id']}/confirm-mapping",
        json={"pioneers_row_number": drop["row_number"], "save": True},
    )
    assert saved_b.status_code == 200, saved_b.text
    rematched = match_variant("VARIANT-A")
    assert rematched["row_number"] == keep["row_number"]
    assert rematched["match_type"] == "SAVED_MAPPING"
    other_unit = match_variant("VARIANT-A", unit="شاخه")
    assert other_unit["row_number"] != keep["row_number"]
    with sqlite3.connect(_db_path(client)) as db:
        count = db.execute(
            "SELECT count(*) FROM product_mappings WHERE supplier_name_raw=?",
            ("کالای مشترک 20",),
        ).fetchone()[0]
    assert count == 2


def test_fresh_process_and_reordered_files(client, tmp_path: Path) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    session = client.app.state.processing.get_session(sid)
    third = next(
        product
        for product in session.products
        if product.row_number not in {keep["row_number"], drop["row_number"]}
    )
    mapped = client.post(
        f"/api/sessions/{sid}/candidates/{keep['candidate_id']}/confirm-mapping",
        json={"pioneers_row_number": third.row_number, "save": True},
    )
    assert mapped.status_code == 200, mapped.text
    unmatched = next(
        row
        for row in client.get(f"/api/sessions/{sid}/candidates").json()["candidates"]
        if row["status"] == "UNMATCHED" and row["supplier_code"]
    )
    for row in (drop, unmatched):
        ignored = client.post(
            f"/api/sessions/{sid}/candidates/{row['candidate_id']}/ignore"
        )
        assert ignored.status_code == 200, ignored.text
    for _ in range(2):
        fresh = client.get(f"/api/sessions/{sid}/candidates").json()["candidates"]
        assert (
            next(row for row in fresh if row["supplier_code"] == drop["supplier_code"])[
                "status"
            ]
            == "IGNORED"
        )
    with sqlite3.connect(_db_path(client)) as db:
        assert (
            db.execute("SELECT count(*) FROM supplier_product_decisions").fetchone()[0]
            == 2
        )
        assert (
            db.execute("SELECT pioneers_code FROM product_mappings").fetchone()[0]
            == third.code
        )

    workbook = load_workbook(tmp_path / "named.xlsx")
    sheet = workbook.active
    values = list(sheet.iter_rows(min_row=2, values_only=True))[::-1]
    for index, row in enumerate(values, 2):
        for column, value in enumerate(row, 1):
            sheet.cell(index, column).value = value
    moved = tmp_path / "moved.xlsx"
    workbook.save(moved)
    workbook.close()
    new_pdf = tmp_path / "newer-reordered.pdf"
    with pymupdf.open(PDF) as doc:
        doc.select(list(range(doc.page_count))[::-1])
        doc.save(new_pdf)
    script = tmp_path / "reopen_check.py"
    script.write_text(_REOPEN, encoding="utf-8")
    cfg = {
        "url": client.app.state.database.url,
        "excel": str(moved),
        "pdf": str(new_pdf),
        "deleted": [drop["supplier_code"], unmatched["supplier_code"]],
        "mapped": keep["supplier_code"],
        "target_code": third.code,
        "old_target_row": third.row_number,
    }
    env = os.environ.copy()
    env["PYTHONPATH"] = str(BACKEND) + os.pathsep + env.get("PYTHONPATH", "")
    result = subprocess.run(
        [sys.executable, str(script), json.dumps(cfg)],
        text=True,
        capture_output=True,
        cwd=BACKEND,
        env=env,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr


def test_export_only_sale_cells_change(client, tmp_path: Path) -> None:
    sid, _keep, _drop = setup_rows(client, tmp_path)
    source = tmp_path / "named.xlsx"
    before_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    assert client.post(f"/api/sessions/{sid}/export").status_code == 200
    response = client.get(f"/api/sessions/{sid}/download")
    assert response.status_code == 200, response.text
    old = load_workbook(source, data_only=False)
    new = load_workbook(BytesIO(response.content), data_only=False)
    assert old.sheetnames == new.sheetnames
    left, right = old.active, new.active
    assert left.max_row > 3
    assert right.max_row == 3
    assert left.max_column == right.max_column
    sale = next(cell.column for cell in left[1] if cell.value == "نرخ فروش اصلی")
    changed: list[str] = []
    for row_idx in range(1, 4):
        for column in range(1, left.max_column + 1):
            source_cell = left.cell(row_idx, column)
            updated = right.cell(row_idx, column)
            if source_cell.value != updated.value:
                changed.append(updated.coordinate)
                assert column == sale
            assert source_cell._style == updated._style
    assert len(changed) == 2, changed
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before_hash
    old.close()
    new.close()


def test_supplier_isolation_via_match_api(client, tmp_path: Path) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    assert (
        client.post(
            f"/api/sessions/{sid}/candidates/{drop['candidate_id']}/ignore"
        ).status_code
        == 200
    )
    assert (
        client.post(
            f"/api/sessions/{sid}/candidates/{keep['candidate_id']}/confirm-mapping",
            json={"pioneers_row_number": drop["row_number"], "save": True},
        ).status_code
        == 200
    )
    session = client.app.state.processing.get_session(sid)
    session.supplier_items = [
        item.model_copy(update={"supplier_id": "newpipe"})
        for item in session.supplier_items
        if item.supplier_code in {keep["supplier_code"], drop["supplier_code"]}
    ]
    session.supplier_key = "newpipe"
    response = client.post(f"/api/sessions/{sid}/match")
    assert response.status_code == 200, response.text
    rows = response.json()["candidates"]
    assert all(row["status"] != "IGNORED" for row in rows)
    kept = next(row for row in rows if row["supplier_code"] == keep["supplier_code"])
    assert kept["row_number"] == keep["row_number"]
    assert kept["match_type"] == "EXACT_NAME"


def test_remap_to_codeless_target_clears_obsolete_code(client, tmp_path: Path) -> None:
    _sid, keep, drop = setup_rows(client, tmp_path)
    workbook = load_workbook(tmp_path / "named.xlsx")
    sheet = workbook.active
    code_col = next(cell.column for cell in sheet[1] if cell.value == HEADER_CODE)
    sheet.cell(4, code_col).value = None
    path = tmp_path / "codeless.xlsx"
    workbook.save(path)
    workbook.close()

    sid = _upload_and_match(client, path, PDF)
    rows = client.get(f"/api/sessions/{sid}/candidates").json()["candidates"]
    candidate = next(row for row in rows if row["supplier_code"] == keep["supplier_code"])
    response = None
    for target in (drop["row_number"], 4):
        response = client.post(
            f"/api/sessions/{sid}/candidates/{candidate['candidate_id']}/confirm-mapping",
            json={"pioneers_row_number": target, "save": True},
        )
        assert response.status_code == 200, response.text
    assert response is not None
    assert response.json()["row_number"] == 4
    assert response.json()["pioneers_code"] is None
    with sqlite3.connect(_db_path(client)) as db:
        stored = db.execute(
            "SELECT pioneers_code, pioneers_name_raw FROM product_mappings"
        ).fetchone()
    assert stored[0] is None

    fresh = _upload_and_match(client, path, PDF)
    row = next(
        item
        for item in client.get(f"/api/sessions/{fresh}/candidates").json()["candidates"]
        if item["supplier_code"] == keep["supplier_code"]
    )
    assert row["status"] == "MATCHED"
    assert row["match_type"] == "SAVED_MAPPING"
    assert row["row_number"] == 4
    assert row["pioneers_code"] is None


def test_removed_saved_target_is_excluded(client, tmp_path: Path) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    target = next(
        product
        for product in client.app.state.processing.get_session(sid).products
        if product.row_number not in {keep["row_number"], drop["row_number"]}
    )
    assert (
        client.post(
            f"/api/sessions/{sid}/candidates/{keep['candidate_id']}/confirm-mapping",
            json={"pioneers_row_number": target.row_number, "save": True},
        ).status_code
        == 200
    )
    workbook = load_workbook(tmp_path / "named.xlsx")
    workbook.active.delete_rows(target.row_number)
    path = tmp_path / "missing.xlsx"
    workbook.save(path)
    workbook.close()
    fresh = _upload_and_match(client, path, PDF)
    row = next(
        item
        for item in client.get(f"/api/sessions/{fresh}/candidates").json()["candidates"]
        if item["supplier_code"] == keep["supplier_code"]
    )
    assert row["status"] != "MATCHED"
    assert row["review_decision"] == "EXCLUDE"
    assert row["exportable"] is False
    assert any("saved Pioneers mapping no longer" in warning for warning in row["warnings"])
    with sqlite3.connect(_db_path(client)) as db:
        assert (
            db.execute("SELECT pioneers_code FROM product_mappings").fetchone()[0]
            == target.code
        )


@pytest.mark.parametrize("action", ["confirm-mapping", "restore"])
def test_commit_failure_preserves_previous_state(
    client, tmp_path: Path, action: str
) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    base = f"/api/sessions/{sid}/candidates/{keep['candidate_id']}"
    if action == "restore":
        assert client.post(base + "/ignore").status_code == 200
    before = client.app.state.processing.get_session(sid).get_candidate(
        keep["candidate_id"]
    )
    error = OperationalError("COMMIT", {}, Exception("injected commit failure"))
    with TestClient(client.app, raise_server_exceptions=False) as no_raise:
        with patch("sqlalchemy.orm.Session.commit", side_effect=error):
            if action == "restore":
                response = no_raise.post(base + "/restore")
            else:
                response = no_raise.post(
                    base + "/confirm-mapping",
                    json={"pioneers_row_number": drop["row_number"], "save": True},
                )
    assert response.status_code >= 400
    after = client.app.state.processing.get_session(sid).get_candidate(
        keep["candidate_id"]
    )
    assert after.status == before.status
    assert after.match_type == before.match_type
    assert after.row_number == before.row_number
    with sqlite3.connect(_db_path(client)) as db:
        if action == "restore":
            assert (
                db.execute(
                    "SELECT decision FROM supplier_product_decisions"
                ).fetchone()[0]
                == "IGNORED"
            )
            assert after.status.value == "IGNORED"
        else:
            assert db.execute("SELECT count(*) FROM product_mappings").fetchone()[0] == 0
            match_value = (
                after.match_type.value if after.match_type is not None else None
            )
            assert match_value != "MANUAL"


@pytest.mark.parametrize("action", ["ignore", "restore", "confirm-mapping"])
def test_flush_failure_preserves_previous_state(
    client, tmp_path: Path, action: str
) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    target = drop if action == "ignore" else keep
    base = f"/api/sessions/{sid}/candidates/{target['candidate_id']}"
    if action == "restore":
        assert client.post(base + "/ignore").status_code == 200
    before = client.app.state.processing.get_session(sid).get_candidate(
        target["candidate_id"]
    )
    error = OperationalError("FLUSH", {}, Exception("injected flush failure"))
    with TestClient(client.app, raise_server_exceptions=False) as no_raise:
        with patch("sqlalchemy.orm.Session.flush", side_effect=error):
            if action == "ignore":
                response = no_raise.post(base + "/ignore")
            elif action == "restore":
                response = no_raise.post(base + "/restore")
            else:
                response = no_raise.post(
                    base + "/confirm-mapping",
                    json={"pioneers_row_number": drop["row_number"], "save": True},
                )
    assert response.status_code >= 400
    after = client.app.state.processing.get_session(sid).get_candidate(
        target["candidate_id"]
    )
    assert after.status == before.status
    assert after.row_number == before.row_number
    with sqlite3.connect(_db_path(client)) as db:
        if action == "confirm-mapping":
            assert db.execute("SELECT count(*) FROM product_mappings").fetchone()[0] == 0
        elif action == "ignore":
            assert (
                db.execute(
                    "SELECT count(*) FROM supplier_product_decisions"
                ).fetchone()[0]
                == 0
            )
        else:
            assert (
                db.execute(
                    "SELECT decision FROM supplier_product_decisions"
                ).fetchone()[0]
                == "IGNORED"
            )


def test_direct_export_uses_corrected_mapping(client, tmp_path: Path) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    excluded = client.post(
        f"/api/sessions/{sid}/candidates/{drop['candidate_id']}/review",
        json={"decision": "EXCLUDE"},
    )
    assert excluded.status_code == 200, excluded.text
    other = _upload_and_match(client, tmp_path / "named.xlsx", PDF)
    rows = client.get(f"/api/sessions/{other}/candidates").json()["candidates"]
    candidate = next(row for row in rows if row["supplier_code"] == keep["supplier_code"])
    target = next(
        product
        for product in client.app.state.processing.get_session(other).products
        if product.row_number not in {keep["row_number"], drop["row_number"]}
    )
    correction = client.post(
        f"/api/sessions/{other}/candidates/{candidate['candidate_id']}/confirm-mapping",
        json={"pioneers_row_number": target.row_number, "save": True},
    )
    assert correction.status_code == 200, correction.text
    exported = client.post(f"/api/sessions/{sid}/export")
    assert exported.status_code == 200, exported.text
    downloaded = client.get(f"/api/sessions/{sid}/download")
    assert downloaded.status_code == 200, downloaded.text
    old = load_workbook(tmp_path / "named.xlsx")
    new = load_workbook(BytesIO(downloaded.content))
    column = next(cell.column for cell in old.active[1] if cell.value == "نرخ فروش اصلی")
    exported_sale = _sale_for_code(new.active, target.code)
    assert exported_sale is not None
    assert exported_sale != old.active.cell(target.row_number, column).value
    assert _sale_for_code(new.active, keep["pioneers_code"]) is None
    old.close()
    new.close()


def test_legacy_mapping_does_not_leak_between_single_variant_uploads(
    client, tmp_path: Path
) -> None:
    sid, keep, _drop = setup_rows(client, tmp_path)
    with client.app.state.database.session() as db:
        supplier = next(
            row
            for row in SupplierRepository(db).list_all()
            if row.parser_type == "dina_polymer"
        )
        MappingRepository(db).create(
            supplier_id=supplier.id,
            supplier_name_raw="کالای مشترک 20",
            pioneers_name_raw=keep["pioneers_name"],
            pioneers_code=keep["pioneers_code"],
            legacy_unscoped=True,
        )

    def match_code(code: str) -> dict:
        session = client.app.state.processing.get_session(sid)
        session.supplier_items = [
            SupplierPriceItem(
                supplier_id="dina_polymer",
                supplier_code=code,
                supplier_name="کالای مشترک 20",
                unit="متر",
                price="123000",
                raw_values={
                    "currency": "rial",
                    "price_basis": "meter",
                },
            )
        ]
        response = client.post(f"/api/sessions/{sid}/match")
        assert response.status_code == 200, response.text
        return response.json()["candidates"][0]

    for code in ("LEGACY-A", "LEGACY-B"):
        row = match_code(code)
        assert row["match_type"] != "SAVED_MAPPING"
        assert row["row_number"] != keep["row_number"]
        assert row["review_decision"] == "EXCLUDE"

    _restart(client)
    fresh = _upload_and_match(client, tmp_path / "named.xlsx", PDF)
    session = client.app.state.processing.get_session(fresh)
    session.supplier_items = [
        SupplierPriceItem(
            supplier_id="dina_polymer",
            supplier_code="LEGACY-C",
            supplier_name="کالای مشترک 20",
            unit="عدد",
            price="123000",
            raw_values={"currency": "rial", "price_basis": "meter"},
        )
    ]
    restarted = client.post(f"/api/sessions/{fresh}/match")
    assert restarted.status_code == 200, restarted.text
    again = restarted.json()["candidates"][0]
    assert again["match_type"] != "SAVED_MAPPING"
    assert again["row_number"] != keep["row_number"]


def test_restore_preserves_unrelated_exclusion(client, tmp_path: Path) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    rematched = client.post(
        f"/api/sessions/{sid}/match",
        json={"suspicious_percent_threshold": "1000"},
    )
    assert rematched.status_code == 200, rematched.text
    rows = rematched.json()["candidates"]
    keep_row = next(row for row in rows if row["supplier_code"] == keep["supplier_code"])
    drop_row = next(row for row in rows if row["supplier_code"] == drop["supplier_code"])
    assert keep_row["warnings"] == []
    keep_url = f"/api/sessions/{sid}/candidates/{keep_row['candidate_id']}"
    drop_url = f"/api/sessions/{sid}/candidates/{drop_row['candidate_id']}"
    assert client.post(keep_url + "/review", json={"decision": "EXCLUDE"}).status_code == 200
    assert client.post(drop_url + "/ignore").status_code == 200
    restored = client.post(drop_url + "/restore")
    assert restored.status_code == 200, restored.text
    current = client.get(f"/api/sessions/{sid}/candidates").json()["candidates"]
    kept = next(row for row in current if row["candidate_id"] == keep_row["candidate_id"])
    assert kept["review_decision"] == "EXCLUDE"
    assert kept["exportable"] is False
    approved = client.post(
        f"/api/sessions/{sid}/candidates/{restored.json()['candidate_id']}/review",
        json={"decision": "APPROVE"},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["exportable"] is True
    exported = client.post(f"/api/sessions/{sid}/export")
    assert exported.status_code == 200, exported.text
    downloaded = client.get(f"/api/sessions/{sid}/download")
    assert downloaded.status_code == 200, downloaded.text
    source = load_workbook(tmp_path / "named.xlsx")
    output = load_workbook(BytesIO(downloaded.content))
    column = next(
        cell.column for cell in source.active[1] if cell.value == "نرخ فروش اصلی"
    )
    assert _sale_for_code(output.active, keep["pioneers_code"]) is None
    restored_sale = _sale_for_code(output.active, drop["pioneers_code"])
    assert restored_sale is not None
    assert restored_sale != source.active.cell(drop["row_number"], column).value
    source.close()
    output.close()


def test_session_only_manual_mapping_survives_export(client, tmp_path: Path) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    excluded = client.post(
        f"/api/sessions/{sid}/candidates/{drop['candidate_id']}/review",
        json={"decision": "EXCLUDE"},
    )
    assert excluded.status_code == 200, excluded.text
    endpoint = f"/api/sessions/{sid}/candidates/{keep['candidate_id']}"
    mapped = client.post(
        endpoint + "/confirm-mapping",
        json={"pioneers_row_number": drop["row_number"], "save": False},
    )
    assert mapped.status_code == 200, mapped.text
    assert mapped.json()["row_number"] == drop["row_number"]
    assert mapped.json()["match_type"] == "MANUAL"
    approved = client.post(endpoint + "/review", json={"decision": "APPROVE"})
    assert approved.status_code == 200, approved.text
    assert approved.json()["exportable"] is True
    with sqlite3.connect(_db_path(client)) as db:
        assert db.execute("SELECT COUNT(*) FROM product_mappings").fetchone()[0] == 0

    def exported_target() -> dict:
        response = client.post(f"/api/sessions/{sid}/export")
        assert response.status_code == 200, response.text
        current = client.get(f"/api/sessions/{sid}/candidates").json()["candidates"]
        return next(row for row in current if row["candidate_id"] == keep["candidate_id"])

    first = exported_target()
    assert first["row_number"] == drop["row_number"]
    assert first["match_type"] == "MANUAL"
    downloaded = client.get(f"/api/sessions/{sid}/download")
    assert downloaded.status_code == 200, downloaded.text
    source = load_workbook(tmp_path / "named.xlsx")
    output = load_workbook(BytesIO(downloaded.content))
    column = next(
        cell.column for cell in source.active[1] if cell.value == "نرخ فروش اصلی"
    )
    assert _sale_for_code(output.active, keep["pioneers_code"]) is None
    mapped_sale = _sale_for_code(output.active, drop["pioneers_code"])
    assert mapped_sale is not None
    assert mapped_sale != source.active.cell(drop["row_number"], column).value
    source.close()
    output.close()
    second = exported_target()
    assert second["row_number"] == drop["row_number"]
    assert second["match_type"] == "MANUAL"
    with sqlite3.connect(_db_path(client)) as db:
        assert db.execute("SELECT COUNT(*) FROM product_mappings").fetchone()[0] == 0


@pytest.mark.parametrize("save", [True, False])
def test_unrelated_ignore_does_not_hide_external_correction(
    client, tmp_path: Path, save: bool
) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    path = f"/api/sessions/{sid}/candidates/{keep['candidate_id']}"
    mapped = client.post(
        path + "/confirm-mapping",
        json={"pioneers_row_number": keep["row_number"], "save": save},
    )
    assert mapped.status_code == 200, mapped.text
    assert mapped.json()["match_type"] == "MANUAL"
    assert client.post(path + "/review", json={"decision": "APPROVE"}).status_code == 200
    other = _upload_and_match(client, tmp_path / "named.xlsx", PDF)
    row = next(
        candidate
        for candidate in client.get(f"/api/sessions/{other}/candidates").json()["candidates"]
        if candidate["supplier_code"] == keep["supplier_code"]
    )
    target = next(
        product
        for product in client.app.state.processing.get_session(other).products
        if product.row_number not in {keep["row_number"], drop["row_number"]}
    )
    corrected = client.post(
        f"/api/sessions/{other}/candidates/{row['candidate_id']}/confirm-mapping",
        json={"pioneers_row_number": target.row_number, "save": True},
    )
    assert corrected.status_code == 200, corrected.text
    deleted = client.post(
        f"/api/sessions/{sid}/candidates/{drop['candidate_id']}/ignore"
    )
    assert deleted.status_code == 200, deleted.text
    with sqlite3.connect(_db_path(client)) as db:
        saved = db.execute("SELECT pioneers_code FROM product_mappings").fetchone()
        assert saved is not None
        assert saved[0] == target.code
    exported = client.post(f"/api/sessions/{sid}/export")
    assert exported.status_code == 200, exported.text
    downloaded = client.get(f"/api/sessions/{sid}/download")
    assert downloaded.status_code == 200, downloaded.text
    before = load_workbook(tmp_path / "named.xlsx")
    after = load_workbook(BytesIO(downloaded.content))
    column = next(
        cell.column for cell in before.active[1] if cell.value == "نرخ فروش اصلی"
    )
    corrected_sale = _sale_for_code(after.active, target.code)
    assert corrected_sale is not None
    assert corrected_sale != before.active.cell(target.row_number, column).value
    assert _sale_for_code(after.active, keep["pioneers_code"]) is None
    before.close()
    after.close()
    if not save:
        with sqlite3.connect(_db_path(client)) as db:
            codes = [
                item[0]
                for item in db.execute(
                    "SELECT pioneers_code FROM product_mappings"
                ).fetchall()
            ]
        assert codes == [target.code]


def test_unrelated_ignore_keeps_unchanged_session_override(
    client, tmp_path: Path
) -> None:
    sid, keep, drop = setup_rows(client, tmp_path)
    rows = client.get(f"/api/sessions/{sid}/candidates").json()["candidates"]
    other = next(
        row
        for row in rows
        if row["candidate_id"] not in {keep["candidate_id"], drop["candidate_id"]}
    )
    assert (
        client.post(
            f"/api/sessions/{sid}/candidates/{drop['candidate_id']}/review",
            json={"decision": "EXCLUDE"},
        ).status_code
        == 200
    )
    endpoint = f"/api/sessions/{sid}/candidates/{keep['candidate_id']}"
    mapped = client.post(
        endpoint + "/confirm-mapping",
        json={"pioneers_row_number": drop["row_number"], "save": False},
    )
    assert mapped.status_code == 200, mapped.text
    assert client.post(endpoint + "/review", json={"decision": "APPROVE"}).status_code == 200
    assert (
        client.post(f"/api/sessions/{sid}/candidates/{other['candidate_id']}/ignore").status_code
        == 200
    )
    exported = client.post(f"/api/sessions/{sid}/export")
    assert exported.status_code == 200, exported.text
    current = client.get(f"/api/sessions/{sid}/candidates").json()["candidates"]
    kept = next(row for row in current if row["candidate_id"] == keep["candidate_id"])
    assert kept["row_number"] == drop["row_number"]
    assert kept["match_type"] == "MANUAL"
    downloaded = client.get(f"/api/sessions/{sid}/download")
    assert downloaded.status_code == 200, downloaded.text
    source = load_workbook(tmp_path / "named.xlsx")
    output = load_workbook(BytesIO(downloaded.content))
    column = next(
        cell.column for cell in source.active[1] if cell.value == "نرخ فروش اصلی"
    )
    assert _sale_for_code(output.active, keep["pioneers_code"]) is None
    mapped_sale = _sale_for_code(output.active, drop["pioneers_code"])
    assert mapped_sale is not None
    assert mapped_sale != source.active.cell(drop["row_number"], column).value
    source.close()
    output.close()
    with sqlite3.connect(_db_path(client)) as db:
        assert db.execute("SELECT COUNT(*) FROM product_mappings").fetchone()[0] == 0
