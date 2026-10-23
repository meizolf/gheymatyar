"""API acceptance: ignores and manual mappings survive a new session."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from gheymatyar.domain.enums import MatchStatus, ReviewDecision
from gheymatyar.repositories import ProductDecisionRepository, SupplierRepository
from gheymatyar.services.processing import ProcessingService
from gheymatyar.services.session_store import SessionStore
from tests.confirmed_metre import prepare_two_exportable_metre_rows


def _restart(client) -> None:
    """Drop in-memory sessions while keeping the same SQLite file."""
    store = SessionStore()
    client.app.state.sessions = store
    client.app.state.processing = ProcessingService(store)


def _upload_and_match(client, excel: Path, pdf: Path) -> str:
    session_id = client.post("/api/sessions").json()["session_id"]
    with excel.open("rb") as handle:
        uploaded = client.post(
            f"/api/sessions/{session_id}/pioneers-excel",
            files={
                "file": (
                    "p.xlsx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert uploaded.status_code == 200, uploaded.text
    with pdf.open("rb") as handle:
        parsed = client.post(
            f"/api/sessions/{session_id}/supplier-pdf",
            data={"supplier_key": "dina_polymer"},
            files={"file": ("dina.pdf", handle, "application/pdf")},
        )
    assert parsed.status_code == 200, parsed.text
    matched = client.post(f"/api/sessions/{session_id}/match")
    assert matched.status_code == 200, matched.text
    return session_id


def test_ignore_matched_and_unmatched_persists_and_skips_export(
    client, pioneers_azin_excel: Path, dina_pdf: Path, tmp_path: Path
) -> None:
    session_id, keep, drop = prepare_two_exportable_metre_rows(
        client, pioneers_azin_excel, dina_pdf, tmp_path
    )
    listed = client.get(f"/api/sessions/{session_id}/candidates").json()
    unmatched = next(
        row for row in listed["candidates"] if row["status"] == "UNMATCHED"
    )

    ignored_matched = client.post(
        f"/api/sessions/{session_id}/candidates/{drop['candidate_id']}/ignore"
    )
    assert ignored_matched.status_code == 200, ignored_matched.text
    assert ignored_matched.json()["status"] == "IGNORED"
    assert ignored_matched.json()["exportable"] is False

    ignored_unmatched = client.post(
        f"/api/sessions/{session_id}/candidates/{unmatched['candidate_id']}/ignore"
    )
    assert ignored_unmatched.status_code == 200, ignored_unmatched.text
    assert ignored_unmatched.json()["status"] == "IGNORED"

    # Repeating ignore is idempotent.
    again = client.post(
        f"/api/sessions/{session_id}/candidates/{drop['candidate_id']}/ignore"
    )
    assert again.status_code == 200, again.text

    after = client.get(f"/api/sessions/{session_id}/candidates").json()
    by_id = {row["candidate_id"]: row for row in after["candidates"]}
    assert by_id[drop["candidate_id"]]["status"] == "IGNORED"
    assert by_id[unmatched["candidate_id"]]["status"] == "IGNORED"
    assert after["ignored_count"] >= 2
    assert (
        after["matched_count"]
        + after["unmatched_count"]
        + after["ambiguous_count"]
        + after["ignored_count"]
        == len(after["candidates"])
    )
    preview = client.get(f"/api/sessions/{session_id}/preview").json()
    preview_ids = {row["candidate_id"] for row in preview["candidates"]}
    assert drop["candidate_id"] not in preview_ids
    assert unmatched["candidate_id"] not in preview_ids

    blocked = client.post(
        f"/api/sessions/{session_id}/candidates/{drop['candidate_id']}/review",
        json={"decision": "APPROVE"},
    )
    assert blocked.status_code == 400
    assert "restored" in blocked.json()["detail"]

    # Stale session: put the ignored row back to an approved match in memory.
    processing = client.app.state.processing
    session = processing.get_session(session_id)
    stale = session.get_candidate(drop["candidate_id"])
    assert stale is not None
    session.replace_candidate(
        stale.model_copy(
            update={
                "status": MatchStatus.MATCHED,
                "review_decision": ReviewDecision.APPROVE,
                "row_number": drop["row_number"],
                "pioneers_name": drop["pioneers_name"],
                "pioneers_code": drop["pioneers_code"],
                "new_price": stale.new_price,
            }
        )
    )

    export = client.post(f"/api/sessions/{session_id}/export")
    assert export.status_code == 200, export.text
    assert export.json()["updated_count"] == 1
    download = client.get(f"/api/sessions/{session_id}/download")
    assert download.status_code == 200
    out = tmp_path / "out.xlsx"
    out.write_bytes(download.content)
    workbook = load_workbook(out, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    sale_idx = headers.index("نرخ فروش اصلی")
    drop_row = [
        cell.value
        for cell in next(
            sheet.iter_rows(min_row=drop["row_number"], max_row=drop["row_number"])
        )
    ]
    assert str(drop_row[sale_idx]).replace(",", "") == str(drop["old_price"]).replace(
        ",", ""
    )
    workbook.close()

    with client.app.state.database.session() as db:
        suppliers = SupplierRepository(db).list_all()
        dina = next(row for row in suppliers if row.parser_type == "dina_polymer")
        other = SupplierRepository(db).create(name="تأمین‌کننده دیگر", parser_type="newpipe")
        decisions = ProductDecisionRepository(db)
        assert decisions.ignored_identity_keys(dina.id)
        assert decisions.ignored_identity_keys(other.id) == set()

    edited = tmp_path / "named.xlsx"
    _restart(client)
    reloaded = _upload_and_match(client, edited, dina_pdf)
    reloaded_rows = client.get(f"/api/sessions/{reloaded}/candidates").json()
    by_code = {
        row["supplier_code"]: row
        for row in reloaded_rows["candidates"]
        if row["supplier_code"]
    }
    assert by_code[drop["supplier_code"]]["status"] == "IGNORED"
    assert by_code[unmatched["supplier_code"]]["status"] == "IGNORED"
    assert by_code[keep["supplier_code"]]["status"] != "IGNORED"

    restored = client.post(
        f"/api/sessions/{reloaded}/candidates/{by_code[drop['supplier_code']]['candidate_id']}/restore"
    )
    assert restored.status_code == 200, restored.text
    assert restored.json()["status"] != "IGNORED"
    final_rows = client.get(f"/api/sessions/{reloaded}/candidates").json()["candidates"]
    restored_row = next(
        row for row in final_rows if row["supplier_code"] == drop["supplier_code"]
    )
    assert restored_row["status"] == "MATCHED"
    still_ignored = next(
        row for row in final_rows if row["supplier_code"] == unmatched["supplier_code"]
    )
    assert still_ignored["status"] == "IGNORED"


def test_corrected_mapping_persists_and_overrides_exact_name(
    client, pioneers_azin_excel: Path, dina_pdf: Path, tmp_path: Path
) -> None:
    session_id, keep, drop = prepare_two_exportable_metre_rows(
        client, pioneers_azin_excel, dina_pdf, tmp_path
    )
    products = client.app.state.processing.get_session(session_id).products
    third = next(
        product
        for product in products
        if product.row_number not in {keep["row_number"], drop["row_number"]}
    )
    first = client.post(
        f"/api/sessions/{session_id}/candidates/{keep['candidate_id']}/confirm-mapping",
        json={"pioneers_row_number": drop["row_number"], "save": True},
    )
    assert first.status_code == 200, first.text
    assert first.json()["match_type"] == "MANUAL"
    assert first.json()["pioneers_code"] == drop["pioneers_code"]

    listed = client.get("/api/suppliers/dina_polymer/mappings").json()["mappings"]
    mine = [row for row in listed if row["supplier_name_raw"] == keep["supplier_name"]]
    assert len(mine) == 1
    mapping_id = mine[0]["id"]

    second = client.post(
        f"/api/sessions/{session_id}/candidates/{keep['candidate_id']}/confirm-mapping",
        json={"pioneers_row_number": third.row_number, "save": True},
    )
    assert second.status_code == 200, second.text
    assert second.json()["row_number"] == third.row_number

    listed_again = client.get("/api/suppliers/dina_polymer/mappings").json()["mappings"]
    mine_again = [
        row for row in listed_again if row["supplier_name_raw"] == keep["supplier_name"]
    ]
    assert len(mine_again) == 1
    assert mine_again[0]["id"] == mapping_id
    assert mine_again[0]["pioneers_code"] == third.code

    _restart(client)
    reloaded = _upload_and_match(client, tmp_path / "named.xlsx", dina_pdf)
    rows = client.get(f"/api/sessions/{reloaded}/candidates").json()["candidates"]
    rematched = next(row for row in rows if row["supplier_code"] == keep["supplier_code"])
    assert rematched["match_type"] == "SAVED_MAPPING"
    assert rematched["status"] == "MATCHED"
    assert rematched["row_number"] == third.row_number
    assert rematched["pioneers_code"] == third.code
    assert rematched["row_number"] != keep["row_number"]
