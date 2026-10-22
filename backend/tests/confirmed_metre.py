"""Build a workbook export that is allowed only after an explicit Rial per-metre confirmation."""

from __future__ import annotations

import shutil
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from gheymatyar.parsers.pioneers_excel_parser import HEADER_NAME
from gheymatyar.suppliers.dina_polymer import DinaPolymerProcessor

METRE_CODES = ("825331050", "825332050")


def enable_rial_per_metre(client: TestClient) -> None:
    response = client.put(
        "/api/settings",
        json={
            "settings": {
                "pioneers_sale_currency": "rial",
                "pioneers_sale_basis": "meter",
            }
        },
    )
    assert response.status_code == 200, response.text


def metre_items(dina_pdf: Path):
    parsed = DinaPolymerProcessor().parse_pdf(dina_pdf)
    by_code = {item.supplier_code: item for item in parsed.items}
    return [by_code[code] for code in METRE_CODES]


def prepare_two_exportable_metre_rows(
    client: TestClient,
    pioneers_excel: Path,
    dina_pdf: Path,
    tmp_path: Path,
) -> tuple[str, dict, dict]:
    """Two Dina per-metre rows, after an explicit Rial per-metre confirmation."""
    enable_rial_per_metre(client)
    items = metre_items(dina_pdf)
    edited = tmp_path / "named.xlsx"
    workbook_with_supplier_names(
        pioneers_excel, edited, [item.supplier_name for item in items]
    )
    session_id = client.post("/api/sessions").json()["session_id"]
    with edited.open("rb") as handle:
        excel = client.post(
            f"/api/sessions/{session_id}/pioneers-excel",
            files={
                "file": (
                    "p.xlsx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert excel.status_code == 200, excel.text
    with dina_pdf.open("rb") as handle:
        pdf = client.post(
            f"/api/sessions/{session_id}/supplier-pdf",
            data={"supplier_key": "dina_polymer"},
            files={"file": ("dina.pdf", handle, "application/pdf")},
        )
    assert pdf.status_code == 200, pdf.text
    assert client.post(f"/api/sessions/{session_id}/match").status_code == 200
    candidates = client.get(f"/api/sessions/{session_id}/candidates").json()["candidates"]
    wanted = {item.supplier_code: item for item in items}
    chosen = [row for row in candidates if row["supplier_code"] in wanted]
    assert len(chosen) == 2
    for row in chosen:
        review = client.post(
            f"/api/sessions/{session_id}/candidates/{row['candidate_id']}/review",
            json={"decision": "APPROVE"},
        )
        assert review.status_code == 200, review.text
    chosen_ids = {row["candidate_id"] for row in chosen}
    for row in candidates:
        if row["candidate_id"] in chosen_ids:
            continue
        excluded = client.post(
            f"/api/sessions/{session_id}/candidates/{row['candidate_id']}/review",
            json={"decision": "EXCLUDE"},
        )
        assert excluded.status_code == 200, excluded.text
    preview = client.get(f"/api/sessions/{session_id}/preview").json()["candidates"]
    keep = next(row for row in preview if row["supplier_code"] == items[0].supplier_code)
    drop = next(row for row in preview if row["supplier_code"] == items[1].supplier_code)
    assert keep["exportable"] is True
    assert drop["exportable"] is True
    assert keep["new_price"] == str(items[0].price)
    return session_id, keep, drop


def workbook_with_supplier_names(source: Path, dest: Path, names: list[str]) -> None:
    shutil.copy(source, dest)
    workbook = load_workbook(dest)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    name_col = headers[HEADER_NAME]
    for offset, name in enumerate(names):
        sheet.cell(2 + offset, name_col).value = name
    workbook.save(dest)
    workbook.close()
