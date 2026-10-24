"""Approved Excel exports contain only those product rows."""

from __future__ import annotations

import hashlib
import shutil
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

from gheymatyar.domain.enums import MatchStatus, ReviewDecision
from gheymatyar.domain.update_candidate import UpdateCandidate
from gheymatyar.parsers.pioneers_excel_parser import (
    HEADER_CODE,
    HEADER_NAME,
    HEADER_SALE_PRICE,
    PRODUCT_SHEET_NAME,
)
from gheymatyar.parsers.pioneers_html_xls import convert_pioneers_html_xls_to_xlsx
from gheymatyar.services.errors import DuplicateUpdateError, PioneersExcelWriteError
from gheymatyar.services.pioneers_excel_writer import update_final_sale_prices
from gheymatyar.services.processing import ProcessingService
from gheymatyar.services.session_store import SessionStore

REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURE = REPO_ROOT / "fixtures" / "pioneers" / "pioneers_sample.xlsx"
HTML_XLS = REPO_ROOT / "fixtures" / "pioneers" / "pioneers_html_sample.xls"
INPUT_ROWS = 6700
APPROVED_SOURCE_ROWS = (2, 10, 50, 100, 500, 1000, 2500, 4000, 6700)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _headers(sheet) -> dict[str, int]:
    return {
        cell.value: cell.column
        for cell in sheet[1]
        if isinstance(cell.value, str) and cell.value.strip()
    }


def _expand_real_catalog(destination: Path) -> list[dict[str, object]]:
    """Copy the real Pioneers header row and grow it to 6,700 product rows."""
    shutil.copy2(FIXTURE, destination)
    workbook = load_workbook(destination)
    sheet = workbook[PRODUCT_SHEET_NAME]
    headers = _headers(sheet)
    code_col = headers[HEADER_CODE]
    name_col = headers[HEADER_NAME]
    sale_col = headers[HEADER_SALE_PRICE]
    template = [
        sheet.cell(2, column).value for column in range(1, sheet.max_column + 1)
    ]
    for row_number in range(3, INPUT_ROWS + 2):
        for column, value in enumerate(template, start=1):
            sheet.cell(row_number, column).value = value
        sheet.cell(row_number, code_col).value = f"{200000 + row_number}"
        sheet.cell(row_number, name_col).value = f"Catalog product {row_number}"
        sheet.cell(row_number, sale_col).value = "1000"
    sheet.cell(2, code_col).value = "200002"
    sheet.cell(2, name_col).value = "Catalog product 2"
    sheet.cell(2, sale_col).value = "1000"

    sheet.row_dimensions[100].hidden = True
    sheet.row_dimensions[4000].hidden = True
    sheet.cell(INPUT_ROWS + 2, 1).value = None
    sheet.auto_filter.ref = f"A1:AN{INPUT_ROWS + 1}"
    table = Table(
        displayName="AllProducts",
        ref=f"A1:AN{INPUT_ROWS + 1}",
    )
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    workbook.defined_names.add(
        DefinedName(
            name="FullCatalog",
            attr_text=f"'{PRODUCT_SHEET_NAME}'!$A$2:$AN${INPUT_ROWS + 1}",
        )
    )
    hidden = workbook.create_sheet("کاتالوگ پنهان")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "نسخه پنهان"
    hidden["A2"] = HEADER_CODE
    hidden["B2"] = HEADER_NAME
    hidden["C2"] = HEADER_SALE_PRICE
    for row_number in range(2, INPUT_ROWS + 2):
        hidden.cell(row_number + 1, 1).value = sheet.cell(row_number, code_col).value
        hidden.cell(row_number + 1, 2).value = sheet.cell(row_number, name_col).value
        hidden.cell(row_number + 1, 3).value = sheet.cell(row_number, sale_col).value
    workbook.save(destination)
    workbook.close()

    selected: list[dict[str, object]] = []
    workbook = load_workbook(destination, read_only=True, data_only=False)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert sheet.max_row >= INPUT_ROWS + 1
        # One forward pass: in read-only mode every sheet.cell() call parses
        # the sheet XML again from the top.
        wanted = set(APPROVED_SOURCE_ROWS)
        saved_rows: dict[int, tuple[object, ...]] = {}
        for row_number, values in enumerate(
            sheet.iter_rows(max_col=40, values_only=True), start=1
        ):
            if row_number in wanted:
                saved_rows[row_number] = values
        for row_number in APPROVED_SOURCE_ROWS:
            values = saved_rows[row_number]
            selected.append(
                {
                    "row_number": row_number,
                    "pioneers_code": str(values[code_col - 1]),
                    "pioneers_name": str(values[name_col - 1]),
                    "old_price": str(values[sale_col - 1]),
                    "values": list(values),
                }
            )
    finally:
        workbook.close()
    assert len(selected) == 9
    return selected


def _product_codes(path: Path) -> list[str]:
    """Code of every product record on every sheet, including hidden catalogs."""
    workbook = load_workbook(path, data_only=False)
    found: list[str] = []
    try:
        for sheet in workbook.worksheets:
            header_row = None
            code_column = None
            for row_number in range(1, (sheet.max_row or 1) + 1):
                values = {
                    str(sheet.cell(row_number, column).value).strip(): column
                    for column in range(1, (sheet.max_column or 1) + 1)
                    if isinstance(sheet.cell(row_number, column).value, str)
                }
                if (
                    HEADER_CODE in values
                    and HEADER_NAME in values
                    and HEADER_SALE_PRICE in values
                ):
                    header_row = row_number
                    code_column = values[HEADER_CODE]
                    break
            if header_row is None or code_column is None:
                continue
            for row_number in range(header_row + 1, (sheet.max_row or header_row) + 1):
                value = sheet.cell(row_number, code_column).value
                if value is not None and str(value).strip():
                    found.append(str(value).strip())
        return found
    finally:
        workbook.close()


def _assert_package_omits(path: Path, token: str) -> None:
    import zipfile

    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith(".xml"):
                continue
            payload = archive.read(name).decode("utf-8", errors="replace")
            assert f">{token}<" not in payload, name


def _product_rows(path: Path) -> list[tuple[object, ...]]:
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook[PRODUCT_SHEET_NAME]
        assert not list(sheet.tables)
        assert "FullCatalog" not in workbook.defined_names
        assert sheet.auto_filter.ref is None
        rows = []
        for row_number in range(2, (sheet.max_row or 1) + 1):
            assert sheet.row_dimensions[row_number].hidden is not True
            values = tuple(
                sheet.cell(row_number, column).value for column in range(1, 41)
            )
            assert any(value not in (None, "") for value in values)
            rows.append(values)
        return rows
    finally:
        workbook.close()


def test_six_thousand_seven_hundred_rows_export_exactly_nine(
    tmp_path: Path,
) -> None:
    fixture_hash = _sha256(FIXTURE)
    source = tmp_path / "catalog.xlsx"
    selected = _expand_real_catalog(source)
    before = _sha256(source)
    output = tmp_path / "approved.xlsx"
    updates = [
        {
            "row_number": item["row_number"],
            "pioneers_name": item["pioneers_name"],
            "pioneers_code": item["pioneers_code"],
            "final_sale_price": Decimal(str(1_000_000 + index)),
        }
        for index, item in enumerate(selected)
    ]
    result = update_final_sale_prices(source, output, updates)

    assert result.updated_count == 9
    assert result.exported_row_count == 9
    rows = _product_rows(output)
    assert len(rows) == 9
    codes = [row[3] for row in rows]
    assert codes == [item["pioneers_code"] for item in selected]
    for index, item in enumerate(selected):
        exported = rows[index]
        for column, source_value in enumerate(item["values"]):
            if column == 36:
                assert exported[column] == str(1_000_000 + index)
                assert exported[column] != item["old_price"]
            else:
                assert exported[column] == source_value
    rejected_code = "200003"
    assert rejected_code not in codes
    saved = load_workbook(output)
    try:
        hidden = saved["کاتالوگ پنهان"]
        assert hidden.sheet_state == "veryHidden"
        assert hidden["A1"].value == "نسخه پنهان"
        assert hidden["A2"].value == HEADER_CODE
        assert hidden.max_row == 2
        assert all(hidden.cell(row, 1).value in (None, "") for row in range(3, 12))
    finally:
        saved.close()
    assert _product_codes(output) == codes
    _assert_package_omits(output, rejected_code)
    assert _sha256(source) == before
    assert _sha256(FIXTURE) == fixture_hash


def test_zero_updates_do_not_copy_the_catalog(tmp_path: Path) -> None:
    source = tmp_path / "catalog.xlsx"
    shutil.copy2(FIXTURE, source)
    before = _sha256(source)
    output = tmp_path / "should-not-exist.xlsx"
    with pytest.raises(PioneersExcelWriteError, match="At least one"):
        update_final_sale_prices(source, output, [])
    assert not output.exists()
    assert _sha256(source) == before


def test_same_code_on_two_rows_is_not_exported(tmp_path: Path) -> None:
    source = tmp_path / "dup.xlsx"
    shutil.copy2(FIXTURE, source)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    sheet["D3"].value = sheet["D2"].value
    sheet["E3"].value = "Other name"
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    output = tmp_path / "out.xlsx"
    with pytest.raises(DuplicateUpdateError, match="cannot be identified uniquely"):
        update_final_sale_prices(
            source,
            output,
            [
                {
                    "row_number": 2,
                    "pioneers_name": "Test Valve 20mm",
                    "pioneers_code": "200002",
                    "final_sale_price": "10",
                },
                {
                    "row_number": 3,
                    "pioneers_name": "Other name",
                    "pioneers_code": "200002",
                    "final_sale_price": "10",
                },
            ],
        )
    assert not output.exists()
    assert _sha256(source) == before


def test_html_xls_export_keeps_only_approved_rows(tmp_path: Path) -> None:
    source = tmp_path / "from-xls.xlsx"
    convert_pioneers_html_xls_to_xlsx(HTML_XLS, source)
    before = _sha256(source)
    original = load_workbook(source, data_only=False)
    try:
        sheet = original.active
        input_rows = sheet.max_row - 1
        assert input_rows > 2
        name = sheet["E2"].value
        code = sheet["D2"].value
    finally:
        original.close()
    output = tmp_path / "out.xlsx"
    result = update_final_sale_prices(
        source,
        output,
        [
            {
                "row_number": 2,
                "pioneers_name": name,
                "pioneers_code": code,
                "final_sale_price": "424242",
            }
        ],
    )
    assert result.updated_count == 1
    assert result.exported_row_count == 1
    exported = load_workbook(output)
    try:
        sheet = exported.active
        assert sheet.max_row == 2
        assert sheet["D2"].value == code
        assert sheet["E2"].value == name
        assert sheet["AK2"].value == "424242"
    finally:
        exported.close()
    assert _sha256(source) == before
    assert _sha256(HTML_XLS) == _sha256(HTML_XLS)


def test_service_exports_only_approved_rows_from_the_large_catalog(
    tmp_path: Path,
) -> None:
    fixture_hash = _sha256(FIXTURE)
    source = tmp_path / "catalog.xlsx"
    selected = _expand_real_catalog(source)
    before = _sha256(source)
    store = SessionStore()
    service = ProcessingService(store)
    session = service.create_session()
    with source.open("rb") as handle:
        service.upload_pioneers_excel(session.session_id, "catalog.xlsx", handle)
    session = service.get_session(session.session_id)
    assert len(session.products) >= INPUT_ROWS - 5
    session.supplier_key = "azin_pipe"
    session.supplier_confirmed = True
    candidates = []
    for index, item in enumerate(selected):
        candidates.append(
            UpdateCandidate(
                candidate_id=f"keep-{index}",
                row_number=int(item["row_number"]),
                pioneers_code=str(item["pioneers_code"]),
                pioneers_name=str(item["pioneers_name"]),
                supplier_code=f"S{index}",
                supplier_name=str(item["pioneers_name"]),
                old_price=Decimal(str(item["old_price"])),
                new_price=Decimal(1_500_000 + index),
                absolute_change=Decimal("100"),
                percentage_change=Decimal("10"),
                match_type=None,
                status=MatchStatus.MATCHED,
                review_decision=ReviewDecision.APPROVE,
                warnings=[],
                suggestions=[],
                source_currency="rial",
                price_basis="printed",
                source_unit=None,
            )
        )
    ignored_product = next(
        product
        for product in session.products
        if product.row_number not in APPROVED_SOURCE_ROWS and product.code
    )
    untouched = [
        product
        for product in session.products
        if product.row_number not in APPROVED_SOURCE_ROWS
        and product.code
        and product.code != ignored_product.code
        and product.row_number not in session.duplicate_codes.get(product.code, ())
    ]
    include_a, include_b, unchanged_product = untouched[:3]
    for index, product in enumerate((include_a, include_b)):
        candidates.append(
            UpdateCandidate(
                candidate_id=f"include-{index}",
                row_number=product.row_number,
                pioneers_code=product.code,
                pioneers_name=product.name,
                supplier_code=f"INC{index}",
                supplier_name=product.name or f"include-{index}",
                old_price=product.sale_price,
                new_price=(product.sale_price or Decimal("1000")) + Decimal("50"),
                absolute_change=Decimal("50"),
                percentage_change=Decimal("5"),
                match_type=None,
                status=MatchStatus.MATCHED,
                review_decision=ReviewDecision.INCLUDE,
                warnings=[],
                suggestions=[],
                source_currency="rial",
                price_basis="printed",
            )
        )
    candidates.append(
        UpdateCandidate(
            candidate_id="unchanged",
            row_number=unchanged_product.row_number,
            pioneers_code=unchanged_product.code,
            pioneers_name=unchanged_product.name,
            supplier_code="SAME",
            supplier_name=unchanged_product.name or "same",
            old_price=unchanged_product.sale_price,
            new_price=unchanged_product.sale_price,
            absolute_change=Decimal("0"),
            percentage_change=Decimal("0"),
            match_type=None,
            status=MatchStatus.MATCHED,
            review_decision=ReviewDecision.APPROVE,
            warnings=[],
            suggestions=[],
            source_currency="rial",
            price_basis="printed",
        )
    )
    candidates.append(
        UpdateCandidate(
            candidate_id="ignored",
            row_number=ignored_product.row_number,
            pioneers_code=ignored_product.code,
            pioneers_name=ignored_product.name,
            supplier_code="IGNORED",
            supplier_name=ignored_product.name or "ignored",
            old_price=ignored_product.sale_price,
            new_price=Decimal("999"),
            absolute_change=None,
            percentage_change=None,
            match_type=None,
            status=MatchStatus.IGNORED,
            review_decision=ReviewDecision.EXCLUDE,
            warnings=[],
            suggestions=[],
            source_currency="rial",
            price_basis="printed",
        )
    )
    candidates.append(
        UpdateCandidate(
            candidate_id="unmatched",
            row_number=None,
            pioneers_code=None,
            pioneers_name=None,
            supplier_code="UNMATCHED",
            supplier_name="no match",
            old_price=None,
            new_price=Decimal("999"),
            absolute_change=None,
            percentage_change=None,
            match_type=None,
            status=MatchStatus.UNMATCHED,
            review_decision=ReviewDecision.EXCLUDE,
            warnings=[],
            suggestions=[],
            source_currency="rial",
            price_basis="printed",
        )
    )
    session.candidates = candidates
    result = service.export_excel(session.session_id)
    assert result.updated_count == 9
    assert result.exported_row_count == 9
    assert session.output_path is not None
    rows = _product_rows(session.output_path)
    assert len(rows) == 9
    codes = {row[3] for row in rows}
    assert codes == {item["pioneers_code"] for item in selected}
    assert ignored_product.code not in codes
    assert include_a.code not in codes
    assert include_b.code not in codes
    assert unchanged_product.code not in codes
    _assert_package_omits(session.output_path, str(include_a.code))
    assert _sha256(source) == before
    assert _sha256(FIXTURE) == fixture_hash


def test_service_blocks_conflicting_prices_and_empty_selection(
    tmp_path: Path,
) -> None:
    source = tmp_path / "small.xlsx"
    shutil.copy2(FIXTURE, source)
    before = _sha256(source)
    store = SessionStore()
    service = ProcessingService(store)
    session = service.create_session()
    with source.open("rb") as handle:
        service.upload_pioneers_excel(session.session_id, "small.xlsx", handle)
    session = service.get_session(session.session_id)
    product = next(
        item
        for item in session.products
        if item.code and session.duplicate_codes.get(item.code, (item.row_number,)) == (item.row_number,)
    )
    session.supplier_key = "azin_pipe"
    session.supplier_confirmed = True

    def candidate(candidate_id: str, price: str, decision: ReviewDecision) -> UpdateCandidate:
        return UpdateCandidate(
            candidate_id=candidate_id,
            row_number=product.row_number,
            pioneers_code=product.code,
            pioneers_name=product.name,
            supplier_code=candidate_id,
            supplier_name=product.name or candidate_id,
            old_price=product.sale_price,
            new_price=Decimal(price),
            absolute_change=Decimal("1"),
            percentage_change=Decimal("1"),
            match_type=None,
            status=MatchStatus.MATCHED,
            review_decision=decision,
            warnings=[],
            suggestions=[],
            source_currency="rial",
            price_basis="printed",
        )

    session.candidates = [
        candidate("a", "2000", ReviewDecision.APPROVE),
        candidate("b", "3000", ReviewDecision.APPROVE),
    ]
    with pytest.raises(ValueError, match="conflicting final sale prices"):
        service.export_excel(session.session_id)
    assert session.output_path is None
    assert not (session.work_dir / "pioneers_updated.xlsx").exists()

    session.candidates = [candidate("only", "2000", ReviewDecision.EXCLUDE)]
    with pytest.raises(ValueError, match="no approved updates"):
        service.export_excel(session.session_id)
    assert not (session.work_dir / "pioneers_updated.xlsx").exists()

    session.candidates = [candidate("plain", "2000", ReviewDecision.INCLUDE)]
    with pytest.raises(ValueError, match="no approved updates"):
        service.export_excel(session.session_id)
    assert not (session.work_dir / "pioneers_updated.xlsx").exists()

    same_price = str(product.sale_price)
    session.candidates = [candidate("same", same_price, ReviewDecision.APPROVE)]
    with pytest.raises(ValueError, match="no approved updates"):
        service.export_excel(session.session_id)
    assert not (session.work_dir / "pioneers_updated.xlsx").exists()
    assert _sha256(source) == before


def _approve_one_changed_row(client, source: Path) -> tuple[str, object]:
    created = client.post("/api/sessions")
    assert created.status_code == 200, created.text
    session_id = created.json()["session_id"]
    with source.open("rb") as handle:
        uploaded = client.post(
            f"/api/sessions/{session_id}/pioneers-excel",
            files={
                "file": (
                    "p.xlsx",
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert uploaded.status_code == 200, uploaded.text
    session = client.app.state.processing.get_session(session_id)
    product = next(
        item
        for item in session.products
        if item.code
        and item.sale_price is not None
        and session.duplicate_codes.get(item.code, (item.row_number,))
        == (item.row_number,)
    )
    session.supplier_confirmed = True
    session.candidates = [
        UpdateCandidate(
            candidate_id="approved",
            row_number=product.row_number,
            pioneers_code=product.code,
            pioneers_name=product.name,
            supplier_code="API",
            supplier_name=product.name or "API",
            old_price=product.sale_price,
            new_price=product.sale_price + Decimal("50"),
            absolute_change=Decimal("50"),
            percentage_change=Decimal("1"),
            match_type=None,
            status=MatchStatus.MATCHED,
            review_decision=ReviewDecision.APPROVE,
            warnings=[],
            suggestions=[],
            source_currency="rial",
            price_basis="printed",
        )
    ]
    return session_id, product


def test_api_drops_partial_hidden_catalog(client, tmp_path: Path) -> None:
    source = tmp_path / "partial.xlsx"
    shutil.copy2(FIXTURE, source)
    workbook = load_workbook(source)
    hidden = workbook.create_sheet("هویت ناقص")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "یادداشت"
    hidden["A2"] = HEADER_CODE
    hidden["B2"] = HEADER_NAME
    hidden["C2"] = "قیمت"
    hidden["A3"] = "ARCH-0001"
    hidden["B3"] = "Arch product"
    hidden["C3"] = "10"
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    session_id, _product = _approve_one_changed_row(client, source)
    exported = client.post(f"/api/sessions/{session_id}/export")
    assert exported.status_code == 200, exported.text
    assert exported.json()["exported_row_count"] == 1
    download = client.get(f"/api/sessions/{session_id}/download")
    assert download.status_code == 200, download.text
    assert download.headers["x-exported-row-count"] == "1"
    saved = load_workbook(BytesIO(download.content))
    try:
        hidden = saved["هویت ناقص"]
        assert hidden.sheet_state == "veryHidden"
        assert hidden["A2"].value == HEADER_CODE
        assert hidden.max_row == 2
        for sheet in saved.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.value != "ARCH-0001"
        assert saved[PRODUCT_SHEET_NAME].max_row == 2
    finally:
        saved.close()
    assert b"ARCH-0001" not in download.content
    assert _sha256(source) == before


def test_api_exports_workbook_with_spaced_headers(client, tmp_path: Path) -> None:
    source = tmp_path / "spaced.xlsx"
    shutil.copy2(FIXTURE, source)
    workbook = load_workbook(source)
    sheet = workbook[PRODUCT_SHEET_NAME]
    code_header = f" {sheet['D1'].value} "
    name_header = f" {sheet['E1'].value} "
    sheet["D1"] = code_header
    sheet["E1"] = name_header
    workbook.save(source)
    workbook.close()
    before = _sha256(source)
    session_id, product = _approve_one_changed_row(client, source)
    exported = client.post(f"/api/sessions/{session_id}/export")
    assert exported.status_code == 200, exported.text
    body = exported.json()
    assert body["exported_row_count"] == 1
    assert body["updated_count"] == 1
    download = client.get(f"/api/sessions/{session_id}/download")
    assert download.status_code == 200, download.text
    assert download.headers["x-exported-row-count"] == "1"
    saved = load_workbook(BytesIO(download.content))
    try:
        exported_sheet = saved[PRODUCT_SHEET_NAME]
        assert exported_sheet["D1"].value == code_header
        assert exported_sheet["E1"].value == name_header
        assert str(exported_sheet["D2"].value) == str(product.code)
        assert exported_sheet.max_row == 2
    finally:
        saved.close()
    assert _sha256(source) == before
