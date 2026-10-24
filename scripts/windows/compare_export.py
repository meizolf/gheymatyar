#!/usr/bin/env python3
"""Compare a Pioneers source workbook to a GheymatYar export.

Asserts:
  * source SHA-256 matches the caller-supplied pre-export hash (immutability)
  * non-product worksheets are unchanged
  * the product sheet keeps its headers and contains only approved product rows
  * within those rows, only ``نرخ فروش اصلی`` may differ
  * at least one sale-price cell changed

For HTML-based ``.xls`` sources, pass ``--html-xls`` to convert with the same
converter the app uses, then compare that workbook to the export.
"""

from __future__ import annotations

import argparse
import hashlib
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

SALE_HEADER = "نرخ فروش اصلی"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cell_key(cell: Cell) -> tuple[object, str | None]:
    """Return (value_or_formula, data_type) for strict comparison."""
    return (cell.value, cell.data_type)


def _header_row(ws: Worksheet) -> list[str | None]:
    cells = next(ws.iter_rows(min_row=1, max_row=1), ())
    return [None if c.value is None else str(c.value) for c in cells]


def _product_sheet_name(wb) -> str:
    matches: list[str] = []
    for name in wb.sheetnames:
        headers = _header_row(wb[name])
        if SALE_HEADER in headers:
            matches.append(name)
    if len(matches) != 1:
        raise SystemExit(
            f"expected exactly one product sheet with {SALE_HEADER!r}, "
            f"found {matches!r}"
        )
    return matches[0]


_CODE_HEADERS = {"کد کالا", "کد"}


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _compare_exact_sheet(src_ws: Worksheet, out_ws: Worksheet) -> None:
    src_dims = (src_ws.max_row or 0, src_ws.max_column or 0)
    out_dims = (out_ws.max_row or 0, out_ws.max_column or 0)
    if src_dims != out_dims:
        raise SystemExit(
            f"sheet {src_ws.title!r} dimensions changed: {src_dims} -> {out_dims}"
        )
    max_row, max_col = src_dims
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            src_cell = src_ws.cell(row=row, column=col)
            out_cell = out_ws.cell(row=row, column=col)
            if _cell_key(src_cell) != _cell_key(out_cell):
                raise SystemExit(
                    f"cell changed at {src_ws.title}!R{row}C{col}: "
                    f"{_cell_key(src_cell)!r} -> {_cell_key(out_cell)!r}"
                )


def _compare_product_sheet(src_ws: Worksheet, out_ws: Worksheet) -> int:
    """Export rows are the approved subset. Only the sale column may differ."""
    src_headers = _header_row(src_ws)
    out_headers = _header_row(out_ws)
    if src_headers != out_headers:
        raise SystemExit(
            f"cell changed at {src_ws.title}!R1: headers "
            f"{src_headers!r} -> {out_headers!r}"
        )
    src_cols = src_ws.max_column or 0
    out_cols = out_ws.max_column or 0
    if src_cols != out_cols:
        raise SystemExit(
            f"sheet {src_ws.title!r} dimensions changed: "
            f"columns {src_cols} -> {out_cols}"
        )
    sale_col = src_headers.index(SALE_HEADER) + 1
    code_col = next(
        (index + 1 for index, header in enumerate(src_headers) if header in _CODE_HEADERS),
        None,
    )
    out_max = out_ws.max_row or 1
    src_max = src_ws.max_row or 1
    if out_max > src_max:
        raise SystemExit(
            f"sheet {src_ws.title!r} dimensions changed: "
            f"rows {src_max} -> {out_max}"
        )

    by_code: dict[str, list[int]] = {}
    for row in range(2, src_max + 1):
        if code_col is None:
            break
        value = src_ws.cell(row, code_col).value
        if _blank(value):
            continue
        by_code.setdefault(str(value).strip(), []).append(row)

    used: set[int] = set()
    changed = 0
    for row in range(2, out_max + 1):
        if row in out_ws.row_dimensions and out_ws.row_dimensions[row].hidden:
            raise SystemExit(f"hidden product row {out_ws.title}!R{row}")
        values = [out_ws.cell(row, col).value for col in range(1, out_cols + 1)]
        if all(_blank(value) for value in values):
            raise SystemExit(f"blank product row {out_ws.title}!R{row}")
        source_row = _match_source_row(
            src_ws,
            out_ws,
            row,
            sale_col=sale_col,
            code_col=code_col,
            by_code=by_code,
            used=used,
            max_col=out_cols,
        )
        used.add(source_row)
        for col in range(1, out_cols + 1):
            src_cell = src_ws.cell(source_row, col)
            out_cell = out_ws.cell(row, col)
            if col == sale_col:
                if out_cell.data_type == "f" and src_cell.data_type != "f":
                    raise SystemExit(
                        f"sale cell became a formula at "
                        f"{out_ws.title}!R{row}C{col}: {out_cell.value!r}"
                    )
                if _cell_key(src_cell) != _cell_key(out_cell):
                    changed += 1
                continue
            if _cell_key(src_cell) != _cell_key(out_cell):
                raise SystemExit(
                    f"cell changed at {out_ws.title}!R{row}C{col}: "
                    f"{_cell_key(src_cell)!r} -> {_cell_key(out_cell)!r}"
                )
    return changed


def _match_source_row(
    src_ws: Worksheet,
    out_ws: Worksheet,
    row: int,
    *,
    sale_col: int,
    code_col: int | None,
    by_code: dict[str, list[int]],
    used: set[int],
    max_col: int,
) -> int:
    if code_col is not None:
        code = out_ws.cell(row, code_col).value
        if not _blank(code):
            matches = [
                source_row
                for source_row in by_code.get(str(code).strip(), [])
                if source_row not in used
            ]
            if len(matches) == 1:
                return matches[0]
            raise SystemExit(
                "selected Pioneers product cannot be identified uniquely: "
                f"{code!r} at {out_ws.title}!R{row}"
            )
    signature = tuple(
        out_ws.cell(row, col).value for col in range(1, max_col + 1) if col != sale_col
    )
    matches = []
    for source_row in range(2, (src_ws.max_row or 1) + 1):
        if source_row in used:
            continue
        source_signature = tuple(
            src_ws.cell(source_row, col).value
            for col in range(1, max_col + 1)
            if col != sale_col
        )
        if source_signature == signature:
            matches.append(source_row)
    if len(matches) != 1:
        raise SystemExit(
            f"cell changed at {out_ws.title}!R{row}: "
            "export row does not match one source product"
        )
    return matches[0]


def _compare_workbooks(source: Path, exported: Path) -> int:
    src_wb = load_workbook(source, data_only=False)
    out_wb = load_workbook(exported, data_only=False)
    try:
        if src_wb.sheetnames != out_wb.sheetnames:
            raise SystemExit(
                f"worksheet names changed\n"
                f"  source: {src_wb.sheetnames}\n"
                f"  export: {out_wb.sheetnames}"
            )

        product_sheet = _product_sheet_name(src_wb)
        changed = 0
        for sheet_name in src_wb.sheetnames:
            if sheet_name == product_sheet:
                changed += _compare_product_sheet(src_wb[sheet_name], out_wb[sheet_name])
            else:
                _compare_exact_sheet(src_wb[sheet_name], out_wb[sheet_name])
        return changed
    finally:
        src_wb.close()
        out_wb.close()


def _prepare_source(source: Path, *, html_xls: bool) -> Path:
    if not html_xls:
        return source
    # Import lazily so non-HTML compares do not need the app on PYTHONPATH.
    from gheymatyar.parsers.pioneers_html_xls import (  # noqa: WPS433
        convert_pioneers_html_xls_to_xlsx,
    )

    tmp = Path(tempfile.mkdtemp(prefix="gheymatyar-html-xls-")) / "converted.xlsx"
    convert_pioneers_html_xls_to_xlsx(source, tmp)
    return tmp


def compare(
    source: Path,
    exported: Path,
    expected_source_sha: str,
    *,
    html_xls: bool = False,
) -> None:
    actual = sha256(source)
    if actual != expected_source_sha.lower():
        raise SystemExit(
            f"source hash changed during processing\n"
            f"  expected: {expected_source_sha}\n"
            f"  actual:   {actual}"
        )

    compare_source = _prepare_source(source, html_xls=html_xls)
    changed = _compare_workbooks(compare_source, exported)
    if changed < 1:
        raise SystemExit("export did not change any نرخ فروش اصلی cells")

    print(
        f"OK compare_export: source immutable sha={actual[:12]}…; "
        f"sheets compared; formulas checked; sale-price cells changed={changed}"
        + (" (html-xls converted)" if html_xls else "")
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--exported", type=Path, required=True)
    parser.add_argument("--source-sha256", required=True)
    parser.add_argument(
        "--html-xls",
        action="store_true",
        help="Treat --source as Pioneers HTML .xls and convert before compare",
    )
    args = parser.parse_args(argv)
    compare(
        args.source,
        args.exported,
        args.source_sha256,
        html_xls=args.html_xls,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
