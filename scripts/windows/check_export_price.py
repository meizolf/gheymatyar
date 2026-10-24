#!/usr/bin/env python3
"""Verify a single approved export product by code, independent of source row."""

from __future__ import annotations

import argparse
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

SALE_HEADER = "نرخ فروش اصلی"
CODE_HEADERS = {"کد کالا", "کد"}


def check(exported: Path, code: str, price: str) -> None:
    wb = load_workbook(exported, read_only=True, data_only=False)
    try:
        sheets = [
            ws
            for ws in wb
            if SALE_HEADER in [
                c.value for c in next(ws.iter_rows(min_row=1, max_row=1), ())
            ]
        ]
        if len(sheets) != 1:
            raise SystemExit("expected exactly one product sheet")
        ws = sheets[0]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        code_idx = next(
            (i for i, header in enumerate(headers) if header in CODE_HEADERS), None
        )
        if code_idx is None:
            raise SystemExit("missing product code column")
        sale_idx = headers.index(SALE_HEADER)
        rows = list(ws.iter_rows(min_row=2))
        if len(rows) != 1:
            raise SystemExit(f"expected exactly one approved product, got {len(rows)}")
        row = next(
            (row for row in rows if str(row[code_idx].value).strip() == code.strip()),
            None,
        )
        if row is None:
            raise SystemExit(f"exported product code does not match approved code {code!r}")
        cell = row[sale_idx]
        if cell.data_type == "f":
            raise SystemExit("exported sale price must not be a formula")
        try:
            got = Decimal(str(cell.value).replace(",", "").strip())
            want = Decimal(price.replace(",", "").strip())
        except InvalidOperation as exc:
            raise SystemExit("invalid sale price") from exc
        if not got.is_finite() or not want.is_finite() or got <= 0 or got != want:
            raise SystemExit(f"sale price mismatch for code {code}: {got!r} != {want!r}")
    finally:
        wb.close()
    print(f"OK check_export_price: only approved code {code} == {want}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exported", type=Path, required=True)
    parser.add_argument("--code", required=True)
    parser.add_argument("--price", required=True)
    args = parser.parse_args(argv)

    check(args.exported, args.code, args.price)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
