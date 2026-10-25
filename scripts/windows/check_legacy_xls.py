#!/usr/bin/env python3
"""Check a GheymatYar Excel 97–2003 ``.xls`` export against its source file.

Reads the export with ``xlrd`` only and does not import the application, so
it checks the app's writer independently.

Asserts:
  * the export starts with the OLE2 signature, is not HTML or a zip/.xlsx,
    and ``xlrd`` reads it as a single-sheet BIFF8 (Excel 97–2003) workbook
  * the sheet has exactly three text headers, in this order:
    ``کد کالا``, ``نام کالا``, ``نرخ فروش اصلی``
  * the export holds exactly ``--rows`` product rows and no other cells
  * every code is text, appears once, and is a code in the source file
  * every name is text and equals that product's source name exactly
  * every price is a positive whole number (Rial)
  * the product with ``--code`` has the price ``--price``
  * the source SHA-256 matches ``--source-sha256`` when given
"""

from __future__ import annotations

import argparse
import hashlib
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from pathlib import Path

import xlrd

SALE_HEADER = "نرخ فروش اصلی"
CODE_HEADER = "کد کالا"
NAME_HEADER = "نـــام کــالا"  # as Pioneers spells it in the source
EXPORT_HEADERS = ("کد کالا", "نام کالا", "نرخ فروش اصلی")
OLE2 = bytes.fromhex("D0CF11E0A1B11AE1")


class _Cells(HTMLParser):
    """Cell text of a Pioneers HTML-based .xls source."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: object) -> None:
        if tag == "tr":
            self.rows.append([])
        elif tag == "td":
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        if tag == "td" and self._cell is not None:
            self.rows[-1].append("".join(self._cell))
            self._cell = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def _source_rows(source: Path) -> list[list[str]]:
    raw = source.read_bytes()
    if raw.startswith(b"\xef\xbb\xbf") or raw.lstrip()[:6].lower() == b"<html>":
        parser = _Cells()
        parser.feed(raw.decode("utf-8-sig"))
        parser.close()
        return parser.rows
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        sheets = [
            ws
            for ws in workbook
            if SALE_HEADER in [c.value for c in next(ws.iter_rows(max_row=1), ())]
        ]
        if len(sheets) != 1:
            raise SystemExit("expected exactly one source product sheet")
        rows = [
            ["" if value is None else str(value) for value in row]
            for row in sheets[0].iter_rows(values_only=True)
        ]
    finally:
        workbook.close()
    width = max(i + 1 for i, value in enumerate(rows[0]) if value)
    return [row[:width] + [""] * (width - len(row[:width])) for row in rows]


def check(
    exported: Path,
    source: Path,
    *,
    code: str,
    price: str,
    rows: int,
    source_sha256: str | None,
) -> None:
    if source_sha256 is not None:
        if hashlib.sha256(source.read_bytes()).hexdigest() != source_sha256:
            raise SystemExit("source file changed during export")
    raw = exported.read_bytes()
    if raw.startswith(b"PK") or raw.lstrip(b"\xef\xbb\xbf \r\n\t")[:1] == b"<":
        raise SystemExit("export is a zip/.xlsx or HTML file, not an Excel 97–2003 .xls")
    if not raw.startswith(OLE2):
        raise SystemExit("export does not start with the OLE2 signature")
    try:
        book = xlrd.open_workbook(file_contents=raw)
    except xlrd.XLRDError as exc:
        raise SystemExit(f"xlrd cannot read the export: {exc}") from exc
    if book.biff_version != 80:
        raise SystemExit(f"export is BIFF {book.biff_version}, expected BIFF8 (80)")
    if book.nsheets != 1:
        raise SystemExit(f"expected one worksheet, found {book.nsheets}")
    sheet = book.sheet_by_index(0)
    if sheet.ncols != len(EXPORT_HEADERS):
        raise SystemExit(f"expected exactly 3 columns, found {sheet.ncols}")
    header = sheet.row(0)
    if [cell.value for cell in header] != list(EXPORT_HEADERS) or any(
        cell.ctype != xlrd.XL_CELL_TEXT for cell in header
    ):
        raise SystemExit(f"header must be {EXPORT_HEADERS!r} as text")
    if sheet.nrows - 1 != rows:
        raise SystemExit(f"expected exactly {rows} product rows, got {sheet.nrows - 1}")

    src_rows = _source_rows(source)
    headers = [value.strip() for value in src_rows[0]]
    src_code, src_name = headers.index(CODE_HEADER), headers.index(NAME_HEADER)
    names = {row[src_code].strip(): row[src_name] for row in src_rows[1:]}
    seen: set[str] = set()
    found = None
    for row_index in range(1, sheet.nrows):
        code_cell, name_cell, price_cell = sheet.row(row_index)
        if code_cell.ctype != xlrd.XL_CELL_TEXT or not code_cell.value:
            raise SystemExit(f"row {row_index + 1}: product code must be text")
        product = code_cell.value
        if product in seen:
            raise SystemExit(f"code {product} appears more than once")
        seen.add(product)
        if product not in names:
            raise SystemExit(f"exported code {product!r} is not in the source")
        if name_cell.ctype != xlrd.XL_CELL_TEXT or name_cell.value != names[product]:
            raise SystemExit(
                f"code {product}: name {name_cell.value!r} is not the source name "
                f"{names[product]!r}"
            )
        value = price_cell.value
        if price_cell.ctype != xlrd.XL_CELL_NUMBER or value <= 0 or not float(value).is_integer():
            raise SystemExit(f"code {product}: price must be a positive whole Rial number")
        if product == code.strip():
            found = int(value)
    if found is None:
        raise SystemExit(f"approved code {code!r} is not in the export")
    try:
        want = Decimal(price.replace(",", "").strip())
    except InvalidOperation as exc:
        raise SystemExit("invalid sale price") from exc
    if Decimal(found) != want:
        raise SystemExit(f"sale price mismatch for code {code}: {found} != {price!r}")
    print(f"OK check_legacy_xls: 3-column BIFF8 .xls, {rows} row(s); code {code} == {want}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exported", type=Path, required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--code", required=True)
    parser.add_argument("--price", required=True)
    parser.add_argument("--rows", type=int, default=1)
    parser.add_argument("--source-sha256")
    args = parser.parse_args(argv)
    check(
        args.exported,
        args.source,
        code=args.code,
        price=args.price,
        rows=args.rows,
        source_sha256=args.source_sha256,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
